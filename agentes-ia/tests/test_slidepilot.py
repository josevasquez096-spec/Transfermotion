from pptx import Presentation

from suite import samples
from suite.datapilot.logic import clean
from suite.slidepilot import logic as L
from suite.slidepilot.strings import SAMPLE_OUTLINE


def test_outline_roundtrip_all_languages():
    for lang, text in SAMPLE_OUTLINE.items():
        deck = L.parse_outline(text)
        assert deck["title"] and len(deck["slides"]) == 7, lang
        assert L.parse_outline(L.deck_to_outline(deck)) == deck, lang


def test_special_layouts_parse():
    deck = L.parse_outline(SAMPLE_OUTLINE["en"])
    by_layout = {s["layout"]: s for s in deck["slides"]}
    assert by_layout["kpis"]["kpis"][0] == {"value": "$694K", "label": "Revenue"}
    chart = by_layout["chart"]["chart"]
    assert chart["categories"] == ["North", "South", "East", "West"]
    assert [s["name"] for s in chart["series"]] == ["2024", "2025"]
    assert by_layout["columns"]["right_title"] == "With the CRM"
    assert by_layout["timeline"]["steps"][0] == "Q1: Discovery"


def test_build_pptx(tmp_path):
    deck = L.parse_outline(SAMPLE_OUTLINE["es"])
    for theme in L.THEMES:
        path = L.build_pptx(deck, tmp_path / f"{theme}.pptx", theme, "es", "Acme", contact="hola@acme.com")
        prs = Presentation(path)
        # portada + agenda + 7 + cierre
        assert len(prs.slides) == 10
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(charts) == 1 and len(charts[0].chart.plots[0].series) == 2
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    assert any("ERP" in n for n in notes)


def test_excel_and_word_to_outline():
    df = clean(samples.sales("fr"))[0]
    deck = L.parse_outline(L.excel_to_outline(df, "Région", ["Total"], "bar", top=3, lang="fr"))
    assert [s["layout"] for s in deck["slides"]] == ["kpis", "chart", "table"]
    assert len(deck["slides"][1]["chart"]["categories"]) == 3
    outline = L.markup_to_outline("# Doc\n## A\n" + "Una frase. " * 30 + "\n- b1\n- b2\n- b3\n- b4\n- b5\n- b6", "es")
    deck = L.parse_outline(outline)
    assert deck["title"] == "Doc" and all(len(s["bullets"]) <= 6 for s in deck["slides"])
