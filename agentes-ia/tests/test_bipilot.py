import json
from datetime import date

from openpyxl import load_workbook

from suite import samples
from suite.bipilot import logic as L


def test_roles_relationships_and_orphans():
    tables = samples.bi_model("en")
    assert L.table_role(tables["Sales"]) == "fact"
    assert L.table_role(tables["Products"]) == "dimension"
    rels = L.suggest_relationships(tables)
    pairs = {(r["one_table"], r["one_col"], r["many_table"], r["many_col"]) for r in rels}
    assert pairs == {("Products", "ProductID", "Sales", "ProductID"), ("Customers", "CustomerID", "Sales", "CustomerID")}
    issues = L.quality_issues(tables, rels)
    assert ("high", "q_orphans", {"table": "Sales", "col": "ProductID", "n": 3, "dim": "Products"}) in issues


def test_dax_measures():
    dax = L.dax_measures("Sales", "Amount", "en", "'Calendar'[Date]", "Products", "Product")
    assert "Total Amount = SUM ( 'Sales'[Amount] )" in dax
    assert "TOTALYTD ( [Total Amount], 'Calendar'[Date] )" in dax
    assert "RANKX ( ALLSELECTED ( 'Products'[Product] )" in dax
    assert "TOTALYTD" not in L.dax_measures("Sales", "Amount", "en")


def test_calendar_outputs():
    start, end = date(2024, 1, 1), date(2024, 12, 31)
    df = L.calendar_df(start, end, "es", fiscal_start=7)
    assert len(df) == 366 and df["Mes"].iloc[0] == "enero"
    assert df.loc[df["Fecha"] == "2024-07-01", "Año fiscal"].item() == "AF 2025"
    dax = L.calendar_dax(start, end, "fr", 7)
    assert dax.startswith("Calendrier =") and '"juillet"' in dax and "CALENDAR ( _Start, _End )" in dax
    m = L.calendar_m(start, end, "en")
    assert m.count("let") == 1 and '"en-US"' in m


def test_m_query_and_theme():
    tables = samples.bi_model("es")
    m = L.m_clean_query(tables["Ventas"], "Ventas", r"C:\data\v.xlsx")
    assert '{"Fecha", type date}' in m and '{"Importe", type number}' in m and "Table.Distinct" in m
    theme = json.loads(L.theme_json("T", L.palette_from("#2B579A")))
    assert len(theme["dataColors"]) == 8 and theme["dataColors"][0] == "#2B579A"


def test_export_model(tmp_path):
    tables = samples.bi_model("es")
    rels = L.suggest_relationships(tables)
    labels = {k: k for k in ("doc", "doc_title", "table", "column", "type", "role", "role_fact", "role_dimension",
                             "role_other", "relationships", "one", "many", "cardinality", "coverage", "issues",
                             "calendar")}
    path = L.export_model(tmp_path / "m.xlsx", tables, rels, "ok", labels, L.calendar_df(date(2024, 1, 1),
                                                                                       date(2024, 1, 31)), "X = 1")
    assert load_workbook(path).sheetnames == ["doc", "DAX", "Ventas", "Productos", "Clientes", "calendar"]
