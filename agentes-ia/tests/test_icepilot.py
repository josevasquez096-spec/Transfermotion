import json
import re
import shutil
import subprocess

import pytest
from docx import Document
from pptx import Presentation

from suite.icepilot import content, games, office
from suite.icepilot.build import build_html
from suite.icepilot.web_text import RULES, UI


def config(game, lang="es", pack="fraud", **options):
    cfg = games.new_config(game, lang)
    cfg["seed"] = 42
    cfg["title"], cfg["subtitle"] = content.pack_titles(pack, lang)
    cfg["items"] = content.pack_items(pack, game, lang, games.GAMES[game]["columns"])
    cfg["options"].update(options)
    return cfg


def payload(page):
    blob = re.search(r'<script type="application/json" id="game-data">(.*?)</script>', page, re.S).group(1)
    return json.loads(blob)


@pytest.mark.parametrize("lang", ["es", "en", "fr"])
@pytest.mark.parametrize("pack", ["fraud", "team"])
def test_every_game_builds_offline(lang, pack):
    for game in games.GAMES:
        cfg = config(game, lang, pack)
        if game == "bingo":
            cfg["options"]["cards"] = 20
        page, warnings = build_html(cfg)
        assert not warnings, (game, warnings)
        assert "Content-Security-Policy" in page and "connect-src 'none'" in page
        external = set(re.findall(r"https?://[^\s\"'<>)]+", page)) - {"http://www.w3.org/2000/svg"}
        assert not external, (game, external)
        data = payload(page)
        assert data["type"] == game and data["ui"]["start"] == UI[lang]["start"]


def test_every_pack_has_content_for_every_game():
    for pack in content.PACKS:
        for lang in ("es", "en", "fr"):
            for game, spec in games.GAMES.items():
                rows = content.pack_items(pack, game, lang, spec["columns"])
                assert rows and set(rows[0]) == set(spec["columns"]), (pack, lang, game)


def test_texts_complete():
    for lang in ("es", "en", "fr"):
        assert set(UI[lang]) == set(UI["es"])
    for game in games.GAMES:
        assert set(RULES[game]) == {"es", "en", "fr"}


@pytest.mark.skipif(not shutil.which("node"), reason="node no disponible")
def test_javascript_syntax(tmp_path):
    seen = set()
    for game, spec in games.GAMES.items():
        if spec["engine"] in seen:
            continue
        seen.add(spec["engine"])
        page, _ = build_html(config(game))
        script = page.split("<script>")[-1].split("</script>")[0]
        js = tmp_path / f"{game}.js"
        js.write_text(script, encoding="utf-8")
        result = subprocess.run(["node", "--check", str(js)], capture_output=True, text=True)
        assert result.returncode == 0, (game, result.stderr)


def test_prepare_validation():
    cfg = games.new_config("quiz")
    cfg["items"] = [
        {"question": "¿2+2?", "a": "3", "b": "4", "correct": "B"},
        {"question": "Abierta", "correct": "Respuesta libre"},
        {"question": "Mala", "a": "x", "b": "y", "correct": "Z"},
        {"question": "Una sola", "a": "x", "correct": "A"},
        {"question": "Por texto", "a": "rojo", "b": "azul", "correct": "azul"},
    ]
    data, warnings = games.prepare(cfg)
    assert [q["q"] for q in data["questions"]] == ["¿2+2?", "Abierta", "Por texto"]
    assert data["questions"][0]["correct"] == 1 and data["questions"][1]["options"] == []
    assert data["questions"][2]["correct"] == 1
    assert [k for k, _ in warnings] == ["w_correct", "w_options"]
    cfg = games.new_config("truefalse", "fr")
    cfg["items"] = [{"statement": "a", "answer": "Vrai"}, {"statement": "b", "answer": "faux"},
                    {"statement": "c", "answer": "?"}]
    data, warnings = games.prepare(cfg)
    assert [q["correct"] for q in data["questions"]] == [0, 1] and data["questions"][0]["options"] == ["Vrai", "Faux"]
    assert warnings == [("w_truefalse", {"n": 3})]
    _, warnings = games.prepare(games.new_config("wheel"))
    assert warnings[0][0] == "w_empty"


def test_snakes_board_is_valid():
    import random

    for cells in ("30", "50", "64", "100"):
        for n in (0, 3, 6, 12):
            b = games.snakes_board(cells, n, n, random.Random(n))
            ends = [x for pair in b["snakes"] + b["ladders"] for x in pair]
            assert len(ends) == len(set(ends)) and 1 not in ends and b["total"] not in ends
            assert all(a < z for a, z in b["ladders"]) and all(a > z for a, z in b["snakes"])
            assert all(1 < x < b["total"] for x in ends)


def test_bingo_cards_unique_with_free_center():
    data, warnings = games.prepare(config("bingo", cards=40))
    assert not warnings and len(data["cards"]) == 40
    assert len({tuple(c) for c in data["cards"]}) == 40
    assert all(c[12] == -1 and len(set(c)) == 25 for c in data["cards"])
    data, warnings = games.prepare({**config("bingo"), "items": [{"term": "x"}] * 5})
    assert warnings[0][0] == "w_bingo_terms"


def test_word_search_placements_match_grid():
    for lang in ("es", "en", "fr"):
        data, warnings = games.prepare(config("wordsearch", lang, reverse=True))
        assert not warnings
        grid = data["grid"]
        for p in data["placements"]:
            letters = "".join(grid[p["r"] + p["dr"] * i][p["c"] + p["dc"] * i] for i in range(len(p["word"])))
            assert letters == p["word"]
        assert {p["word"] for p in data["placements"]} == {w["word"] for w in data["words"]}


def test_scramble_changes_words():
    data, _ = games.prepare(config("scramble"))
    assert all(r["prompt"].replace(" / ", " ") != r["answer"].upper() for r in data["rounds"])


def test_office_exports(tmp_path):
    cfg = config("quiz")
    path, _ = office.build_pptx(cfg, tmp_path / "q.pptx")
    prs = Presentation(path)
    n = len(cfg["items"])
    assert len(prs.slides) == 2 + 2 * n + 1  # portada, reglas, pregunta+respuesta, cierre
    assert "Respuesta:" in prs.slides[2].notes_slide.notes_text_frame.text
    cfg = config("bingo", cards=6)
    path, _ = office.build_docx(cfg, tmp_path / "b.docx")
    doc = Document(path)
    assert len(doc.tables) == 1 + 6 * 2  # contenido + (cabecera BINGO + cartón) × 6
    for game, spec in games.GAMES.items():
        office.build_docx(config(game), tmp_path / f"{game}.docx")
        if spec["pptx"]:
            office.build_pptx(config(game), tmp_path / f"{game}.pptx")


def test_browser_play(tmp_path):
    """Juega a todos los juegos en Chromium (solo si Playwright está instalado)."""
    playwright = pytest.importorskip("playwright.sync_api")
    import os

    exe = os.environ.get("PILOT_CHROMIUM") or ("/opt/pw-browsers/chromium" if os.path.exists("/opt/pw-browsers/chromium")
                                                else None)
    with playwright.sync_playwright() as p:
        try:
            browser = p.chromium.launch(executable_path=exe) if exe else p.chromium.launch()
        except Exception as exc:  # navegador no instalado
            pytest.skip(str(exc))
        for game in games.GAMES:
            page = browser.new_page(viewport={"width": 1366, "height": 768})
            errors, requests = [], []
            page.on("pageerror", lambda e: errors.append(str(e)))
            page.on("request", lambda r: requests.append(r.url) if not r.url.startswith(("file:", "data:")) else None)
            html, _ = build_html(config(game, cards=5) if game == "bingo" else config(game))
            file = tmp_path / f"{game}.html"
            file.write_text(html, encoding="utf-8")
            page.goto(file.as_uri())
            start = page.locator(".intro .btn.primary.big")
            if start.count():
                start.first.click()
            page.wait_for_timeout(200)
            assert not errors and not requests, (game, errors, requests)
            page.close()
        browser.close()
