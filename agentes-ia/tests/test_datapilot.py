import pandas as pd
import pytest
from openpyxl import load_workbook

from suite import samples
from suite.datapilot import logic as L
from suite.tabular import read_table


@pytest.fixture
def sales():
    return L.clean(samples.sales("es"))[0]


@pytest.mark.parametrize("text,value", [("1.234,56", 1234.56), ("$ 1,234.5", 1234.5), ("(300)", -300.0),
                                        ("1.234.567", 1234567.0), ("12", 12.0), ("abc", None), ("", None)])
def test_parse_number(text, value):
    assert L.parse_number(text) == value


def test_clean_report_and_types():
    df, report = L.clean(samples.sales("es"))
    assert report["duplicates"] == 6 and report["empty_rows"] == 3 and report["unified"] > 0
    assert set(report["numbers"]) == {"Unidades", "Precio", "Total"} and report["dates"] == ["Fecha"]
    assert pd.api.types.is_datetime64_any_dtype(df["Fecha"])
    assert sorted(df["Región"].unique()) == ["Centro", "Este", "Norte", "Oeste", "Sur"]


def test_clean_after_reading_real_excel(tmp_path):
    path = tmp_path / "s.xlsx"
    samples.sales("en").to_excel(path, index=False)
    df, report = L.clean(read_table(path))
    assert report["unified"] > 0 and report["trimmed"] > 0
    assert df["Region"].nunique() == 5


def test_run_plan_group_filter_sort(sales):
    plan = {"filters": [{"column": "categoría", "op": "==", "value": "software", "value2": ""}],
            "group_by": [{"column": "Fecha", "date_part": "quarter"}],
            "metrics": [{"column": "Total", "agg": "sum", "alias": "Ventas"}],
            "select_columns": [], "sort_by": "Ventas", "ascending": False, "limit": 2}
    result = L.run_plan(sales, plan)
    assert list(result.columns) == ["Fecha (quarter)", "Ventas"] and len(result) == 2
    assert result["Ventas"].is_monotonic_decreasing


def test_run_plan_totals_and_unknown_column(sales):
    plan = {"filters": [], "group_by": [], "metrics": [{"column": "Total", "agg": "sum", "alias": "T"}],
            "select_columns": [], "sort_by": "", "ascending": False, "limit": 0}
    assert L.run_plan(sales, plan)["T"][0] == pytest.approx(sales["Total"].sum())
    with pytest.raises(KeyError):
        L.run_plan(sales, {**plan, "metrics": [{"column": "Nope", "agg": "sum", "alias": "x"}]})


def test_filters_in_between_contains(sales):
    base = {"group_by": [], "metrics": [], "select_columns": [], "sort_by": "", "ascending": True, "limit": 0}
    r = L.run_plan(sales, {**base, "filters": [{"column": "Región", "op": "in", "value": "norte|SUR", "value2": ""}]})
    assert set(r["Región"]) == {"Norte", "Sur"}
    r = L.run_plan(sales, {**base, "filters": [{"column": "Total", "op": "between", "value": "100", "value2": "200"}]})
    assert r["Total"].between(100, 200).all()
    r = L.run_plan(sales, {**base, "filters": [{"column": "Producto", "op": "contains", "value": "erp", "value2": ""}]})
    assert (r["Producto"] == "Licencia ERP").all()


@pytest.mark.parametrize("formula", ['=SUMIFS(D2:D100,B2:B100,"Norte, Sur",C2:C100,">=1.5")',
                                     "=IFERROR(VLOOKUP(A2,'Hoja.1'!$A:$C,3,FALSE),0.25)",
                                     "=YEAR(TODAY())-1", "=SUM({1,2,3})*2.5"])
def test_formula_translation_roundtrip(formula):
    es = L.translate_formula(formula, "en", "es")
    fr = L.translate_formula(formula, "en", "fr")
    assert L.translate_formula(es, "es", "en") == formula
    assert L.translate_formula(fr, "fr", "es") == es


def test_formula_translation_examples():
    assert L.translate_formula("=IFERROR(VLOOKUP(A2,B:C,2,0),0.5)", "en", "es") == \
        "=SI.ERROR(BUSCARV(A2;B:C;2;0);0,5)"
    assert L.translate_formula("=SOMME.SI.ENS(C:C;A:A;\"x\")", "fr", "en") == '=SUMIFS(C:C,A:A,"x")'


def test_export_report(tmp_path, sales):
    labels = {"data": "Datos", "summary": "Resumen", "profile": "Perfil", "insights": "IA", "disclaimer": "d"}
    path = L.export_report(tmp_path / "r.xlsx", sales, labels, L.pivot(sales, ["Región"], "Total"), "Hallazgos\n- uno")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Datos", "Resumen", "Perfil", "IA"]
    assert wb["Resumen"]._charts and wb["Datos"].auto_filter.ref


def test_merge_files(tmp_path):
    a, b = tmp_path / "a.csv", tmp_path / "b.xlsx"
    pd.DataFrame({"x": [1, 2], "y": ["a", "b"]}).to_csv(a, sep=";", index=False)
    pd.DataFrame({"x": [3], "z": [True]}).to_excel(b, index=False)
    merged = L.merge_files([a, b], "Origen")
    assert len(merged) == 3 and set(merged.columns) == {"Origen", "x", "y", "z"}
