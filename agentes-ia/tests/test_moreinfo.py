from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image

from suite.moreinfo import logic


def shots(tmp_path, sections=("Transactions", "Transactions", "Profile", "Cards", "Claim information")):
    out = []
    for n, section in enumerate(sections):
        path = tmp_path / f"shot{n}.png"
        Image.new("RGB", (1600 if n == 0 else 700, 300 + 250 * n), (200, 220, 240)).save(path)
        out.append(logic.Shot(section, str(path), f"caption {n}" if n % 2 else ""))
    return out


def test_clean_case_and_numbering(tmp_path):
    assert logic.clean_case(' 1234/56 "x" ') == "123456x"
    a, b = tmp_path / "a", tmp_path / "b"
    a.mkdir()
    b.mkdir()
    assert logic.next_names("C1", "MOREINFO", [a, b], [".pdf", ".pdf"]) == ["C1,MOREINFO.pdf", "C1,MOREINFO1.pdf"]
    (a / "C1,MOREINFO.pdf").touch()
    (b / "c1,moreinfo2.xlsx").touch()
    (b / "C10,MOREINFO1.pdf").touch()  # otro caso
    assert logic.next_names("C1", "MOREINFO", [a, b, tmp_path / "missing"], [".pdf", ".docx"]) == [
        "C1,MOREINFO1.pdf", "C1,MOREINFO3.docx"]


def test_pdf_package_is_numbered_and_copied(tmp_path):
    share1, share2 = tmp_path / "share1", tmp_path / "share2"
    share1.mkdir()
    share2.mkdir()
    (share2 / "777,MOREINFO.pdf").write_bytes(b"%PDF")
    extra_img = tmp_path / "extra.png"
    Image.new("RGB", (300, 200), "red").save(extra_img)
    extra_doc = tmp_path / "statement.xlsx"
    Workbook().save(extra_doc)
    result = logic.create_package("777", "Ana", shots(tmp_path), [str(extra_img), (str(extra_doc), True)],
                                  share1, [share1, share2], use_word=False, keep_docx=True)
    assert result["engine"] == "builtin" and not result["errors"] and result["kept"] is None
    names = result["names"]
    assert names == ["777,MOREINFO1.pdf", "777,MOREINFO1.docx", "777,MOREINFO2.pdf", "777,MOREINFO3.xlsx"]
    assert sorted(p.name for p in share1.iterdir()) == sorted(names)
    # la carta (marcada) va también a la segunda carpeta de cartas; el resto solo a la principal
    assert sorted(p.name for p in share2.iterdir()) == ["777,MOREINFO.pdf", "777,MOREINFO3.xlsx"]
    pdf = (share1 / names[0]).read_bytes()
    assert pdf.startswith(b"%PDF") and pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page\r") >= 1
    doc = Document(share1 / names[1])
    assert len(doc.inline_shapes) == 5
    assert [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")] == [
        "Transactions", "Profile", "Cards", "Claim information"]


def test_existing_file_continues_numbering(tmp_path):
    share = tmp_path / "share"
    logic.create_package("888", "", shots(tmp_path)[:1], [], share, use_word=False)
    again = logic.create_package("888", "", shots(tmp_path)[:1], [], share, use_word=False)
    assert again["names"] == ["888,MOREINFO1.pdf"] and not again["errors"]
    assert sorted(p.name for p in share.iterdir()) == ["888,MOREINFO.pdf", "888,MOREINFO1.pdf"]


def test_case_documents_keep_or_convert_format(tmp_path):
    share = tmp_path / "share"
    letter = tmp_path / "No pay letter.docx"
    doc = Document()
    doc.add_paragraph("We have completed the review of your claim.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Amount", "$120.00"
    doc.save(letter)
    pdf = tmp_path / "footprint.pdf"
    logic.image_to_pdf(_png(tmp_path), pdf)
    sheet = tmp_path / "voice match.xlsx"
    Workbook().save(sheet)
    email = tmp_path / "approval.msg"
    email.write_bytes(b"msg")
    text = tmp_path / "web extraction.txt"
    text.write_text("extract", encoding="utf-8")
    old_doc = tmp_path / "old.doc"
    old_doc.write_bytes(b"doc")
    documents = [str(x) for x in (letter, pdf, sheet, email, text, old_doc)]
    result = logic.create_package("4321", "", shots(tmp_path)[:1], documents, share, use_word=False)
    assert result["names"] == ["4321,MOREINFO.pdf", "4321,MOREINFO1.pdf", "4321,MOREINFO2.pdf",
                               "4321,MOREINFO3.xlsx", "4321,MOREINFO4.msg", "4321,MOREINFO5.txt",
                               "4321,MOREINFO6.doc"]
    assert result["notes"] == [("doc_text_only", "No pay letter.docx"), ("doc_kept", "old.doc")]
    assert (share / "4321,MOREINFO1.pdf").read_bytes().startswith(b"%PDF")
    only = logic.create_package("4321", "", [], [str(text)], share, include_moreinfo=False)
    assert only["names"] == ["4321,MOREINFO7.txt"] and only["engine"] is None
    with pytest.raises(ValueError):
        logic.create_package("4321", "", [], [], share, include_moreinfo=False)


def _png(tmp_path, name="img.png"):
    path = tmp_path / name
    Image.new("RGB", (200, 100), "white").save(path)
    return path


def test_excel_package_adds_tab(tmp_path):
    wb = Workbook()
    wb.active.title = "Transactions"
    wb.active.append(["Date", "Merchant", "Amount"])
    wb.active.append(["2026-10-01", "Shop", 120.5])
    src = tmp_path / "txns.xlsx"
    wb.save(src)
    result = logic.create_package("999", "Ana", shots(tmp_path)[2:], [], tmp_path / "out", mode="excel",
                                  excel_source=str(src))
    book = load_workbook(tmp_path / "out" / "999,MOREINFO.xlsx")
    assert result["names"] == ["999,MOREINFO.xlsx"]
    assert book.sheetnames == ["Transactions", "Additional info"] and book["Transactions"]["C2"].value == 120.5
    assert len(book["Additional info"]._images) == 3 and book["Additional info"]["A1"].value == \
        "Additional info – Case 999"


def test_validation(tmp_path):
    cases = [(" ", shots(tmp_path), tmp_path, {}), ("1", shots(tmp_path), "", {}), ("1", [], tmp_path, {}),
             ("1", [], tmp_path, {"mode": "excel"})]
    for case, items, folders, kwargs in cases:
        with pytest.raises(ValueError):
            logic.create_package(case, "", items, [], folders, **kwargs)


def test_unreachable_folder_is_reported(tmp_path):
    blocker = tmp_path / "file.txt"
    blocker.write_text("x")
    result = logic.create_package("5", "", shots(tmp_path)[:1], [], blocker / "sub", use_word=False)
    assert result["errors"] and (Path(result["kept"]) / "5,MOREINFO.pdf").exists()  # nada se pierde


# ---------------------------------------------------------------- No Pay Decision Letter template

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _legacy_template(path):
    """Plantilla con un campo de formulario antiguo, un desplegable antiguo, un REF que repite el nombre y un
    marcador partido en dos fragmentos de texto."""
    from lxml import etree

    doc = Document()
    p = doc.add_paragraph()
    xml = (
        f'<w:p xmlns:w="{W_NS}"><w:r><w:t xml:space="preserve">Name: </w:t></w:r>'
        '<w:bookmarkStart w:id="1" w:name="CustName"/>'
        '<w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="CustName"/><w:enabled/>'
        '<w:textInput/></w:ffData></w:fldChar></w:r><w:r><w:instrText> FORMTEXT </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>     </w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r><w:bookmarkEnd w:id="1"/></w:p>')
    p._p.getparent().replace(p._p, etree.fromstring(xml))
    p = doc.add_paragraph()
    xml = (
        f'<w:p xmlns:w="{W_NS}"><w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="Outcome"/>'
        '<w:ddList><w:listEntry w:val="Denied"/><w:listEntry w:val="Partially approved"/></w:ddList>'
        '</w:ffData></w:fldChar></w:r><w:r><w:instrText> FORMDROPDOWN </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')
    p._p.getparent().replace(p._p, etree.fromstring(xml))
    p = doc.add_paragraph()
    xml = (
        f'<w:p xmlns:w="{W_NS}"><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText> REF CustName \\h </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>old</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')
    p._p.getparent().replace(p._p, etree.fromstring(xml))
    p = doc.add_paragraph("Case: ")
    p.add_run("[Case")
    p.add_run(" number] end")
    doc.save(path)
    return path


def _all_text(path):
    import zipfile

    from lxml import etree

    tree = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    return ["".join(t.text or "" for t in p.iter(f"{{{W_NS}}}t")) for p in tree.iter(f"{{{W_NS}}}p")]


def test_letter_content_controls(tmp_path):
    from suite.moreinfo import letter

    template = letter.make_sample_template(tmp_path / "blank.docx")
    fields = {f.key: f for f in letter.scan(template)}
    assert fields["cc:customername"].count == 3 and fields["cc:reason"].kind == "dropdown"
    assert fields["cc:letterdate"].kind == "date" and fields["cc:channel"].kind == "combo"
    assert len(fields["cc:reason"].options) == 4 and "ph:claim number" in fields
    values = {"cc:customername": "Laura Méndez", "cc:reason": fields["cc:reason"].options[1],
              "cc:letterdate": "October 30, 2026", "ph:claim number": "R-2026-0042", "cc:channel": "Online"}
    out = letter.fill(template, values, tmp_path / "letter.docx")
    text = _all_text(out)
    assert sum(t.count("Laura Méndez") for t in text) == 3  # también como título al final
    assert "Claim number: R-2026-0042" in text and any("Card was in the customer" in t for t in text)
    assert text[-1] == "Laura Méndez"
    again = {f.key: f for f in letter.scan(out)}
    assert again["cc:customername"].value == "Laura Méndez"
    # solo quedan con el texto gris de ejemplo los 3 campos que no se rellenaron (importe, comercio, analista)
    assert __import__("zipfile").ZipFile(out).read("word/document.xml").decode().count("PlaceholderText") == 3


def test_letter_legacy_fields_ref_and_split_placeholder(tmp_path):
    from suite.moreinfo import letter

    template = _legacy_template(tmp_path / "legacy.docx")
    fields = {f.key: f for f in letter.scan(template)}
    assert set(fields) == {"ff:custname", "ff:outcome", "ph:case number"}
    assert fields["ff:outcome"].options == ["Denied", "Partially approved"]
    out = letter.fill(template, {"ff:custname": "Emma Wilson", "ff:outcome": "Partially approved",
                                 "ph:case number": "R-77"}, tmp_path / "out.docx")
    text = _all_text(out)
    assert text[0].startswith("Name: ") and "Emma Wilson" in text[0]
    assert text[2] == "Emma Wilson"  # el REF repite el nombre
    assert text[3] == "Case: R-77 end"
    xml = __import__("zipfile").ZipFile(out).read("word/document.xml").decode()
    assert 'w:result w:val="1"' in xml


def test_letter_bound_to_custom_xml(tmp_path):
    import zipfile

    from suite.moreinfo import letter

    template = tmp_path / "bound.docx"
    letter.make_sample_template(template)
    data = {n: zipfile.ZipFile(template).read(n) for n in zipfile.ZipFile(template).namelist()}
    xml = data["word/document.xml"].decode().replace(
        '<w:tag w:val="Merchant"/>', '<w:tag w:val="Merchant"/><w:dataBinding w:prefixMappings="xmlns:ns0=\'urn:x\'" '
        'w:xpath="/ns0:letter[1]/ns0:merchant[1]" w:storeItemID="{11111111-1111-1111-1111-111111111111}"/>')
    data["word/document.xml"] = xml.encode()
    data["customXml/item1.xml"] = b'<letter xmlns="urn:x"><merchant>old</merchant></letter>'
    with zipfile.ZipFile(template, "w") as zf:
        for name, content in data.items():
            zf.writestr(name, content)
    out = letter.fill(template, {"cc:merchant": "TechStore"}, tmp_path / "out.docx")
    assert b"TechStore" in zipfile.ZipFile(out).read("customXml/item1.xml")
