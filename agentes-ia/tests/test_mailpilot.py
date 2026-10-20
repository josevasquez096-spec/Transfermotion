import random
import zipfile
from datetime import datetime

import pytest
from openpyxl import load_workbook

import fake_outlook
from suite.mailpilot import analysis, assist, detectors, english, rules
from suite.mailpilot.backend_base import BackendError, Progress
from suite.mailpilot.backend_demo import DRAFTS, DemoBackend
from suite.mailpilot.backend_outlook import OutlookBackend
from suite.mailpilot.edits import apply_changes, diff_changes, split_quoted, word_text
from suite.mailpilot.model import Message
from suite.mailpilot.prefs import Prefs
from suite.mailpilot.query import Query, QueryError, fold, outlook_conditions

NOW = datetime(2026, 10, 20, 12, 0)
INTERNAL = ["banco-demo.com"]


def mail(key, **kw):
    base = dict(id=key, store_id="s", folder_id="f", folder_path="Bandeja de entrada/Reclamaciones")
    base.update(kw)
    return Message(**base)


MESSAGES = {
    "claim": mail("claim", subject="Reclamación R-2026-0042 cargo no reconocido", sender_name="Laura Méndez",
                  sender_email="laura.mendez@gmail.com", to="Fraude", received=datetime(2026, 10, 15, 9),
                  unread=True, categories=["Urgente", "Contracargo"], has_attachments=True, size=3_000_000,
                  body="Tengo un cargo  no\nreconocido en mi tarjeta 4111 1111 1111 1111. "
                       "IBAN ES91 2100 0418 4502 0005 1332", attachments=["captura.png"]),
    "visa": mail("visa", subject="Visa Alert: chargeback 7781", sender_name="Visa Alerts",
                 sender_email="alerts@visa.com", received=datetime(2026, 8, 1), body="Case CLM-778100.", importance=2),
    "phish": mail("phish", subject="URGENT: verify your account", sender_name="IT Service Desk",
                  sender_email="it@banco-dem0.com", received=datetime(2026, 10, 19),
                  body="Verify your account within 24 hours", reply_to="x@evil.top",
                  html='<a href="http://192.168.4.4/login">https://www.banco-demo.com</a>'),
    "internal": mail("internal", subject="Reunión de equipo", sender_name="Ana Torres",
                     sender_email="ana.torres@banco-demo.com", received=datetime(2026, 10, 20, 8),
                     body="Nos vemos a las 10:30 en la sala 3.", flagged=True),
}


def matching(text, **options):
    query = Query(text, now=NOW, **options)
    return {k for k, m in MESSAGES.items() if query.matches(m, internal_domains=INTERNAL)}


@pytest.mark.parametrize("text,expected", [
    ("reconocido", {"claim"}),
    ('"cargo no reconocido"', {"claim"}),
    ("RECLAMACION", {"claim"}),
    ("from:visa.com", {"visa"}),
    ("de:laura", {"claim"}),
    ("asunto:alert", {"visa"}),
    ("-is:unread", {"visa", "phish", "internal"}),
    ("chargeback OR contracargo", {"visa", "claim"}),
    ("(chargeback OR contracargo) is:unread", {"claim"}),
    ("-(visa OR banco)", {"claim"}),
    ("has:card", {"claim"}),
    ("tiene:iban", {"claim"}),
    ("has:claim", {"claim", "visa"}),
    ("has:attachment size:>1mb", {"claim"}),
    ("is:external", {"claim", "visa", "phish"}),
    ("is:internal", {"internal"}),
    ("is:phishing", {"phish"}),
    ("after:2026-10-01", {"claim", "phish", "internal"}),
    ("desde:01/10/2026 hasta:19/10/2026", {"claim"}),
    ("newer:2d", {"phish", "internal"}),
    ("older:30d", {"visa"}),
    ("on:2026-10-20", {"internal"}),
    ("fecha:hoy", {"internal"}),
    ("etiqueta:contra", {"claim"}),
    ("-has:category", {"visa", "phish", "internal"}),
    ("importancia:alta", {"visa"}),
    (r"re:R-\d{4}-\d+", {"claim"}),
    ("subject:/^visa/", {"visa"}),
    ("domain:gmail.com", {"claim"}),
    ("filename:png", {"claim"}),
    ("10:30", {"internal"}),
    ("subject:(verify OR reunión)", {"phish", "internal"}),
    ("is:flagged", {"internal"}),
    ("", set(MESSAGES)),
])
def test_query_matches(text, expected):
    assert matching(text) == expected


def test_query_options_and_folding():
    assert matching("carg", whole_word=True) == set()
    assert matching("cargo", whole_word=True) == {"claim"}
    assert matching("Reclamación", exact=True) == {"claim"}
    assert matching("reclamación", exact=True) == set()
    assert fold("Reclamación  Nº 5 – Peña’s") == "reclamacion nº 5 - pena's"
    assert Query("después:2026-09-01 fraude", now=NOW).date_floor() == datetime(2026, 9, 1)
    patterns = [p.pattern for p in Query('reclamacion "cargo no" -visa', now=NOW).highlights()]
    assert len(patterns) == 2 and "[oóòôöõ]" in patterns[0]


@pytest.mark.parametrize("text,code", [("has:foo", "has"), ("is:bar", "is"), ("after:31/02/2026", "date"),
                                       ("size:big", "size"), ("(a OR b", "paren"), ("a)", "paren"),
                                       ("re:([", "regex"), ("after:99999d", "date"),
                                       ("importance:max", "importance")])
def test_query_errors(text, code):
    with pytest.raises(QueryError) as info:
        Query(text, now=NOW)
    assert info.value.code == code


def test_dasl_prefilter_is_never_stricter_than_python():
    # Sin «exacto», un texto con letras puede llevar acentos: no se manda a Outlook.
    assert Query("reclamacion", now=NOW).to_dasl() == ""
    # Los fragmentos sin letras sí (números de reclamación, importes…).
    dasl = Query("R-2026-0042", now=NOW).to_dasl()
    assert "'%2026%'" in dasl and "'%0042%'" in dasl and "R-" not in dasl
    exact = Query("Reclamación", exact=True, now=NOW).to_dasl()
    assert exact.startswith("@SQL=") and "%Reclamación%" in exact and "textdescription" in exact
    assert "read\" = 0" in Query("is:unread fraude", now=NOW).to_dasl()
    assert Query("-is:unread", now=NOW).to_dasl() == ""
    assert Query("o'brien", exact=True, now=NOW).to_dasl().count("o''brien") >= 1


def test_body_is_loaded_only_when_needed():
    bodies = {"a": "La tarjeta 4111111111111111", "b": "Nada", "c": "fraude en curso"}
    loads = []

    def loader(msg, html=False):
        loads.append(msg.id)
        msg.body, msg.attachments = bodies[msg.id], []
        if html:
            msg.html = ""

    def fresh():
        return [Message(k, "s", "f", "Inbox", subject=f"Asunto {k}",
                        sender_email=f"x@{'visa.com' if k == 'a' else 'banco-demo.com'}") for k in "abc"]

    query = Query("from:visa.com has:card")
    assert [m.id for m in fresh() if query.matches(m, loader)] == ["a"] and loads == ["a"]
    loads.clear()
    assert not [m for m in fresh() if Query("from:nadie fraude").matches(m, loader)] and loads == []


def test_outlook_rule_conversion():
    spec = outlook_conditions(Query("subject:(fraude OR fraud) from:visa.com -body:prueba"))
    assert spec == {"conditions": {"subject": ["fraude", "fraud"], "sender_address": ["visa.com"]},
                    "exceptions": {"body": ["prueba"]}}
    assert outlook_conditions(Query("has:attachment is:important domain:visa.com"))["conditions"] == {
        "has_attachment": True, "importance": 2, "sender_address": ["visa.com"]}
    for text in ("fraude tarjeta", "after:2026-01-01 fraude", 'from:"Juan Perez"', "-fraude", "re:abc"):
        assert outlook_conditions(Query(text)) is None, text


# ------------------------------------------------------------------ corrector de inglés

GOOD = """I have attached the statement for your review.
Could you please send me the signed form by Friday?
Did you receive my previous email?
It's important that we close the case this week.
Its status changed to resolved yesterday.
Is there money left in the account?
Is important information missing from the file?
You're right about the deadline.
I look forward to your reply.
He has worked here since 2019.
We have write access to the shared folder.
We have know-how in fraud prevention.
The branch is close to my house.
I have to send the report today.
The card was blocked yesterday.
I don't know why nobody answered the phone.
Actually, I agree with you.
See https://intranet.bank.com/policies or write to fraud.team@bank.com.
The total was 1,250.00 USD (approx.)."""

BAD = [("I recieved the documents yesterday.", "recieved"), ("I didn't received the email.", "received"),
       ("Please explain me the process.", "explain me"), ("I look forward to hear from you.", "hear"),
       ("We need more informations.", "informations"), ("I am waiting since 3 days.", "since 3 days"),
       ("We have send the documents.", "send"), ("The documents were send yesterday.", "send"),
       ("Is very important for our client.", "Is very important"), ("We have received you claim.", "you"),
       ("Actually the case is in revision.", "Actually"), ("Please dont share your PIN with nobody.", "nobody"),
       ("Its a fraud case.", "Its"), ("Their are two cards.", "Their"), ("This is a important case.", "a"),
       ("We speak spanish and english.", "spanish"), ("Please do the needful.", "do the needful")]


@pytest.mark.parametrize("line", GOOD.splitlines())
def test_english_no_false_positives(line):
    issues = [i for i in english.check(line) if i.kind != "clarity"]
    assert not issues, [(i.rule, i.text) for i in issues]


@pytest.mark.parametrize("sentence,expected", BAD)
def test_english_detects_mistakes(sentence, expected):
    assert any(expected.lower() in i.text.lower() or i.text.lower() in expected.lower()
               for i in english.check(sentence) if i.text), sentence


def test_english_missing_subject_statement_and_question():
    fix = lambda s: english.apply(s, [i for i in english.check(s) if i.rule == "missing_subject"])[0]  # noqa: E731
    assert fix("Is possible to change the address?") == "Is it possible to change the address?"
    assert fix("Is possible to change the address.") == "It is possible to change the address."
    assert fix("We need it because is necessary.") == "We need it because it is necessary."


def test_english_apply_safe_and_subject():
    text, _ = english.apply_safe("i am writting to you. Thank you for contact us.", "RE: RE: claim")
    assert text.startswith("I am writing") and "contacting" in text
    issues = english.check("Hello team", subject="URGENT CLAIM UPDATE NOW")
    assert any(i.rule == "subject_caps" for i in issues)
    assert all(i.rule in english.MESSAGES for i in english.check(BAD[0][0] + " " + GOOD))


# ------------------------------------------------------------------ cambios en borradores

def test_diff_changes_roundtrip_random():
    rnd = random.Random(11)
    words = ["the", "a", "claim", "card", "i", "is", "fraud", ",", ".", "\n", "  ", "the the", "data", "\n\n"]
    for _ in range(3000):
        old = " ".join(rnd.choice(words) for _ in range(rnd.randint(1, 20)))
        tokens = old.split(" ")
        for _ in range(rnd.randint(1, 4)):
            k = rnd.randrange(len(tokens))
            if rnd.random() < 0.5:
                tokens[k] = rnd.choice(words)
            else:
                tokens.insert(k, rnd.choice(words))
        new = " ".join(tokens)
        changes = diff_changes(old, new)
        assert all(find for find, _, _ in changes)
        assert apply_changes(old, changes)[0] == new, (old, new)


def test_diff_changes_on_checker_fixes():
    for _, _, body in DRAFTS:
        fixed, _ = english.apply(body, english.check(body))
        changes = diff_changes(body, fixed)
        assert changes and apply_changes(body, changes) == (fixed, len(changes), 0)


def test_split_quoted_and_word_text():
    own, quoted = split_quoted("Hi Ana,\nThanks.\n\nFrom: Juan <j@x.com>\nSent: Monday\nTo: Ana\n\nold")
    assert own == "Hi Ana,\nThanks.\n\n" and quoted.startswith("From: Juan")
    assert split_quoted("Hola\n\nEl lun, 5 oct 2026 a las 10:00, Juan escribió:\n> hola")[0] == "Hola\n\n"
    assert split_quoted("Bonjour\r\n________________________________\r\nDe : Juan")[0] == "Bonjour\r\n"
    assert word_text("a\rb\x0bc\x07d\x13e") == "a\nb\nc\nd e"


# ------------------------------------------------------------------ detectores

def test_detectors():
    assert detectors.find_cards("4111 1111 1111 1111 and 4111 1111 1111 1112") == ["4111 1111 1111 1111"]
    assert detectors.find_ibans("IBAN ES91 2100 0418 4502 0005 1332.") == ["ES91 2100 0418 4502 0005 1332"]
    assert detectors.find_claims("Caso R-2026-0042 y CLM-778100", r"(R|CLM)-\d+(-\d+)?") == ["R-2026-0042",
                                                                                        "CLM-778100"]
    assert detectors.find_amounts("1.250,00 €. $1,250.00 y 300 EUR") == ["1.250,00 €", "$1,250.00", "300 EUR"]
    assert detectors.mask_iban("ES91 2100 0418 4502 0005 1332") == "ES91 **** **** 1332"
    for domain in ("banco-dem0.com", "bancodemo-secure.com", "banco-demo.com.verify.top"):
        assert detectors.lookalike(domain, {"banco-demo.com"}), domain
    for domain in ("visa.com", "mail.banco-demo.com", "gmail.com"):
        assert not detectors.lookalike(domain, {"banco-demo.com"}), domain


@pytest.mark.parametrize("key,phishing", [("phish", True), ("visa", False), ("claim", False), ("internal", False)])
def test_phishing_score(key, phishing):
    msg = MESSAGES[key]
    score, flags = detectors.phishing_score(msg, msg.body, msg.html or "", INTERNAL)
    assert (score >= detectors.PHISHING_THRESHOLD) == phishing, (score, flags)


# ------------------------------------------------------------------ buzón de demostración

@pytest.fixture
def demo():
    backend = DemoBackend("es", now=NOW)
    backend.connect()
    return backend


def all_folders(backend):
    return [f.id for s in backend.stores() for f in backend.folders(s.id) if f.depth > 0]


@pytest.mark.parametrize("lang", ["es", "en", "fr"])
def test_demo_mailbox_contents(lang):
    backend = DemoBackend(lang, now=NOW)
    stores = backend.connect()
    assert [s.kind for s in stores] == ["primary", "shared", "pst"]
    found, _ = backend.search(all_folders(backend), Query("is:phishing"), internal_domains=INTERNAL)
    assert len(found) == 4  # suplantación de TI y del banco, fraude del CEO y factura falsa
    assert backend.search(all_folders(backend), Query("has:card"))[0]
    assert all(m.received < NOW for m in backend._messages.values())


def test_demo_folder_operations(demo):
    folders = demo.folders("demo-primary")
    inbox = next(f for f in folders if f.special == "inbox")
    new = demo.create_folder(inbox.id, "Casos VIP")
    with pytest.raises(BackendError) as info:
        demo.create_folder(inbox.id, "casos vip")
    assert info.value.code == "folder_exists"
    with pytest.raises(BackendError):
        demo.rename_folder(inbox.id, "Otra")  # carpeta del sistema
    with pytest.raises(BackendError) as info:
        demo.move_folder(inbox.id, new.id)
    assert info.value.code == "folder_special"
    sub = demo.create_folder(new.id, "2026")
    with pytest.raises(BackendError) as info:
        demo.move_folder(new.id, sub.id)
    assert info.value.code == "folder_loop"
    demo.rename_folder(new.id, "VIP")
    assert any(f.path == "Bandeja de entrada/VIP/2026" for f in demo.folders("demo-primary"))
    demo.delete_folder(new.id)
    assert any(f.path == "Elementos eliminados/VIP" for f in demo.folders("demo-primary"))


def test_demo_message_operations(demo, tmp_path):
    folders = demo.folders("demo-primary")
    inbox = next(f for f in folders if f.special == "inbox")
    archive = next(f for f in folders if f.path == "Archivo")
    messages, _ = demo.messages(inbox.id, limit=5)
    assert all(m.body is None for m in messages)  # el cuerpo se carga aparte
    assert demo.details(messages[0]).body
    assert demo.set_categories(messages[:2], add=["Revisar"], remove=["Urgente"]) == 2
    assert demo.mark_read(messages[:2]) == 2 and demo.set_flag(messages[:1]) == 1
    assert demo.move_messages(messages[:1], archive.id) == 1
    assert demo.delete_messages(messages[1:2]) == 1
    saved = demo.save_messages(messages[2:4], tmp_path)
    assert len(saved) == 2 and saved[0].suffix == ".eml"
    store = demo.open_shared("pagos@banco-demo.com")
    assert store.kind == "shared" and demo.folders(store.id)[1].total == 8
    with pytest.raises(BackendError):
        demo.open_shared("nadie@banco-demo.com")


def test_demo_categories_and_drafts(demo):
    demo.create_category("VIP", 4)
    with pytest.raises(BackendError):
        demo.create_category("vip", 1)
    demo.update_category("VIP", new_name="Clientes VIP", color=5)
    assert ("Clientes VIP", 5) in [(c.name, c.color) for c in demo.categories()]
    demo.delete_category("Clientes VIP")
    draft = demo.current_draft()
    fixed, _ = english.apply(draft["body"], english.check(draft["body"]))
    applied, failed, current = demo.apply_to_current(diff_changes(draft["body"], fixed), "New subject")
    assert applied and not failed and current == fixed
    created = demo.create_draft("x@y.com", "Hola", "Texto")
    assert created.subject == "Hola"


def test_search_limit_and_cancel(demo):
    found, truncated = demo.search(all_folders(demo), Query(""), limit=10)
    assert len(found) == 10 and truncated
    progress = Progress()
    progress.cancel.set()
    found, stopped = demo.search(all_folders(demo), Query(""), progress=progress)
    assert stopped and not found


# ------------------------------------------------------------------ Outlook (modelo simulado)

@pytest.fixture
def outlook():
    app = fake_outlook.build()
    backend = OutlookBackend(dispatch=lambda: app, com_init=False, separator=";")
    backend.connect()
    yield backend, app
    backend.close()


def test_outlook_stores_folders_and_messages(outlook):
    backend, _ = outlook
    stores = backend.stores()
    assert [(s.name, s.kind) for s in stores] == [("ana.torres@banco-demo.com", "primary"), ("Archivo 2025", "pst")]
    shared = backend.open_shared("fraude@banco-demo.com")
    assert shared.kind == "shared" and len(backend.stores()) == 3
    with pytest.raises(BackendError) as info:
        backend.open_shared("nadie@x.com")
    assert info.value.code == "shared_failed"
    folders = backend.folders(stores[0].id)
    assert [f.special for f in folders if f.special] == ["inbox", "sent", "drafts", "deleted", "junk"]
    assert "Calendar" not in [f.name for f in folders]
    inbox = next(f for f in folders if f.special == "inbox")
    messages, truncated = backend.messages(inbox.id)
    assert [m.subject for m in messages][0] == "URGENT: verify your account" and not truncated
    visa = next(m for m in messages if m.sender_email == "disputes@visa.com")
    assert visa.categories == ["Urgente", "Seguimiento"] and visa.unread and visa.has_attachments
    exchange = next(m for m in messages if m.sender_name == "Marta Gil")
    assert exchange.sender_email == "" and exchange.flagged  # dirección X500 = interna


def test_outlook_search_uses_and_survives_prefilter(outlook):
    backend, app = outlook
    backend.open_shared("fraude@banco-demo.com")
    ids = [f.id for s in backend.stores() for f in backend.folders(s.id)]
    found, _ = backend.search(ids, Query("is:phishing"), internal_domains=INTERNAL)
    assert [m.subject for m in found] == ["URGENT: verify your account"]
    found, _ = backend.search(ids, Query("has:card"))
    assert [m.subject for m in found] == ["Reclamación R-2026-0042"]
    found, _ = backend.search(ids, Query("is:unread"))
    assert app.ns.DefaultStore.default_folders[6].tables[-1].startswith("@SQL=")
    app.ns.DefaultStore.reject_filters = True  # filtro no admitido: se evalúa todo en Python
    assert [m.subject for m in backend.search(ids, Query("is:unread"))[0]] == [m.subject for m in found]


def test_outlook_message_operations(outlook, tmp_path):
    backend, app = outlook
    folders = backend.folders(backend.stores()[0].id)
    inbox = next(f for f in folders if f.special == "inbox")
    claims = next(f for f in folders if f.name == "Reclamaciones")
    visa = next(m for m in backend.messages(inbox.id)[0] if m.sender_email == "disputes@visa.com")
    assert backend.set_categories([visa], add=["Contracargo"], remove=["Urgente"]) == 1
    assert app.ns.GetItemFromID(visa.id).Categories == "Seguimiento; Contracargo"
    assert backend.move_messages([visa], claims.id) == 1
    moved = backend.messages(claims.id)[0]
    assert backend.mark_read(moved) == 2 and backend.set_flag(moved) == 2
    assert [p.suffix for p in backend.save_messages(moved, tmp_path)] == [".msg", ".msg"]
    assert [p.name for p in backend.save_attachments(moved, tmp_path)] == ["notice.pdf"]
    assert backend.delete_messages(moved[:1]) == 1
    with pytest.raises(BackendError) as info:
        backend.move_messages([visa], claims.id)  # el identificador cambió al mover
    assert info.value.code == "com"


def test_outlook_folders_categories_and_rules(outlook):
    backend, app = outlook
    store = backend.stores()[0]
    inbox = next(f for f in backend.folders(store.id) if f.special == "inbox")
    new = backend.create_folder(inbox.id, "Contracargos")
    with pytest.raises(BackendError):
        backend.create_folder(inbox.id, "contracargos")
    backend.rename_folder(new.id, "Chargebacks")
    claims = next(f for f in backend.folders(store.id) if f.name == "Reclamaciones")
    backend.move_folder(new.id, claims.id)
    assert any(f.path == "Inbox/Reclamaciones/Chargebacks" for f in backend.folders(store.id))
    backend.delete_folder(new.id)
    assert any(f.path == "Deleted Items/Chargebacks" for f in backend.folders(store.id))
    backend.create_category("Phishing", 9)
    backend.update_category("Phishing", new_name="Sospechoso", color=1)
    backend.delete_category("Seguimiento")
    assert [(c.name, c.color) for c in backend.categories()] == [("Urgente", 1), ("Sospechoso", 1)]
    spec = outlook_conditions(Query("subject:(fraude OR fraud) from:visa.com -body:prueba"))
    backend.create_rule("MailPilot - visa", spec, store.id, claims.id, ["Contracargo"], flag=True)
    rule = app.ns.DefaultStore.rules.items[0]
    assert rule.Conditions.Subject.Text == ["fraude", "fraud"] and rule.Conditions.Subject.Enabled
    assert rule.Conditions.SenderAddress.Address == ["visa.com"] and rule.Exceptions.Body.Text == ["prueba"]
    assert rule.Actions.MoveToFolder.Folder.Name == "Reclamaciones" and rule.Actions.MarkAsTask.Enabled
    assert rule.Enabled and app.ns.DefaultStore.rules.saved == 1


def test_outlook_draft_is_corrected_in_word(outlook):
    backend, app = outlook
    item = app.ns.DefaultStore.default_folders[16].items[0]
    app.explorer.Selection.items = [item]
    draft = backend.current_draft()
    assert draft["editable"] and "From: Visa" not in draft["body"]
    fixed, _ = english.apply(draft["body"], english.check(draft["body"]))
    applied, failed, current = backend.apply_to_current(diff_changes(draft["body"], fixed), "Chargeback docs (v2)")
    assert applied == 3 and not failed and current == fixed
    assert item.document.text.startswith("Dear team,\rI am writing to inform you") and "From: Visa" in \
        item.document.text
    assert item.Subject == "Chargeback docs (v2)" and item.saved
    backend.create_draft("x@y.com", "Hola", "Line 1\nLine <2>")
    assert "Line &lt;2&gt;</p><p>--<br>Firma" in app.created[-1].HTMLBody


def test_outlook_errors_are_translated():
    def missing():
        raise fake_outlook.ComError(-2147221005, "Invalid class string", None, None)

    backend = OutlookBackend(dispatch=missing, com_init=False)
    with pytest.raises(BackendError) as info:
        backend.connect()
    assert info.value.code == "no_outlook"
    backend.close()
    app = fake_outlook.build()
    backend = OutlookBackend(dispatch=lambda: app, com_init=False)
    backend.connect()
    with pytest.raises(BackendError) as info:
        backend.current_draft()
    assert info.value.code == "no_current"
    backend.close()


# ------------------------------------------------------------------ reglas, análisis y preferencias

def test_rules_plan_execute_and_outlook(demo):
    stores = demo.stores()
    folders = {s.id: demo.folders(s.id) for s in stores}
    news = rules.MailRule("Boletines", "from:newsletter@fraudinsights.com OR from:noreply@cyberbulletin.org",
                          folder_store=stores[0].name, folder_path="Boletines", mark_read=True)
    visa = rules.MailRule("Visa", "from:visa.com", categories=["Revisar"], flag=True)

    def resolve(rule):
        return rules.find_folder(stores, folders, rule.folder_store, rule.folder_path)

    scope = [f.id for f in folders["demo-primary"] if f.depth > 0]
    candidates, _ = demo.search(scope, Query(""))
    planned = rules.plan([news, visa], candidates, resolve, demo.details, INTERNAL)
    assert {tuple(p.rules) for p in planned} == {("Boletines",), ("Visa",)}
    counts = rules.execute(demo, planned)
    assert counts["flagged"] == 9 and counts["categorized"] == 9 and counts["moved"] >= 1
    assert rules.plan([news, visa], demo.search(scope, Query(""))[0], resolve, demo.details, INTERNAL) == []
    assert rules.to_outlook(visa)[0]["conditions"] == {"sender_address": ["visa.com"]}
    assert rules.to_outlook(news)[1] == ["rule_note_read"]
    assert rules.to_outlook(rules.MailRule("x", "fraude tarjeta", categories=["a"]))[0] is None
    with pytest.raises(QueryError) as info:
        rules.compile_rules([rules.MailRule("Mala", "has:nada")])
    assert "Mala" in info.value.detail
    assert rules.MailRule.from_dict({**visa.to_dict(), "extra": 1}) == visa


def test_analysis_tables_and_excel(demo, tmp_path):
    messages, _ = demo.search(all_folders(demo), Query(""), limit=0)
    overview = analysis.overview(messages)
    assert overview["total"] == len(messages) > 100
    assert list(analysis.senders(messages, 5).columns)[0] == "col_sender"
    domains = analysis.domains(messages, INTERNAL)
    assert domains.loc[domains["col_domain"] == "banco-demo.com", "col_internal"].all()
    keys = {s["key"] for s in analysis.suggestions(messages, INTERNAL, now=NOW)}
    assert {"sug_newsletter", "sug_domain", "sug_claims"} <= keys
    full = demo.scan(messages, html=True)
    assert len(analysis.phishing(full, INTERNAL)) == 4
    sensitive = analysis.sensitive(full)
    assert {"card", "iban"} <= {kind for kind, _, _ in sensitive}
    assert all("*" in value for _, value, _ in sensitive)
    claims = analysis.claims(full, internal=INTERNAL)
    assert claims["col_claim"].str.match(r"R-20\d\d-\d+").all()
    evil = analysis.message_table([Message("x", "s", "f", "Inbox", subject='=HYPERLINK("http://evil.top","x")')])
    path = analysis.export_excel(tmp_path / "a.xlsx", [("Claims", claims), ("Messages", evil)])
    assert load_workbook(path)["Messages"]["D2"].data_type == "s"
    assert "<f>" not in zipfile.ZipFile(path).read("xl/worksheets/sheet2.xml").decode()


def test_prefs_roundtrip(tmp_path):
    prefs = Prefs.load(tmp_path / "mp.json")
    assert prefs.domains(["banco-demo.com"]) == ["banco-demo.com"]
    prefs["internal_domains"] = "@banco.com; filial.com"
    prefs["user_words"].append("chargeback")
    prefs.save()
    again = Prefs.load(tmp_path / "mp.json")
    assert again.domains(["x.com"]) == ["banco.com", "filial.com"] and again["user_words"] == ["chargeback"]
    (tmp_path / "bad.json").write_text("[1, 2]", encoding="utf-8")
    assert Prefs.load(tmp_path / "bad.json")["rules"] == []


def test_assist_helpers():
    from suite.privacy import Anonymizer

    anon = Anonymizer()
    masked = anon.mask("Write to emma@gmail.com about IBAN ES91 2100 0418 4502 0005 1332")
    assert "emma@gmail.com" not in masked
    data = assist.unmask_all({"body": masked, "changes": [masked], "n": 1}, anon)
    assert "emma@gmail.com" in data["body"] and "emma@gmail.com" in data["changes"][0]
    assert "Tone: firm but polite" in assist.email_prompt("Hi", "Body", "firm")
