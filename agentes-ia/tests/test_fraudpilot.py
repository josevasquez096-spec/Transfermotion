from datetime import date

from openpyxl import load_workbook

from suite import samples
from suite.fraudpilot import logic as L
from suite.privacy import Anonymizer


def claim(**kw):
    c = L.new_claim()
    ind = kw.pop("indicators", {})
    c.update(kw)
    c["indicators"].update(ind)
    return c


def test_classification_rules():
    cases = [
        ({"goods_not_received": True}, {}, "merchant_dispute"),
        ({"not_opened_by_customer": True}, {}, "identity_theft"),
        ({"authorized_payment": True, "contacted_by_fraudster": True}, {}, "app_scam"),
        ({"sim_issue": True}, {}, "sim_swap"),
        ({"shared_credentials": True}, {}, "phishing_vishing"),
        ({"new_device": True}, {"product": "online_banking"}, "account_takeover"),
        ({"card_lost_stolen": True}, {}, "lost_stolen_card"),
        ({"card_in_possession": True}, {"channel": "card_present"}, "counterfeit_skimming"),
        ({}, {"channel": "card_not_present"}, "card_not_present"),
        ({}, {"channel": "branch"}, "undetermined"),
    ]
    for indicators, fields, expected in cases:
        assert L.classify(claim(indicators=indicators, **fields)) == expected


def test_evaluate_priority_flags_and_deadlines():
    c = claim(amount=6000, n_txns=3, received="2026-10-16", txn_date="2026-10-15", channel="mobile_app",
              indicators={"authorized_payment": True})
    ev = L.evaluate(c)
    assert ev["typology"] == "app_scam" and ev["priority"] == "high"
    assert ("flag_recall_window", {}) in ev["flags"]
    # viernes 16 → 2 días hábiles = martes 20; 10 hábiles = viernes 30; 45 naturales = 30 nov
    assert ev["deadlines"] == {"ack": "2026-10-20", "provisional": "2026-10-30", "resolution": "2026-11-30"}
    assert "act_recall_funds" in ev["actions"] and "doc_police" in ev["documents"]


def test_consistency_flags():
    c = claim(amount=50, received="2026-10-01", txn_date="2026-03-01", prior_claims=2,
              indicators={"card_in_possession": True, "card_lost_stolen": True})
    keys = [k for k, _ in L.evaluate(c)["flags"]]
    assert {"flag_very_late", "flag_repeat", "flag_contradiction_card"} <= set(keys)
    assert L.evaluate(c)["consistency_risk"] >= 70


def test_case_prompt_is_anonymized():
    c = claim(customer_name="Laura Méndez", customer_id="C-1007", account="ES91 2100 0418 4502 0005 1332",
              narrative="Soy Laura Méndez, mi correo es laura@mail.com", amount=10)
    text = L.case_prompt(c, L.evaluate(c), Anonymizer())
    assert "Laura" not in text and "laura@mail.com" not in text and "C-1007" not in text and "ES91" not in text


def test_store_roundtrip(tmp_path):
    store = L.CaseStore(tmp_path / "c.db")
    c = claim(id=L.new_id(), customer_name="Ana", amount=10)
    c["evaluation"] = L.evaluate(c)
    store.save(c)
    assert store.get(c["id"])["customer_name"] == "Ana" and L.new_id(store.ids()) != c["id"]
    c["status"] = "closed"
    store.save(c)
    assert len(store.get(c["id"])["history"]) == 2
    rows = L.case_rows(store.all(), today=date.today())
    assert rows[0]["days_left"] is None  # cerrado
    store.set_meta("sla", {"ack_days": 3})
    assert store.get_meta("sla") == {"ack_days": 3}
    store.delete(c["id"])
    assert store.all() == []


def test_batch_detects_patterns(tmp_path):
    for lang in ("es", "en", "fr"):
        df = samples.fraud_claims(lang)
        mapping = L.guess_mapping(df.columns)
        assert set(mapping) == set(L.FIELDS), lang
        detail, result = L.analyze_batch(df, mapping)
        s = result["summary"]
        assert s["cpp_merchants"] == ["Gasolinera Ruta 9"] and s["duplicates"] >= 2 and s["late"] >= 1
        assert detail["_score"].is_monotonic_decreasing
    labels = {k: k for k in ("sheet_triage", "score", "flags", "sheet_summary", "k_claims", "k_amount", "k_flagged",
                             "k_duplicates", "k_repeat", "k_late", "k_high", "k_cpp", "channel", "sheet_merchants",
                             "merchant", "customers", "peak_30d", "peak_start", "suspect")}
    path = L.export_batch(tmp_path / "b.xlsx", detail, result, labels, {"b_duplicate": "Dup"})
    assert load_workbook(path).sheetnames == ["sheet_triage", "sheet_summary", "sheet_merchants"]
    assert "b_duplicate" in L.batch_context(result, mapping) or "Duplicates" in L.batch_context(result, mapping)
