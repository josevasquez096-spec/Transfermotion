from datetime import date

from docx import Document

from suite import samples
from suite.docpilot import logic as L


def test_markup_roundtrip(tmp_path):
    text = "# Informe\n> Acme\n## Resultados\nVentas **altas**.\n- uno\n1. paso"
    path = L.markup_to_docx(text, tmp_path / "a.docx", "es", "Acme")
    doc = Document(path)
    assert doc.paragraphs[0].style.name == "Title"
    assert any(r.bold for p in doc.paragraphs for r in p.runs if r.text == "altas")
    assert "NUMPAGES" in doc.sections[0].footer._element.xml
    assert L.read_docx_text(path).splitlines()[:3] == ["# Informe", "Acme", "## Resultados"]


def test_placeholder_split_across_runs_is_replaced(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Hola ")
    p.add_run("{{Nom")
    p.add_run("bre}}").bold = True
    p.add_run(", su saldo es {{Importe}}.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Fecha: {{Fecha}}"
    doc.sections[0].header.paragraphs[0].text = "Ref {{Nombre}}"
    tpl = tmp_path / "t.docx"
    doc.save(tpl)
    assert L.find_placeholders(tpl) == ["Nombre", "Importe", "Fecha"]
    import pandas as pd

    df = pd.DataFrame({"Nombre": ["Ana"], "Importe": [1234.5], "Fecha": [pd.Timestamp("2026-11-15")]})
    [out] = L.mail_merge(tpl, df, tmp_path / "out", "es", name_col="Nombre")
    result = Document(out)
    assert result.paragraphs[0].text == "Hola Ana, su saldo es 1.234,50."
    assert result.tables[0].cell(0, 0).text == "Fecha: 15 de noviembre de 2026"
    assert result.sections[0].header.paragraphs[0].text == "Ref Ana"
    assert out.name == "Ana.docx"


def test_mail_merge_single_file_and_sample(tmp_path):
    for lang in ("es", "en", "fr"):
        tpl = L.sample_template(tmp_path / f"t_{lang}.docx", lang)
        df = samples.merge_data(lang)
        assert set(L.find_placeholders(tpl)) == {str(c) for c in df.columns}
        [single] = L.mail_merge(tpl, df, tmp_path / lang, lang, single_file=True)
        text = L.read_docx_text(single)
        assert "{{" not in text and all(name in text for name in df.iloc[:, 0])


def test_formats_and_readability():
    assert L.format_number(1234.5, "es") == "1.234,50"
    assert L.format_number(1234.5, "en") == "1,234.50"
    assert L.format_date(date(2026, 3, 5), "fr") == "5 mars 2026"
    easy = L.text_stats("El gato come. La casa es azul. Hoy hace sol.", "es")
    hard = L.text_stats("La implementación de la metodología organizacional interdepartamental requiere "
                        "considerablemente una coordinación administrativa extraordinariamente meticulosa.", "es")
    assert easy["score"] > hard["score"] and easy["sentences"] == 3
    assert L.readability_level(90) == "very_easy" and L.readability_level(10) == "hard"


def test_skeleton_and_draft_markup():
    sk = L.skeleton("minutes", "en", "Weekly sync")
    assert sk.startswith("# Weekly sync") and "## Action items (owner and date)" in sk
    data = {"title": "T", "subtitle": "S", "sections": [{"heading": "H", "paragraphs": ["p"], "bullets": ["b"]}]}
    assert L.draft_to_markup(data) == "# T\n> S\n\n## H\np\n- b"
