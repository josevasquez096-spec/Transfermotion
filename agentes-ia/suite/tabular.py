"""Lectura de Excel/CSV y escritura de hojas de Excel con formato profesional."""

import csv
import io
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EXCEL_EXT = (".xlsx", ".xlsm", ".xls")
CSV_EXT = (".csv", ".txt", ".tsv")


def list_sheets(path):
    path = Path(path)
    if path.suffix.lower() in EXCEL_EXT:
        return list(pd.ExcelFile(path).sheet_names)
    return []


def read_csv(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    return pd.read_csv(io.StringIO(text), sep=delimiter)


def read_table(path, sheet=None):
    path = Path(path)
    if path.suffix.lower() in EXCEL_EXT:
        df = pd.read_excel(path, sheet_name=sheet if sheet is not None else 0)
    else:
        df = read_csv(path)
    df.columns = [str(c).strip() for c in df.columns]
    return df


def read_all_sheets(path):
    """{nombre de hoja: DataFrame} de un Excel, o {nombre de archivo: DataFrame} de un CSV."""
    path = Path(path)
    if path.suffix.lower() in EXCEL_EXT:
        frames = pd.read_excel(path, sheet_name=None)
        out = {}
        for name, df in frames.items():
            if df.dropna(how="all").empty:
                continue
            df.columns = [str(c).strip() for c in df.columns]
            out[str(name)] = df
        return out
    return {path.stem: read_table(path)}


def is_text(series):
    """Columna de texto (object en pandas 2, StringDtype por defecto en pandas 3)."""
    return series.dtype == object or isinstance(series.dtype, pd.StringDtype)


THIN = Side(style="thin", color="D0D7DE")


def write_df(ws, df, accent="217346", start_row=1, start_col=1, number_format="#,##0.00", as_table=True):
    """Escribe un DataFrame con encabezado de color, filtros, bordes y anchos ajustados."""
    header_fill = PatternFill("solid", fgColor=accent)
    header_font = Font(bold=True, color="FFFFFF")
    zebra = PatternFill("solid", fgColor="F3F6F9")
    for j, col in enumerate(df.columns):
        cell = ws.cell(row=start_row, column=start_col + j, value=str(col))
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            cell = ws.cell(row=start_row + i, column=start_col + j, value=excel_value(value))
            cell.border = Border(bottom=THIN)
            if isinstance(cell.value, float):
                cell.number_format = number_format
            elif hasattr(cell.value, "year") and hasattr(cell.value, "month"):
                cell.number_format = "yyyy-mm-dd"
            if i % 2 == 0:
                cell.fill = zebra
    if as_table and len(df.columns):
        last = get_column_letter(start_col + len(df.columns) - 1)
        ws.auto_filter.ref = f"{get_column_letter(start_col)}{start_row}:{last}{start_row + len(df)}"
        ws.freeze_panes = ws.cell(row=start_row + 1, column=start_col)
    autofit(ws)


def excel_value(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):  # tipos de numpy
        return value.item()
    return value


def autofit(ws, min_width=8, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = cell.value.strftime("%Y-%m-%d") if hasattr(cell.value, "strftime") else str(cell.value)
            length = max(len(line) for line in text.split("\n")) if text else 0
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), length)
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = max(min_width, min(max_width, width + 2))


def title_cell(ws, ref, text, size=16, color="1F2937"):
    cell = ws[ref]
    cell.value = text
    cell.font = Font(bold=True, size=size, color=color)
    return cell
