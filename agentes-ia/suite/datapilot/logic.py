"""DataPilot: lógica sin interfaz (perfil, limpieza, consultas, fórmulas e informe de Excel)."""

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..ai import STR, STR_LIST, obj
from ..tabular import is_text, read_table, title_cell, write_df

# ---------------------------------------------------------------- perfil


def col_kind(series):
    s = series.dropna()
    if s.empty:
        return "empty"
    if pd.api.types.is_bool_dtype(s):
        return "bool"
    if pd.api.types.is_numeric_dtype(s):
        return "number"
    if pd.api.types.is_datetime64_any_dtype(s):
        return "date"
    return "text"


def _example(values):
    if not values.size:
        return ""
    first = values.iloc[0]
    return first.strftime("%Y-%m-%d") if isinstance(first, pd.Timestamp) else str(first)


def profile(df):
    rows = []
    for col in df.columns:
        s = df[col]
        kind = col_kind(s)
        non_null = s.dropna()
        row = {
            "column": str(col), "kind": kind, "filled": int(non_null.size), "empty": int(s.isna().sum()),
            "unique": int(non_null.astype(str).nunique()), "example": _example(non_null),
            "min": "", "max": "", "mean": "",
        }
        if kind == "number":
            row.update(min=float(non_null.min()), max=float(non_null.max()), mean=float(non_null.mean()))
        elif kind == "date":
            row.update(min=non_null.min().strftime("%Y-%m-%d"), max=non_null.max().strftime("%Y-%m-%d"))
        rows.append(row)
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- limpieza

_CURRENCY = re.compile(r"[\s$€£¥%]|USD|EUR|MXN|COP|CAD| ")


def parse_number(value):
    """'1.234,56' → 1234.56, '$ 1,234.5' → 1234.5, '12' → 12.0; None si no es un número."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, np.integer, np.floating)):
        return None if pd.isna(value) else float(value)
    text = _CURRENCY.sub("", str(value).strip())
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    if "," in text and "." in text:
        decimal = "," if text.rfind(",") > text.rfind(".") else "."
        thousands = "." if decimal == "," else ","
        text = text.replace(thousands, "").replace(decimal, ".")
    elif "," in text:
        text = text.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", text) else text.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3}){2,}", text):
        text = text.replace(".", "")
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


CLEAN_DEFAULTS = {"trim": True, "drop_empty": True, "dedupe": True, "numbers": True, "dates": True, "unify": True,
                  "case": "none"}


def clean(df, options=None, dayfirst=True):
    """Devuelve (df_limpio, informe) sin modificar el original."""
    opts = {**CLEAN_DEFAULTS, **(options or {})}
    out = df.copy()
    report = {"trimmed": 0, "empty_rows": 0, "empty_cols": 0, "duplicates": 0, "numbers": [], "dates": [], "unified": 0,
              "case": 0}

    text_cols = [c for c in out.columns if is_text(out[c])]
    if opts["trim"]:
        for c in text_cols:
            mask = out[c].map(lambda v: isinstance(v, str))
            before = out.loc[mask, c]
            after = before.str.strip().str.replace(r"\s{2,}", " ", regex=True)
            report["trimmed"] += int((before != after).sum())
            out.loc[mask, c] = after
            out.loc[mask & (out[c] == ""), c] = None
    if opts["drop_empty"]:
        n = len(out)
        out = out.dropna(how="all")
        report["empty_rows"] = n - len(out)
        empty_cols = [c for c in out.columns if out[c].isna().all()]
        out = out.drop(columns=empty_cols)
        report["empty_cols"] = len(empty_cols)
    if opts["numbers"]:
        for c in [c for c in out.columns if is_text(out[c])]:
            values = out[c].dropna()
            if values.empty or values.map(lambda v: isinstance(v, (pd.Timestamp,))).any():
                continue
            parsed = values.map(parse_number)
            if parsed.notna().mean() >= 0.9:
                out[c] = out[c].map(parse_number).astype(float)
                if (out[c].dropna() % 1 == 0).all():
                    out[c] = out[c].astype("Int64")
                report["numbers"].append(str(c))
    if opts["dates"]:
        for c in [c for c in out.columns if is_text(out[c])]:
            values = out[c].dropna()
            if values.empty:
                continue
            if values.map(lambda v: isinstance(v, (pd.Timestamp,)) or hasattr(v, "year")).all():
                out[c] = pd.to_datetime(out[c], errors="coerce")
                report["dates"].append(str(c))
                continue
            if not values.map(lambda v: isinstance(v, str) and bool(re.search(r"\d", v))).all():
                continue
            parsed = pd.to_datetime(values, errors="coerce", dayfirst=dayfirst, format="mixed")
            if parsed.notna().mean() >= 0.9:
                out[c] = pd.to_datetime(out[c], errors="coerce", dayfirst=dayfirst, format="mixed")
                report["dates"].append(str(c))
    if opts["unify"]:  # "NORTE", "norte" y "Norte" → la forma más frecuente
        for c in [c for c in out.columns if is_text(out[c])]:
            mask = out[c].map(lambda v: isinstance(v, str))
            values = out.loc[mask, c]
            if values.empty:
                continue
            best = values.groupby(values.str.lower()).agg(lambda v: v.value_counts().index[0])
            after = values.str.lower().map(best)
            report["unified"] += int((values != after).sum())
            out.loc[mask, c] = after
    if opts["case"] in ("title", "upper", "lower"):
        for c in [c for c in out.columns if is_text(out[c])]:
            mask = out[c].map(lambda v: isinstance(v, str))
            before = out.loc[mask, c]
            after = getattr(before.str, opts["case"])()
            report["case"] += int((before != after).sum())
            out.loc[mask, c] = after
    if opts["dedupe"]:
        n = len(out)
        out = out.drop_duplicates()
        report["duplicates"] = n - len(out)
    return out.reset_index(drop=True), report


def merge_files(paths, source_label="Source"):
    """Une varios Excel/CSV en una sola tabla (unión de columnas) y añade el archivo de origen."""
    frames = []
    for path in paths:
        df = read_table(path)
        df.insert(0, source_label, Path(path).name)
        frames.append(df)
    return pd.concat(frames, ignore_index=True, sort=False) if frames else pd.DataFrame()


# ---------------------------------------------------------------- tabla dinámica y consultas

AGGS = ("sum", "mean", "count", "min", "max", "median", "nunique")
DATE_PARTS = ("none", "year", "quarter", "month", "year_month", "weekday")
OPS = ("==", "!=", ">", ">=", "<", "<=", "contains", "not_contains", "in", "between", "is_null", "not_null")


def pivot(df, rows, value, agg="sum", columns=None):
    if agg == "count":
        result = df.pivot_table(index=rows, columns=columns or None, values=value, aggfunc="count", fill_value=0)
    else:
        data = df.copy()
        data[value] = pd.to_numeric(data[value], errors="coerce")
        result = data.pivot_table(index=rows, columns=columns or None, values=value, aggfunc=agg, fill_value=0)
    if isinstance(result, pd.Series):
        result = result.to_frame()
    if columns:
        result.columns = [str(c) for c in result.columns]
    else:
        result.columns = [f"{agg}({value})"]
    result = result.reset_index()
    return result


def resolve_column(df, name):
    if name in df.columns:
        return name
    norm = str(name).strip().lower()
    for col in df.columns:
        if str(col).strip().lower() == norm:
            return col
    raise KeyError(name)


def _coerce(series, value):
    if value is None:
        return None
    if pd.api.types.is_numeric_dtype(series):
        number = parse_number(value)
        return number if number is not None else value
    if pd.api.types.is_datetime64_any_dtype(series):
        try:
            return pd.Timestamp(value)
        except (ValueError, TypeError):
            return value
    return value


def _date_part(series, part):
    s = pd.to_datetime(series, errors="coerce")
    if part == "year":
        return s.dt.year
    if part == "quarter":
        return s.dt.year.astype("Int64").astype(str) + "-Q" + s.dt.quarter.astype("Int64").astype(str)
    if part == "month":
        return s.dt.month
    if part == "year_month":
        return s.dt.strftime("%Y-%m")
    if part == "weekday":
        return s.dt.dayofweek + 1
    return series


def apply_filter(df, flt):
    col = resolve_column(df, flt["column"])
    s = df[col]
    op = flt["op"]
    value = _coerce(s, flt.get("value"))
    if op == "is_null":
        return df[s.isna()]
    if op == "not_null":
        return df[s.notna()]
    if op in ("contains", "not_contains"):
        mask = s.astype(str).str.contains(str(flt.get("value", "")), case=False, regex=False, na=False)
        return df[mask if op == "contains" else ~mask]
    if op == "in":
        options = [_coerce(s, v.strip()) for v in str(flt.get("value", "")).split("|")]
        if is_text(s):
            lowered = {str(o).lower() for o in options}
            return df[s.astype(str).str.lower().isin(lowered)]
        return df[s.isin(options)]
    if op == "between":
        upper = _coerce(s, flt.get("value2"))
        return df[(s >= value) & (s <= upper)]
    if is_text(s) and op in ("==", "!="):
        mask = s.astype(str).str.strip().str.lower() == str(value).strip().lower()
        return df[mask if op == "==" else ~mask]
    ops = {"==": s.__eq__, "!=": s.__ne__, ">": s.__gt__, ">=": s.__ge__, "<": s.__lt__, "<=": s.__le__}
    return df[ops[op](value).fillna(False)]


def run_plan(df, plan):
    """Ejecuta un plan de consulta (generado por la IA) con pandas, sin ejecutar código arbitrario."""
    data = df.copy()
    for flt in plan.get("filters") or []:
        data = apply_filter(data, flt)
    keys = []
    for g in plan.get("group_by") or []:
        col = resolve_column(data, g["column"])
        part = g.get("date_part", "none")
        if part and part != "none":
            name = f"{col} ({part})"
            data[name] = _date_part(data[col], part)
            keys.append(name)
        else:
            keys.append(col)
    metrics = plan.get("metrics") or []
    if metrics:
        named = {}
        for m in metrics:
            agg = m["agg"] if m["agg"] in AGGS else "sum"
            col = resolve_column(data, m["column"]) if m.get("column") else data.columns[0]
            if agg not in ("count", "nunique"):
                data[col] = pd.to_numeric(data[col], errors="coerce")
            alias = m.get("alias") or f"{agg}({col})"
            named[alias] = pd.NamedAgg(column=col, aggfunc=agg)
        if keys:
            result = data.groupby(keys, dropna=False).agg(**named).reset_index()
        else:
            result = pd.DataFrame([{k: data[v.column].agg(v.aggfunc) for k, v in named.items()}])
    else:
        result = data
        cols = [resolve_column(data, c) for c in plan.get("select_columns") or []]
        if keys:
            result = data.groupby(keys, dropna=False).size().reset_index(name="count")
        elif cols:
            result = data[cols]
    sort_by = plan.get("sort_by")
    if sort_by:
        try:
            col = resolve_column(result, sort_by)
            result = result.sort_values(col, ascending=bool(plan.get("ascending", False)))
        except KeyError:
            pass
    limit = plan.get("limit") or 0
    if limit and limit > 0:
        result = result.head(int(limit))
    return result.reset_index(drop=True)


PLAN_SCHEMA = obj({
    "explanation": STR,
    "filters": {"type": "array", "items": obj({
        "column": STR, "op": {"type": "string", "enum": list(OPS)}, "value": STR, "value2": STR,
    })},
    "group_by": {"type": "array", "items": obj({
        "column": STR, "date_part": {"type": "string", "enum": list(DATE_PARTS)},
    })},
    "metrics": {"type": "array", "items": obj({
        "column": STR, "agg": {"type": "string", "enum": list(AGGS)}, "alias": STR,
    })},
    "select_columns": STR_LIST,
    "sort_by": STR,
    "ascending": {"type": "boolean"},
    "limit": {"type": "integer"},
})

PLAN_SYSTEM = """You are DataPilot, an analyst that turns business questions about a spreadsheet into a query plan.
Use ONLY the column names listed, spelled exactly. The plan is executed with pandas:
filters (AND) → group_by (optionally by a date part) → metrics (aggregations) → sort → limit.
For "in" filters separate values with "|". Use value2 only for "between"; otherwise empty string.
If no grouping/metrics are needed, list the columns to show in select_columns.
sort_by must be a column or metric alias of the result (or empty). limit 0 means no limit.
In `explanation`, briefly say what the result shows and any assumption you made."""


def schema_summary(df, max_values=12):
    lines = [f"Rows: {len(df)}"]
    for col in df.columns:
        s = df[col].dropna()
        kind = col_kind(df[col])
        line = f"- {col} ({kind})"
        if kind == "number" and not s.empty:
            line += f": min {s.min():g}, max {s.max():g}"
        elif kind == "date" and not s.empty:
            line += f": from {s.min():%Y-%m-%d} to {s.max():%Y-%m-%d}"
        elif kind == "text" and not s.empty:
            uniques = s.astype(str).unique()
            sample = ", ".join(uniques[:max_values])
            line += f": {len(uniques)} distinct values, e.g. {sample}"
        lines.append(line)
    return "\n".join(lines)


def insights_context(df, max_groups=8):
    """Resumen estadístico compacto para que la IA redacte conclusiones (sin enviar filas completas)."""
    parts = [schema_summary(df)]
    numeric = [c for c in df.columns if col_kind(df[c]) == "number"]
    texts = [c for c in df.columns if col_kind(df[c]) == "text" and 1 < df[c].nunique() <= 30]
    dates = [c for c in df.columns if col_kind(df[c]) == "date"]
    if numeric:
        parts.append("Numeric totals:\n" + df[numeric].sum(numeric_only=True).round(2).to_string())
    for t in texts[:4]:
        for n in numeric[:2]:
            top = df.groupby(t)[n].sum().sort_values(ascending=False).head(max_groups).round(2)
            parts.append(f"{n} by {t}:\n{top.to_string()}")
    if dates and numeric:
        monthly = df.set_index(pd.to_datetime(df[dates[0]]))[numeric[0]].resample("MS").sum().round(2)
        monthly.index = monthly.index.strftime("%Y-%m")
        parts.append(f"{numeric[0]} per month ({dates[0]}):\n{monthly.tail(24).to_string()}")
    return "\n\n".join(parts)


INSIGHTS_SYSTEM = """You are DataPilot, a senior business analyst. From the dataset summary, write a concise executive
analysis: 1) three to six key findings with figures, 2) trends or anomalies worth checking, 3) concrete recommended
actions. Use short headings and bullet points in plain text (no markdown tables). Only use the numbers provided."""

# ---------------------------------------------------------------- fórmulas

FUNCTIONS = [
    ("SUM", "SUMA", "SOMME"), ("SUMIF", "SUMAR.SI", "SOMME.SI"), ("SUMIFS", "SUMAR.SI.CONJUNTO", "SOMME.SI.ENS"),
    ("SUMPRODUCT", "SUMAPRODUCTO", "SOMMEPROD"), ("AVERAGE", "PROMEDIO", "MOYENNE"),
    ("AVERAGEIF", "PROMEDIO.SI", "MOYENNE.SI"), ("AVERAGEIFS", "PROMEDIO.SI.CONJUNTO", "MOYENNE.SI.ENS"),
    ("COUNT", "CONTAR", "NB"), ("COUNTA", "CONTARA", "NBVAL"), ("COUNTBLANK", "CONTAR.BLANCO", "NB.VIDE"),
    ("COUNTIF", "CONTAR.SI", "NB.SI"), ("COUNTIFS", "CONTAR.SI.CONJUNTO", "NB.SI.ENS"),
    ("MAX", "MAX", "MAX"), ("MIN", "MIN", "MIN"), ("MAXIFS", "MAX.SI.CONJUNTO", "MAX.SI.ENS"),
    ("MINIFS", "MIN.SI.CONJUNTO", "MIN.SI.ENS"), ("MEDIAN", "MEDIANA", "MEDIANE"),
    ("ROUND", "REDONDEAR", "ARRONDI"), ("ROUNDUP", "REDONDEAR.MAS", "ARRONDI.SUP"),
    ("ROUNDDOWN", "REDONDEAR.MENOS", "ARRONDI.INF"), ("ABS", "ABS", "ABS"),
    ("IF", "SI", "SI"), ("IFS", "SI.CONJUNTO", "SI.CONDITIONS"), ("IFERROR", "SI.ERROR", "SIERREUR"),
    ("IFNA", "SI.ND", "SI.NON.DISP"), ("AND", "Y", "ET"), ("OR", "O", "OU"), ("NOT", "NO", "NON"),
    ("VLOOKUP", "BUSCARV", "RECHERCHEV"), ("HLOOKUP", "BUSCARH", "RECHERCHEH"),
    ("XLOOKUP", "BUSCARX", "RECHERCHEX"), ("LOOKUP", "BUSCAR", "RECHERCHE"), ("INDEX", "INDICE", "INDEX"),
    ("MATCH", "COINCIDIR", "EQUIV"), ("XMATCH", "COINCIDIRX", "EQUIVX"), ("FILTER", "FILTRAR", "FILTRE"),
    ("SORT", "ORDENAR", "TRIER"), ("UNIQUE", "UNICOS", "UNIQUE"), ("CONCATENATE", "CONCATENAR", "CONCATENER"),
    ("CONCAT", "CONCAT", "CONCAT"), ("TEXTJOIN", "UNIRCADENAS", "JOINDRE.TEXTE"), ("LEFT", "IZQUIERDA", "GAUCHE"),
    ("RIGHT", "DERECHA", "DROITE"), ("MID", "EXTRAE", "STXT"), ("LEN", "LARGO", "NBCAR"),
    ("TRIM", "ESPACIOS", "SUPPRESPACE"), ("UPPER", "MAYUSC", "MAJUSCULE"), ("LOWER", "MINUSC", "MINUSCULE"),
    ("PROPER", "NOMPROPIO", "NOMPROPRE"), ("SUBSTITUTE", "SUSTITUIR", "SUBSTITUE"), ("TEXT", "TEXTO", "TEXTE"),
    ("VALUE", "VALOR", "CNUM"), ("FIND", "ENCONTRAR", "TROUVE"), ("SEARCH", "HALLAR", "CHERCHE"),
    ("TODAY", "HOY", "AUJOURDHUI"), ("NOW", "AHORA", "MAINTENANT"), ("DATE", "FECHA", "DATE"),
    ("YEAR", "AÑO", "ANNEE"), ("MONTH", "MES", "MOIS"), ("DAY", "DIA", "JOUR"), ("WEEKDAY", "DIASEM", "JOURSEM"),
    ("EOMONTH", "FIN.MES", "FIN.MOIS"), ("EDATE", "FECHA.MES", "MOIS.DECALER"),
    ("NETWORKDAYS", "DIAS.LAB", "NB.JOURS.OUVRES"), ("WORKDAY", "DIA.LAB", "SERIE.JOUR.OUVRE"),
    ("DATEDIF", "SIFECHA", "DATEDIF"), ("ISBLANK", "ESBLANCO", "ESTVIDE"), ("ISNUMBER", "ESNUMERO", "ESTNUM"),
    ("ISERROR", "ESERROR", "ESTERREUR"), ("INDIRECT", "INDIRECTO", "INDIRECT"), ("OFFSET", "DESREF", "DECALER"),
    ("ROW", "FILA", "LIGNE"), ("COLUMN", "COLUMNA", "COLONNE"), ("RANK.EQ", "JERARQUIA.EQV", "EQUATION.RANG"),
    ("SUBTOTAL", "SUBTOTALES", "SOUS.TOTAL"), ("AGGREGATE", "AGREGAR", "AGREGAT"),
]
_LANG_INDEX = {"en": 0, "es": 1, "fr": 2}
_SEP = {"en": (",", "."), "es": (";", ","), "fr": (";", ",")}  # (separador de argumentos, decimal)
_NAME = re.compile(r"[^\W\d][\w.]*(?=\()", re.UNICODE)
_NUMBER = {".": re.compile(r"\d+(?:\.\d+)?"), ",": re.compile(r"\d+(?:,\d+)?")}


def translate_formula(formula, src, dst):
    """Traduce nombres de funciones y separadores de una fórmula de Excel entre EN, ES y FR."""
    if src == dst:
        return formula
    names = {row[_LANG_INDEX[src]].upper(): row[_LANG_INDEX[dst]] for row in FUNCTIONS}
    src_arg, src_dec = _SEP[src]
    dst_arg, dst_dec = _SEP[dst]
    out, i, depth_brace = [], 0, 0
    while i < len(formula):
        ch = formula[i]
        if ch in "\"'":  # texto o nombre de hoja: se copia tal cual (las comillas dobles se escapan duplicándolas)
            j = i + 1
            while j < len(formula):
                if formula[j] == ch:
                    if j + 1 < len(formula) and formula[j + 1] == ch:
                        j += 2
                        continue
                    break
                j += 1
            out.append(formula[i:j + 1])
            i = j + 1
            continue
        if ch == "{":
            depth_brace += 1
        elif ch == "}":
            depth_brace = max(0, depth_brace - 1)
        prev = formula[i - 1] if i else ""
        if not (prev.isalnum() or prev in "$_."):
            m = _NAME.match(formula, i)
            if m:
                name = m.group(0)
                out.append(names.get(name.upper(), name))
                i = m.end()
                continue
            m = _NUMBER[src_dec].match(formula, i)
            if m and depth_brace == 0:
                out.append(m.group(0).replace(src_dec, dst_dec))
                i = m.end()
                continue
        if ch == src_arg and depth_brace == 0:
            out.append(dst_arg)
        else:
            out.append(ch)
        i += 1
    return "".join(out)


FORMULA_SCHEMA = obj({"formula_en": STR, "explanation": STR, "steps": STR_LIST, "example": STR})

FORMULA_SYSTEM = """You are DataPilot, an Excel formula expert. Write ONE formula that solves the request, using
English function names, comma argument separators and dot decimals (it is translated automatically to Spanish and
French Excel). Start it with "=". Prefer modern functions (XLOOKUP, FILTER, SUMIFS, LET) unless the user needs
compatibility with old Excel. If a sheet layout is given, use those column letters and assume headers in row 1.
`steps`: how it works, piece by piece. `example`: where to put it and what result to expect."""

EXPLAIN_SYSTEM = """You are DataPilot, an Excel formula expert. Explain the given formula to a business user:
what it returns, each part step by step, common errors and one improvement if useful. Plain text, short bullets."""


def sheet_layout(df):
    return ", ".join(f"{get_column_letter(i + 1)}={c}" for i, c in enumerate(df.columns))


# ---------------------------------------------------------------- informe de Excel


def export_report(path, df, labels, pivot_df=None, insights=None, accent="217346"):
    """Libro con: datos con formato, resumen con gráfico, perfil de columnas y conclusiones de la IA."""
    wb = Workbook()
    ws = wb.active
    ws.title = labels["data"][:31]
    write_df(ws, df, accent=accent)

    if pivot_df is not None and not pivot_df.empty:
        ws2 = wb.create_sheet(labels["summary"][:31])
        title_cell(ws2, "A1", labels["summary"])
        write_df(ws2, pivot_df, accent=accent, start_row=3, as_table=False)
        n = len(pivot_df)
        if 1 <= n <= 60 and len(pivot_df.columns) >= 2:
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = labels["summary"]
            chart.height, chart.width = 9, 18
            last_col = min(len(pivot_df.columns), 7)
            data = Reference(ws2, min_col=2, max_col=last_col, min_row=3, max_row=3 + n)
            cats = Reference(ws2, min_col=1, min_row=4, max_row=3 + n)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws2.add_chart(chart, f"{get_column_letter(len(pivot_df.columns) + 2)}3")

    ws3 = wb.create_sheet(labels["profile"][:31])
    prof = profile(df)
    prof.columns = [labels.get(f"p_{c}", c) for c in prof.columns]
    write_df(ws3, prof, accent=accent)

    if insights:
        ws4 = wb.create_sheet(labels["insights"][:31])
        title_cell(ws4, "A1", labels["insights"])
        ws4.column_dimensions["A"].width = 110
        for i, line in enumerate(insights.splitlines(), start=3):
            cell = ws4.cell(row=i, column=1, value=line)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if line and not line.startswith(("-", "•", " ")) and len(line) < 80:
                cell.font = Font(bold=True)
        ws4.cell(row=len(insights.splitlines()) + 5, column=1, value=labels["disclaimer"]).font = Font(
            italic=True, color="6B7280")
    wb.save(path)
    return path
