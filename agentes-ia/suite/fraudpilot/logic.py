"""FraudPilot: lógica sin interfaz para la gestión de reclamaciones de fraude.

- Evaluación por reglas de una reclamación: tipología, prioridad, señales de alerta, acciones,
  documentos y plazos (configurables según la normativa de cada entidad).
- Registro local de casos en SQLite.
- Análisis masivo de un Excel de reclamaciones: duplicados, reclamantes recurrentes, avisos tardíos,
  importes atípicos y comercios con concentración de casos (posible punto de compromiso).

Las reglas ayudan a priorizar; la decisión final siempre es de una persona.
"""

import json
import re
import sqlite3
from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

from ..ai import STR, STR_LIST, obj
from ..tabular import title_cell, write_df

PRODUCTS = ("credit_card", "debit_card", "transfer", "online_banking", "wallet", "check", "other")
CHANNELS = ("card_present", "card_not_present", "atm", "online_banking", "mobile_app", "phone", "branch")
INDICATORS = ("card_in_possession", "card_lost_stolen", "contacted_by_fraudster", "shared_credentials",
              "authorized_payment", "sim_issue", "new_device", "foreign_txn", "recognizes_merchant",
              "goods_not_received", "not_opened_by_customer", "new_account", "police_report")
TYPOLOGIES = ("merchant_dispute", "identity_theft", "app_scam", "sim_swap", "phishing_vishing", "account_takeover",
              "lost_stolen_card", "counterfeit_skimming", "card_not_present", "undetermined")
STATUSES = ("received", "investigating", "pending_customer", "provisional_credit", "resolved_favor",
            "resolved_rejected", "closed")
OPEN_STATUSES = ("received", "investigating", "pending_customer", "provisional_credit")
DEFAULT_SLA = {"ack_days": 2, "provisional_days": 10, "resolution_days": 45, "late_days": 60, "high_amount": 1000}


def new_claim():
    today = date.today()
    return {
        "id": "", "received": today.isoformat(), "customer_name": "", "customer_id": "", "account": "",
        "product": "credit_card", "channel": "card_not_present", "txn_date": today.isoformat(), "amount": 0.0,
        "currency": "USD", "n_txns": 1, "merchant": "", "narrative": "", "prior_claims": 0,
        "indicators": {k: False for k in INDICATORS}, "status": "received", "notes": "", "evaluation": {},
        "ai_analysis": "",
    }


def new_id(existing=()):
    stamp = date.today().strftime("%Y%m%d")
    n = 1
    while f"FR-{stamp}-{n:03d}" in existing:
        n += 1
    return f"FR-{stamp}-{n:03d}"


def to_date(value):
    if isinstance(value, date):
        return value if not isinstance(value, datetime) else value.date()
    try:
        return date.fromisoformat(str(value).strip()[:10])
    except ValueError:
        return None


def add_business_days(start, days):
    current, added = start, 0
    while added < days:
        current += timedelta(days=1)
        if current.weekday() < 5:
            added += 1
    return current


# ---------------------------------------------------------------- evaluación por reglas

def classify(claim):
    i = claim.get("indicators", {})
    channel, product = claim.get("channel"), claim.get("product")
    if i.get("goods_not_received") or (i.get("recognizes_merchant") and not i.get("contacted_by_fraudster")):
        return "merchant_dispute"
    if i.get("not_opened_by_customer"):
        return "identity_theft"
    if i.get("authorized_payment"):
        return "app_scam"
    if i.get("sim_issue"):
        return "sim_swap"
    if i.get("shared_credentials") or i.get("contacted_by_fraudster"):
        return "phishing_vishing"
    if i.get("new_device") and (product in ("transfer", "online_banking", "wallet") or
                                channel in ("online_banking", "mobile_app")):
        return "account_takeover"
    if i.get("card_lost_stolen"):
        return "lost_stolen_card"
    if channel in ("card_present", "atm") and i.get("card_in_possession", True):
        return "counterfeit_skimming"
    if channel == "card_not_present":
        return "card_not_present"
    return "undetermined"


ACTIONS = {
    "common": ["act_verify_identity", "act_statement"],
    "card_not_present": ["act_block_card", "act_auth_logs", "act_merchant_data", "act_chargeback_fraud"],
    "counterfeit_skimming": ["act_block_card", "act_entry_mode", "act_cpp", "act_chargeback_fraud"],
    "lost_stolen_card": ["act_block_card", "act_loss_timeline", "act_pin_usage", "act_chargeback_fraud"],
    "phishing_vishing": ["act_reset_credentials", "act_auth_logs", "act_negligence", "act_report_phishing"],
    "account_takeover": ["act_lock_access", "act_device_ip", "act_recall_funds", "act_reset_credentials"],
    "sim_swap": ["act_lock_access", "act_carrier", "act_auth_logs", "act_recall_funds"],
    "app_scam": ["act_recall_funds", "act_mule", "act_reimbursement_rules", "act_report_phishing"],
    "identity_theft": ["act_freeze_product", "act_kyc_review", "act_bureau", "act_police"],
    "merchant_dispute": ["act_contact_merchant", "act_dispute_route", "act_chargeback_dispute"],
    "undetermined": ["act_more_info"],
}
DOCUMENTS = {
    "common": ["doc_claim_form", "doc_id", "doc_statement_marked"],
    "card_not_present": [],
    "counterfeit_skimming": ["doc_card_possession"],
    "lost_stolen_card": ["doc_loss_report"],
    "phishing_vishing": ["doc_evidence_messages"],
    "account_takeover": ["doc_evidence_messages"],
    "sim_swap": ["doc_carrier_letter", "doc_police"],
    "app_scam": ["doc_evidence_messages", "doc_transfer_receipt", "doc_police"],
    "identity_theft": ["doc_identity_affidavit", "doc_police"],
    "merchant_dispute": ["doc_merchant_contact", "doc_purchase_proof"],
    "undetermined": [],
}


def evaluate(claim, sla=None):
    """Tipología, prioridad, señales de alerta, acciones, documentos y plazos de una reclamación."""
    sla = {**DEFAULT_SLA, **(sla or {})}
    i = claim.get("indicators", {})
    typology = classify(claim)
    amount = float(claim.get("amount") or 0)
    n_txns = int(claim.get("n_txns") or 1)
    received = to_date(claim.get("received")) or date.today()
    txn_date = to_date(claim.get("txn_date"))
    flags = []  # (clave, parámetros, peso) — señales que piden revisar la reclamación con más cuidado
    urgency = 0
    if amount >= sla["high_amount"] * 5:
        urgency += 35
    elif amount >= sla["high_amount"]:
        urgency += 25
    elif amount >= sla["high_amount"] / 4:
        urgency += 10
    if n_txns >= 3:
        urgency += 15
    if typology in ("account_takeover", "sim_swap", "app_scam", "identity_theft"):
        urgency += 20  # la exposición puede seguir abierta
    if typology == "app_scam" and txn_date and (received - txn_date).days <= 2:
        urgency += 20
        flags.append(("flag_recall_window", {}, 0))
    if i.get("foreign_txn"):
        urgency += 5
    delay = (received - txn_date).days if txn_date else 0
    if delay > sla["late_days"] * 2:
        flags.append(("flag_very_late", {"days": delay}, 25))
    elif delay > sla["late_days"]:
        flags.append(("flag_late", {"days": delay}, 15))
    prior = int(claim.get("prior_claims") or 0)
    if prior >= 2:
        flags.append(("flag_repeat", {"n": prior}, 25))
    elif prior == 1:
        flags.append(("flag_prior", {}, 10))
    if i.get("recognizes_merchant") and typology != "merchant_dispute":
        flags.append(("flag_recognizes", {}, 20))
    if i.get("card_in_possession") and i.get("card_lost_stolen"):
        flags.append(("flag_contradiction_card", {}, 20))
    if i.get("new_account"):
        flags.append(("flag_new_account", {}, 10))
    if i.get("shared_credentials"):
        flags.append(("flag_negligence", {}, 0))
    if txn_date and txn_date > received:
        flags.append(("flag_future_date", {}, 15))
    if typology == "merchant_dispute":
        flags.append(("flag_not_fraud", {}, 0))
    consistency = min(100, sum(w for _, _, w in flags))
    priority = "high" if urgency >= 45 else "medium" if urgency >= 20 else "low"
    actions = ACTIONS["common"] + ACTIONS[typology]
    if not i.get("police_report") and typology in ("identity_theft", "sim_swap", "app_scam"):
        actions = actions + ["act_police"] if "act_police" not in actions else actions
    documents = DOCUMENTS["common"] + DOCUMENTS[typology]
    if i.get("police_report"):
        documents = [d for d in documents if d != "doc_police"] + ["doc_police_attached"]
    deadlines = {
        "ack": add_business_days(received, sla["ack_days"]).isoformat(),
        "provisional": add_business_days(received, sla["provisional_days"]).isoformat(),
        "resolution": (received + timedelta(days=sla["resolution_days"])).isoformat(),
    }
    return {
        "typology": typology, "priority": priority, "urgency": min(100, urgency), "consistency_risk": consistency,
        "flags": [(k, p) for k, p, _ in flags], "actions": list(dict.fromkeys(actions)),
        "documents": list(dict.fromkeys(documents)), "deadlines": deadlines, "delay_days": delay,
    }


# ---------------------------------------------------------------- IA

CASE_SCHEMA = obj({
    "summary": STR,
    "typology": {"type": "string", "enum": list(TYPOLOGIES)},
    "risk_level": {"type": "string", "enum": ["low", "medium", "high"]},
    "key_facts": STR_LIST,
    "inconsistencies": STR_LIST,
    "missing_information": STR_LIST,
    "recommended_actions": STR_LIST,
    "liability_view": STR,
    "questions_for_customer": STR_LIST,
})
CASE_SYSTEM = """You are FraudPilot, a senior fraud-claims investigator at a bank / payments company. Analyze the
customer's fraud claim (form data, indicators and the customer's own statement). Personal data may be replaced by
tokens like [CLIENTE_1]; keep the tokens as they are. Be factual and neutral: do not accuse the customer.
- typology: the most likely fraud pattern (merchant_dispute = not fraud, a commercial dispute).
- risk_level: urgency/exposure for the institution and the customer.
- inconsistencies: contradictions between statement, dates, channel and indicators (empty if none).
- missing_information / questions_for_customer: what an investigator should obtain next.
- recommended_actions: concrete next steps (blocks, logs to review, recall requests, chargeback route…).
- liability_view: a preliminary, non-binding view of who likely bears the loss under typical rules, noting that it
  depends on local regulation and the institution's policy.
This is decision support only; a human investigator decides."""

LETTER_SYSTEM = """You are FraudPilot. Improve the customer letter: clear, empathetic and professional, keeping every
fact, date, amount, reference number and token like [CLIENTE_1] exactly as written. Do not promise outcomes or
add commitments that are not in the original. Keep the same language and the simple markup (# title, paragraphs,
- bullets). Return only the letter."""

BATCH_SYSTEM = """You are FraudPilot, a fraud analytics lead. From the aggregated statistics of a fraud-claims
portfolio, write a brief for management: 1) main patterns (channels, merchants, timing), 2) possible points of
compromise or organized activity, 3) signs of first-party / friendly fraud to review, 4) operational
recommendations (controls, rules, customer education). Use short headings and bullets, plain text."""


def case_prompt(claim, evaluation, anonymizer=None, t=None):
    lines = [
        f"Claim ID: {claim.get('id')}",
        f"Received: {claim.get('received')} | Transaction date: {claim.get('txn_date')}",
        f"Customer: {claim.get('customer_name')} ({claim.get('customer_id')}) | Account: {claim.get('account')}",
        f"Product: {claim.get('product')} | Channel: {claim.get('channel')}",
        f"Amount: {claim.get('amount')} {claim.get('currency')} in {claim.get('n_txns')} transaction(s)",
        f"Merchant / beneficiary: {claim.get('merchant')}",
        f"Prior claims (12 months): {claim.get('prior_claims')}",
        "Indicators: " + ", ".join(f"{k}={'yes' if v else 'no'}" for k, v in claim.get("indicators", {}).items()),
        f"Rule-based pre-assessment: typology={evaluation['typology']}, priority={evaluation['priority']}, "
        f"flags={[k for k, _ in evaluation['flags']]}",
        f"Customer statement:\n<<<\n{claim.get('narrative') or '-'}\n>>>",
    ]
    text = "\n".join(lines)
    if anonymizer is not None:
        anonymizer.add_known(claim.get("customer_name"), "CLIENTE")
        anonymizer.add_known(claim.get("customer_id"), "ID_CLIENTE")
        anonymizer.add_known(claim.get("account"), "CUENTA")
        text = anonymizer.mask(text)
    return text


# ---------------------------------------------------------------- base de datos local

class CaseStore:
    def __init__(self, path):
        self.path = str(path)
        with self._conn() as db:
            db.execute("CREATE TABLE IF NOT EXISTS cases (id TEXT PRIMARY KEY, data TEXT NOT NULL, "
                       "updated TEXT NOT NULL)")
            db.execute("CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT NOT NULL)")

    def _conn(self):
        return sqlite3.connect(self.path)

    def save(self, claim):
        claim = dict(claim)
        previous = self.get(claim["id"]) or {}
        history = list(claim.get("history") or previous.get("history") or [])
        history.append({"at": datetime.now().isoformat(timespec="seconds"), "status": claim.get("status")})
        claim["history"] = history[-50:]
        with self._conn() as db:
            db.execute("INSERT OR REPLACE INTO cases (id, data, updated) VALUES (?, ?, ?)",
                       (claim["id"], json.dumps(claim, ensure_ascii=False), datetime.now().isoformat()))
        return claim

    def get(self, case_id):
        with self._conn() as db:
            row = db.execute("SELECT data FROM cases WHERE id = ?", (case_id,)).fetchone()
        return json.loads(row[0]) if row else None

    def all(self):
        with self._conn() as db:
            rows = db.execute("SELECT data FROM cases ORDER BY updated DESC").fetchall()
        return [json.loads(r[0]) for r in rows]

    def ids(self):
        with self._conn() as db:
            return [r[0] for r in db.execute("SELECT id FROM cases").fetchall()]

    def delete(self, case_id):
        with self._conn() as db:
            db.execute("DELETE FROM cases WHERE id = ?", (case_id,))

    def get_meta(self, key, default=None):
        with self._conn() as db:
            row = db.execute("SELECT value FROM meta WHERE key = ?", (key,)).fetchone()
        return json.loads(row[0]) if row else default

    def set_meta(self, key, value):
        with self._conn() as db:
            db.execute("INSERT OR REPLACE INTO meta (key, value) VALUES (?, ?)", (key, json.dumps(value)))


def case_rows(cases, today=None):
    today = today or date.today()
    rows = []
    for c in cases:
        ev = c.get("evaluation") or {}
        due = to_date((ev.get("deadlines") or {}).get("resolution"))
        open_case = c.get("status") in OPEN_STATUSES
        rows.append({
            "id": c.get("id"), "received": c.get("received"), "customer": c.get("customer_name"),
            "typology": ev.get("typology", ""), "amount": float(c.get("amount") or 0), "currency": c.get("currency"),
            "priority": ev.get("priority", ""), "status": c.get("status"), "due": due.isoformat() if due else "",
            "days_left": (due - today).days if due and open_case else None,
        })
    return rows


# ---------------------------------------------------------------- análisis masivo

FIELDS = ("claim_id", "customer_id", "claim_date", "txn_date", "amount", "channel", "merchant")
SYNONYMS = {
    "claim_id": ["claim", "reclam", "caso", "case", "ticket", "dossier", "réclamation", "folio"],
    "customer_id": ["customer id", "id cliente", "cliente id", "client id", "id client", "customer", "cliente",
                    "client"],
    "claim_date": ["claim date", "fecha reclam", "date réclam", "fecha de reclam", "reported", "fecha aviso",
                   "date de réclamation"],
    "txn_date": ["transaction date", "fecha trans", "date trans", "fecha operación", "fecha operacion",
                 "date opération", "txn"],
    "amount": ["amount", "importe", "monto", "montant", "valor", "value"],
    "channel": ["channel", "canal"],
    "merchant": ["merchant", "comercio", "commerçant", "commercant", "beneficiar", "establecimiento"],
}


def guess_mapping(columns):
    mapping, used = {}, set()
    norm = {c: re.sub(r"[_\-.]+", " ", str(c)).lower().strip() for c in columns}
    for field in FIELDS:
        for syn in SYNONYMS[field]:
            hit = next((c for c in columns if c not in used and syn in norm[c]), None)
            if hit is not None:
                mapping[field] = hit
                used.add(hit)
                break
    return mapping


def analyze_batch(df, mapping, sla=None, cluster_min=5):
    """Añade puntuación de revisión y señales a cada reclamación; devuelve (detalle, resumen)."""
    sla = {**DEFAULT_SLA, **(sla or {})}
    m = mapping
    out = df.copy()
    amount = pd.to_numeric(out[m["amount"]], errors="coerce").fillna(0) if m.get("amount") else pd.Series(0, out.index)
    claim_date = pd.to_datetime(out[m["claim_date"]], errors="coerce") if m.get("claim_date") else None
    txn_date = pd.to_datetime(out[m["txn_date"]], errors="coerce") if m.get("txn_date") else None
    flags = [[] for _ in range(len(out))]
    score = pd.Series(0, index=out.index, dtype=float)

    def mark(mask, key, weight):
        nonlocal score
        mask = mask.fillna(False).astype(bool)
        score = score + mask * weight
        for pos in [out.index.get_loc(i) for i in out.index[mask]]:
            flags[pos].append(key)

    if m.get("customer_id"):
        cust = out[m["customer_id"]].astype(str)
        dup_keys = [cust, amount.round(2).astype(str)]
        if txn_date is not None:
            dup_keys.append(txn_date.dt.strftime("%Y-%m-%d"))
        key = pd.concat(dup_keys, axis=1).agg("|".join, axis=1)
        mark(key.duplicated(keep=False), "b_duplicate", 30)
        when = claim_date if claim_date is not None else txn_date
        if when is not None:  # 3 o más reclamaciones del mismo cliente en 90 días
            frame = pd.DataFrame({"c": cust, "d": when}).dropna()
            near = pd.Series(0, index=out.index)
            for _, g in frame.groupby("c"):
                if len(g) < 3:
                    continue
                days = g["d"].to_numpy(dtype="datetime64[D]").astype("int64")
                for idx, day in zip(g.index, days):
                    near[idx] = int((np.abs(days - day) <= 90).sum())
            mark(near >= 3, "b_repeat", 25)
        else:
            mark(cust.map(cust.value_counts()) >= 3, "b_repeat", 25)
    if claim_date is not None and txn_date is not None:
        delay = (claim_date - txn_date).dt.days
        out["_delay_days"] = delay
        mark(delay > sla["late_days"], "b_late", 15)
        mark(delay < 0, "b_date_error", 10)
    threshold = max(sla["high_amount"], float(amount.quantile(0.95)) if len(amount) else 0)
    mark(amount >= threshold, "b_high_amount", 15)
    mark((amount >= 500) & (amount % 100 == 0), "b_round", 5)
    cpp = pd.DataFrame()
    if m.get("merchant"):
        merchant = out[m["merchant"]].astype(str)
        stats = out.assign(_m=merchant, _a=amount)
        group = stats.groupby("_m")
        cpp = pd.DataFrame({
            "claims": group.size(),
            "customers": group[m["customer_id"]].nunique() if m.get("customer_id") else group.size(),
            "amount": group["_a"].sum().round(2),
        })
        if txn_date is not None:
            # Pico de reclamaciones en 30 días: muchos clientes distintos en poco tiempo en el mismo comercio
            # sugiere un punto de compromiso (skimming, filtración de datos…).
            stats["_t"] = txn_date
            peaks = {}
            for name, g in stats.dropna(subset=["_t"]).groupby("_m"):
                g = g.sort_values("_t")
                days = g["_t"].to_numpy(dtype="datetime64[D]").astype("int64")
                ends = np.searchsorted(days, days + 30, side="left")
                counts = ends - np.arange(len(days))
                best = int(counts.argmax())
                window = g.iloc[best:ends[best]]
                people = window[m["customer_id"]].nunique() if m.get("customer_id") else len(window)
                peaks[name] = (int(counts[best]), int(people), g["_t"].iloc[best].strftime("%Y-%m-%d"))
            cpp["peak_30d"] = [peaks.get(x, (0, 0, ""))[0] for x in cpp.index]
            cpp["peak_start"] = [peaks.get(x, (0, 0, ""))[2] for x in cpp.index]
            people = pd.Series([peaks.get(x, (0, 0, ""))[1] for x in cpp.index], index=cpp.index)
            baseline = float(cpp["peak_30d"].median()) if len(cpp) else 0
            cpp["suspect"] = (cpp["peak_30d"] >= max(cluster_min, 2 * baseline)) & (people >= 3)
            cpp = cpp.sort_values("peak_30d", ascending=False)
        else:
            limit = max(cluster_min, cpp["claims"].mean() + 2 * cpp["claims"].std(ddof=0)) if len(cpp) else 0
            cpp["suspect"] = (cpp["claims"] >= limit) & (cpp["customers"] >= max(3, cluster_min - 1))
            cpp = cpp.sort_values("claims", ascending=False)
        suspects = set(cpp.index[cpp["suspect"]])
        mark(merchant.isin(suspects), "b_cpp", 10)
    out["_score"] = score.clip(upper=100).astype(int)
    out["_flags"] = [", ".join(f) for f in flags]
    out = out.sort_values("_score", ascending=False)
    summary = {
        "claims": len(out), "amount": float(amount.sum()), "flagged": int((out["_score"] >= 25).sum()),
        "duplicates": sum("b_duplicate" in f for f in flags), "repeat": sum("b_repeat" in f for f in flags),
        "late": sum("b_late" in f for f in flags), "high": sum("b_high_amount" in f for f in flags),
        "cpp_merchants": [str(x) for x in cpp.index[cpp["suspect"]]] if len(cpp) else [],
    }
    by_channel = pd.DataFrame()
    if m.get("channel"):
        by_channel = out.assign(_a=amount).groupby(out[m["channel"]].astype(str)).agg(
            claims=("_score", "size"), amount=("_a", "sum")).sort_values("claims", ascending=False)
    by_month = pd.DataFrame()
    if txn_date is not None:
        by_month = out.assign(_a=amount, _mo=txn_date.dt.strftime("%Y-%m")).groupby("_mo").agg(
            claims=("_score", "size"), amount=("_a", "sum"))
    return out, {"summary": summary, "cpp": cpp, "by_channel": by_channel, "by_month": by_month}


def batch_context(result, mapping):
    """Estadísticas agregadas (sin datos personales) para que la IA redacte el informe."""
    s = result["summary"]
    parts = [f"Claims: {s['claims']}, total amount: {s['amount']:.2f}, flagged for review: {s['flagged']}",
             f"Duplicates: {s['duplicates']}, repeat claimants (3+ claims): {s['repeat']}, late reports: {s['late']}, "
             f"high amounts: {s['high']}",
             f"Merchants with unusual concentration of claims: {', '.join(s['cpp_merchants']) or 'none'}"]
    if len(result["cpp"]):
        parts.append("Top merchants by claims:\n" + result["cpp"].head(10).drop(columns=["suspect"]).to_string())
    if len(result["by_channel"]):
        parts.append("By channel:\n" + result["by_channel"].round(2).to_string())
    if len(result["by_month"]):
        parts.append("By month:\n" + result["by_month"].round(2).to_string())
    return "\n\n".join(parts)


def export_batch(path, detail, result, labels, flag_names):
    wb = Workbook()
    ws = wb.active
    ws.title = labels["sheet_triage"][:31]
    view = detail.drop(columns=[c for c in detail.columns if c == "_delay_days"]).rename(
        columns={"_score": labels["score"], "_flags": labels["flags"]})
    view[labels["flags"]] = view[labels["flags"]].map(
        lambda v: ", ".join(flag_names.get(x.strip(), x.strip()) for x in v.split(",") if x.strip()))
    write_df(ws, view, accent="7B1E3A")
    col = list(view.columns).index(labels["score"]) + 1
    letter = ws.cell(row=1, column=col).column_letter
    rng = f"{letter}2:{letter}{len(view) + 1}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["50"],
                                                  fill=PatternFill("solid", fgColor="F8B4B4")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["25", "49"],
                                                  fill=PatternFill("solid", fgColor="FDE68A")))

    ws2 = wb.create_sheet(labels["sheet_summary"][:31])
    title_cell(ws2, "A1", labels["sheet_summary"])
    s = result["summary"]
    rows = [(labels["k_claims"], s["claims"]), (labels["k_amount"], round(s["amount"], 2)),
            (labels["k_flagged"], s["flagged"]), (labels["k_duplicates"], s["duplicates"]),
            (labels["k_repeat"], s["repeat"]), (labels["k_late"], s["late"]), (labels["k_high"], s["high"]),
            (labels["k_cpp"], ", ".join(s["cpp_merchants"]) or "—")]
    for i, (k, v) in enumerate(rows, start=3):
        ws2.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws2.cell(row=i, column=2, value=v)
    ws2.column_dimensions["A"].width = 42
    ws2.column_dimensions["B"].width = 40
    if len(result["by_channel"]):
        start = len(rows) + 5
        ch = result["by_channel"].reset_index()
        ch.columns = [labels["channel"], labels["k_claims"], labels["k_amount"]]
        write_df(ws2, ch, accent="7B1E3A", start_row=start, as_table=False)
        chart = BarChart()
        chart.title = labels["channel"]
        chart.height, chart.width = 8, 16
        chart.add_data(Reference(ws2, min_col=2, min_row=start, max_row=start + len(ch)), titles_from_data=True)
        chart.set_categories(Reference(ws2, min_col=1, min_row=start + 1, max_row=start + len(ch)))
        ws2.add_chart(chart, "E3")
    if len(result["cpp"]):
        ws3 = wb.create_sheet(labels["sheet_merchants"][:31])
        cpp = result["cpp"].reset_index().rename(columns={"_m": labels["merchant"]})
        cpp = cpp.rename(columns={"claims": labels["k_claims"], "customers": labels["customers"],
                                  "amount": labels["k_amount"], "peak_30d": labels["peak_30d"],
                                  "peak_start": labels["peak_start"], "suspect": labels["suspect"]})
        cpp[labels["suspect"]] = cpp[labels["suspect"]].map({True: "⚠", False: ""})
        write_df(ws3, cpp, accent="7B1E3A")
    wb.save(path)
    return path


def export_cases(path, rows, labels):
    wb = Workbook()
    ws = wb.active
    ws.title = labels["sheet_cases"][:31]
    df = pd.DataFrame(rows)
    write_df(ws, df, accent="7B1E3A")
    wb.save(path)
    return path
