"""SlidePilot: lógica sin interfaz (esquema ↔ presentación, temas y generación del .pptx)."""

import re
from datetime import date

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from ..ai import STR, STR_LIST, obj

LAYOUTS = ("bullets", "columns", "kpis", "quote", "timeline", "section", "chart", "table")
CHART_TYPES = ("column", "bar", "line", "pie")

THEMES = {
    "corporate": {"bg": "FFFFFF", "primary": "1F3A5F", "accent": "2E86DE", "text": "1F2937", "muted": "6B7280",
                  "panel": "F1F5F9", "on_primary": "FFFFFF",
                  "series": ["2E86DE", "1F3A5F", "60A5FA", "F59E0B", "10B981", "EF4444"]},
    "dark": {"bg": "0F172A", "primary": "020617", "accent": "38BDF8", "text": "E2E8F0", "muted": "94A3B8",
             "panel": "1E293B", "on_primary": "FFFFFF",
             "series": ["38BDF8", "A78BFA", "34D399", "FBBF24", "F87171", "F472B6"]},
    "minimal": {"bg": "FFFFFF", "primary": "111111", "accent": "E63946", "text": "222222", "muted": "6B6B6B",
                "panel": "F5F5F5", "on_primary": "FFFFFF",
                "series": ["E63946", "111111", "8D99AE", "457B9D", "F4A261", "2A9D8F"]},
    "green": {"bg": "FFFFFF", "primary": "14532D", "accent": "16A34A", "text": "1F2937", "muted": "6B7280",
              "panel": "ECFDF5", "on_primary": "FFFFFF",
              "series": ["16A34A", "14532D", "86EFAC", "0EA5E9", "F59E0B", "64748B"]},
    "energy": {"bg": "FFFFFF", "primary": "7C2D12", "accent": "EA580C", "text": "1F2937", "muted": "6B7280",
               "panel": "FFF7ED", "on_primary": "FFFFFF",
               "series": ["EA580C", "7C2D12", "FDBA74", "0284C7", "65A30D", "6B7280"]},
    "creative": {"bg": "FFFFFF", "primary": "4C1D95", "accent": "DB2777", "text": "1F2937", "muted": "6B7280",
                 "panel": "F5F3FF", "on_primary": "FFFFFF",
                 "series": ["7C3AED", "DB2777", "F59E0B", "06B6D4", "10B981", "4C1D95"]},
}
FONT = "Segoe UI"
CLOSING = {"es": "¡Gracias!", "en": "Thank you!", "fr": "Merci !"}
AGENDA = {"es": "Agenda", "en": "Agenda", "fr": "Sommaire"}
CONT = {"es": "(cont.)", "en": "(cont.)", "fr": "(suite)"}

# ---------------------------------------------------------------- esquema de texto ↔ presentación
#   # Título de la presentación
#   Subtítulo (primera línea normal tras el título)
#   ## Título de diapositiva               → viñetas
#   ## [section] / [kpis] / [quote] / [timeline] / [columns] / [chart:column] / [table] Título
#   - viñeta      (  - subviñeta con sangría)
#   ### Encabezado de columna               (solo en [columns])
#   - 98% | Satisfacción                    (en [kpis]; en [chart] «- Categoría | valor | valor»;
#   - * | Serie 1 | Serie 2                  la línea «* |» da los nombres de las series)
#   > nota del orador


def empty_slide(layout="bullets", title=""):
    return {"layout": layout, "title": title, "bullets": [], "left_title": "", "left_bullets": [], "right_title": "",
            "right_bullets": [], "kpis": [], "quote": "", "author": "", "steps": [],
            "chart": {"type": "column", "categories": [], "series": []}, "table": [], "notes": ""}


def _num(text):
    text = str(text).strip().replace(" ", "").replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", text) or re.fullmatch(r"-?\d+,\d+", text):
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def parse_outline(text):
    deck = {"title": "", "subtitle": "", "slides": []}
    slide = None
    column = None
    series_names = None
    for raw in text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# ") and not deck["title"] and slide is None:
            deck["title"] = stripped[2:].strip()
            continue
        if stripped.startswith("## "):
            head = stripped[3:].strip()
            layout, chart_type = "bullets", "column"
            m = re.match(r"^\[(\w+)(?::(\w+))?\]\s*(.*)$", head)
            if m and m.group(1).lower() in LAYOUTS:
                layout = m.group(1).lower()
                chart_type = (m.group(2) or "column").lower()
                head = m.group(3)
            slide = empty_slide(layout, head)
            if layout == "chart":
                slide["chart"]["type"] = chart_type if chart_type in CHART_TYPES else "column"
            deck["slides"].append(slide)
            column, series_names = None, None
            continue
        if slide is None:
            if not deck["subtitle"] and not stripped.startswith(("-", ">")):
                deck["subtitle"] = stripped
            continue
        if stripped.startswith(">"):
            slide["notes"] = (slide["notes"] + "\n" + stripped[1:].strip()).strip()
            continue
        if stripped.startswith("### ") and slide["layout"] == "columns":
            column = "left" if column is None else "right"
            slide[f"{column}_title"] = stripped[4:].strip()
            continue
        bullet = re.match(r"^(\s*)[-•*]\s+(.*)$", line)
        content = bullet.group(2).strip() if bullet else stripped
        level = 1 if bullet and len(bullet.group(1).replace("\t", "  ")) >= 2 else 0
        layout = slide["layout"]
        if layout == "kpis":
            value, _, label = content.partition("|")
            slide["kpis"].append({"value": value.strip(), "label": label.strip()})
        elif layout == "quote":
            if not slide["quote"]:
                slide["quote"] = content.strip('"“”«» ')
            else:
                slide["author"] = content.lstrip("—–- ").strip()
        elif layout == "timeline":
            slide["steps"].append(content)
        elif layout == "columns":
            slide[f"{column or 'left'}_bullets"].append(content)
        elif layout == "chart":
            parts = [p.strip() for p in content.split("|")]
            if parts[0] == "*":
                series_names = parts[1:]
                continue
            values = [_num(p) for p in parts[1:]]
            if len(parts) < 2 or any(v is None for v in values):
                slide["bullets"].append(content)
                continue
            chart = slide["chart"]
            if not chart["series"]:
                names = series_names or [slide["title"] or "Serie"]
                chart["series"] = [{"name": names[i] if i < len(names) else f"S{i + 1}", "values": []}
                                   for i in range(len(values))]
            chart["categories"].append(parts[0])
            for i, s in enumerate(chart["series"]):
                s["values"].append(values[i] if i < len(values) else 0)
        elif layout == "table":
            slide["table"].append([p.strip() for p in content.split("|")])
        else:
            slide["bullets"].append(("  " if level else "") + content)
    return deck


def _fmt_value(v):
    return f"{v:g}" if isinstance(v, (int, float)) else str(v)


def deck_to_outline(deck):
    lines = [f"# {deck.get('title', '').strip()}"]
    if deck.get("subtitle"):
        lines.append(deck["subtitle"].strip())
    for s in deck.get("slides", []):
        layout = s.get("layout", "bullets")
        tag = ""
        if layout == "chart":
            tag = f"[chart:{s.get('chart', {}).get('type', 'column')}] "
        elif layout != "bullets":
            tag = f"[{layout}] "
        lines += ["", f"## {tag}{s.get('title', '').strip()}"]
        if layout == "kpis":
            lines += [f"- {k.get('value', '')} | {k.get('label', '')}" for k in s.get("kpis", [])]
        elif layout == "quote":
            lines += [f"- {s.get('quote', '')}", f"- — {s.get('author', '')}" if s.get("author") else ""]
        elif layout == "timeline":
            lines += [f"- {x}" for x in s.get("steps", [])]
        elif layout == "columns":
            lines += [f"### {s.get('left_title', '') or '-'}"] + [f"- {x}" for x in s.get("left_bullets", [])]
            lines += [f"### {s.get('right_title', '') or '-'}"] + [f"- {x}" for x in s.get("right_bullets", [])]
        elif layout == "chart":
            chart = s.get("chart", {})
            series = chart.get("series", [])
            if len(series) > 1:
                lines.append("- * | " + " | ".join(x.get("name", "") for x in series))
            for i, cat in enumerate(chart.get("categories", [])):
                vals = [_fmt_value(x["values"][i]) if i < len(x.get("values", [])) else "0" for x in series]
                lines.append(f"- {cat} | " + " | ".join(vals))
        elif layout == "table":
            lines += ["- " + " | ".join(str(c) for c in row) for row in s.get("table", [])]
        else:
            for b in s.get("bullets", []):
                lines.append(f"  - {b.strip()}" if b.startswith("  ") else f"- {b.strip()}")
        for note in (s.get("notes") or "").splitlines():
            if note.strip():
                lines.append(f"> {note.strip()}")
    return "\n".join(l for l in lines if l is not None)


def markup_to_outline(markup, lang="es", max_bullets=6):
    """Convierte un documento (marcado de DocPilot / texto de un .docx) en esquema de diapositivas."""
    title, subtitle, slides, current = "", "", [], None
    for raw in markup.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# ") and not title:
            title = line[2:].strip()
            continue
        if line.startswith(("## ", "### ")):
            current = empty_slide("bullets", line.lstrip("#").strip())
            slides.append(current)
            continue
        if current is None:
            if not subtitle:
                subtitle = line.lstrip("> ").strip()
            continue
        text = re.sub(r"^[-•*\d.)\s]+", "", line).strip() if re.match(r"^([-•*]|\d+[.)])\s", line) else line
        pieces = [text] if len(text) <= 160 else [p.strip() for p in re.split(r"(?<=[.!?])\s+", text) if p.strip()]
        for piece in pieces:
            if len(current["bullets"]) >= max_bullets:
                current = empty_slide("bullets", f"{slides[-1]['title'].replace(CONT[lang], '').strip()} {CONT[lang]}")
                slides.append(current)
            current["bullets"].append(piece if len(piece) <= 200 else piece[:197] + "…")
    deck = {"title": title or (slides[0]["title"] if slides else ""), "subtitle": subtitle, "slides": slides}
    return deck_to_outline(deck)


def excel_to_outline(df, category, values, chart_type="column", top=12, title="", lang="es"):
    """Agrupa por categoría (suma) y crea diapositivas de gráfico, KPI y tabla."""
    data = df.copy()
    for v in values:
        data[v] = pd.to_numeric(data[v], errors="coerce")
    grouped = data.groupby(category)[values].sum().sort_values(values[0], ascending=False).head(top)
    totals = data[values].sum()
    labels = {"es": ("Resumen", "Total", "Detalle", "por"), "en": ("Summary", "Total", "Detail", "by"),
              "fr": ("Synthèse", "Total", "Détail", "par")}[lang]
    deck = {"title": title or f"{', '.join(values)} {labels[3]} {category}", "subtitle": date.today().isoformat(),
            "slides": []}
    kpi = empty_slide("kpis", labels[0])
    kpi["kpis"] = [{"value": _short(totals[v]),
                    "label": v if str(v).lower().startswith(labels[1].lower()) else f"{labels[1]} {v}"}
                   for v in values[:3]]
    leader = grouped.index[0] if len(grouped) else ""
    if leader != "":
        kpi["kpis"].append({"value": str(leader)[:18], "label": f"#1 {category}"})
    deck["slides"].append(kpi)
    chart = empty_slide("chart", f"{', '.join(values)} {labels[3]} {category}")
    chart["chart"] = {"type": chart_type, "categories": [str(c) for c in grouped.index],
                      "series": [{"name": v, "values": [round(float(x), 2) for x in grouped[v]]} for v in values]}
    deck["slides"].append(chart)
    table = empty_slide("table", f"{labels[2]} {labels[3]} {category}")
    table["table"] = [[category] + values] + [[str(i)] + [_short(r[v]) for v in values] for i, r in grouped.iterrows()]
    deck["slides"].append(table)
    return deck_to_outline(deck)


def _short(value):
    value = float(value)
    for limit, suffix in ((1e9, "B"), (1e6, "M"), (1e3, "K")):
        if abs(value) >= limit:
            return f"{value / limit:.1f}{suffix}"
    return f"{value:,.0f}" if value.is_integer() else f"{value:,.2f}"


DECK_SCHEMA = obj({
    "title": STR,
    "subtitle": STR,
    "slides": {"type": "array", "items": obj({
        "layout": {"type": "string", "enum": list(LAYOUTS)},
        "title": STR,
        "bullets": STR_LIST,
        "left_title": STR, "left_bullets": STR_LIST, "right_title": STR, "right_bullets": STR_LIST,
        "kpis": {"type": "array", "items": obj({"value": STR, "label": STR})},
        "quote": STR, "author": STR,
        "steps": STR_LIST,
        "chart": obj({
            "type": {"type": "string", "enum": list(CHART_TYPES)},
            "categories": STR_LIST,
            "series": {"type": "array", "items": obj({"name": STR, "values": {"type": "array",
                                                                                "items": {"type": "number"}}})},
        }),
        "table": {"type": "array", "items": STR_LIST},
        "notes": STR,
    })},
})
DECK_SYSTEM = """You are SlidePilot, an expert presentation designer and business storyteller. Build a clear,
persuasive deck for the given topic, audience and goal: one idea per slide, action titles (the takeaway, not just a
label), 3-5 short bullets (max ~12 words each), and variety of layouts:
bullets, columns (comparison/before-after), kpis (2-4 big numbers), quote, timeline (3-6 steps), section (divider),
chart (ONLY with numbers given in the source material; never invent figures), table (only with given data).
Do not add a cover or a closing slide (they are added automatically). Fill unused fields with empty strings/arrays
and chart type "column". `notes`: 2-4 sentences the speaker can say. Use only facts from the user's material; when
data is missing, keep statements qualitative."""


# ---------------------------------------------------------------- generación del .pptx

W, H = Inches(13.333), Inches(7.5)


def _rgb(hexcolor):
    return RGBColor.from_string(hexcolor)


def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _text(slide, x, y, w, h, text, size=18, color="1F2937", bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          italic=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    f = run.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, FONT
    f.color.rgb = _rgb(color)
    return box


def _bullet_char(paragraph, color, level):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.32 + 0.35 * level))))
    pPr.set("indent", str(-Emu(Inches(0.3))))
    for tag in ("a:buNone", "a:buChar", "a:buClr"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = pPr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": color})
    clr.append(srgb)
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "–"}))


def _bullets(slide, x, y, w, h, items, theme, size=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    total = sum(len(i) for i in items)
    size = size or (28 if len(items) <= 3 and total < 160 else 24 if len(items) <= 5 and total < 360 else 20
                    if len(items) <= 7 and total < 600 else 16)
    for i, item in enumerate(items):
        level = 1 if item.startswith("  ") else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        _bullet_char(p, theme["accent"], level)
        _rich(p, item.strip(), size - 2 * level, theme["text"])
    return box


def _rich(paragraph, text, size, color):
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text):
        if m.start() > pos:
            _run(paragraph, text[pos:m.start()], size, color)
        _run(paragraph, m.group(1), size, color, bold=True)
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], size, color)


def _run(paragraph, text, size, color, bold=False):
    run = paragraph.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
    run.font.color.rgb = _rgb(color)


def _title(slide, text, theme):
    _rect(slide, Inches(0.6), Inches(0.55), Inches(0.12), Inches(0.75), theme["accent"])
    _text(slide, Inches(0.9), Inches(0.42), Inches(11.8), Inches(1.0), text, size=30 if len(text) < 60 else 24,
          color=theme["text"], bold=True, anchor=MSO_ANCHOR.MIDDLE)


def _footer(slide, theme, deck_title, number):
    _text(slide, Inches(0.6), Inches(6.95), Inches(9), Inches(0.35), deck_title, size=10, color=theme["muted"])
    _text(slide, Inches(11.9), Inches(6.95), Inches(0.9), Inches(0.35), str(number), size=10, color=theme["muted"],
          align=PP_ALIGN.RIGHT)


def _new(prs, theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(theme["bg"])
    return slide


def _cover(prs, theme, title, subtitle, author):
    slide = _new(prs, theme)
    _rect(slide, 0, 0, W, H, theme["primary"])
    _rect(slide, 0, Inches(5.9), W, Inches(0.12), theme["accent"])
    _rect(slide, Inches(0.9), Inches(2.0), Inches(1.4), Inches(0.12), theme["accent"])
    _text(slide, Inches(0.85), Inches(2.25), Inches(11.5), Inches(1.9), title, size=44 if len(title) < 50 else 34,
          color=theme["on_primary"], bold=True, anchor=MSO_ANCHOR.TOP)
    if subtitle:
        _text(slide, Inches(0.9), Inches(4.25), Inches(11), Inches(0.9), subtitle, size=20, color="CBD5E1")
    _text(slide, Inches(0.9), Inches(6.25), Inches(11), Inches(0.5), author, size=14, color="CBD5E1")
    return slide


def _closing(prs, theme, text, contact):
    slide = _new(prs, theme)
    _rect(slide, 0, 0, W, H, theme["primary"])
    _text(slide, 0, Inches(2.6), W, Inches(1.4), text, size=54, color=theme["on_primary"], bold=True,
          align=PP_ALIGN.CENTER)
    if contact:
        _text(slide, 0, Inches(4.1), W, Inches(0.6), contact, size=18, color="CBD5E1", align=PP_ALIGN.CENTER)
    _rect(slide, Inches(6.07), Inches(4.0), Inches(1.2), Inches(0.08), theme["accent"])
    return slide


def _render_slide(prs, theme, s, number, deck_title):
    layout = s.get("layout", "bullets")
    slide = _new(prs, theme)
    if layout == "section":
        _rect(slide, 0, 0, Inches(4.2), H, theme["accent"])
        _text(slide, Inches(0.5), Inches(2.6), Inches(3.2), Inches(2), f"{number:02d}", size=80, color="FFFFFF",
              bold=True)
        _text(slide, Inches(4.8), Inches(2.7), Inches(8), Inches(2), s.get("title", ""), size=40, color=theme["text"],
              bold=True, anchor=MSO_ANCHOR.MIDDLE)
        if s.get("bullets"):
            _text(slide, Inches(4.8), Inches(4.6), Inches(8), Inches(1.5), " · ".join(b.strip() for b in s["bullets"]),
                  size=18, color=theme["muted"])
    else:
        _title(slide, s.get("title", ""), theme)
        top, height = Inches(1.65), Inches(5.1)
        if layout == "columns":
            for i, side in enumerate(("left", "right")):
                x = Inches(0.6 + i * 6.2)
                _rect(slide, x, top, Inches(5.9), height, theme["panel"], MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.04
                _text(slide, x + Inches(0.3), top + Inches(0.2), Inches(5.3), Inches(0.7), s.get(f"{side}_title", ""),
                      size=22, color=theme["accent"], bold=True)
                _bullets(slide, x + Inches(0.3), top + Inches(0.95), Inches(5.3), height - Inches(1.1),
                         s.get(f"{side}_bullets", []), theme, size=20)
        elif layout == "kpis":
            kpis = s.get("kpis", [])[:4] or [{"value": "—", "label": ""}]
            n = len(kpis)
            card_w = (12.1 - 0.3 * (n - 1)) / n
            for i, k in enumerate(kpis):
                x = Inches(0.6 + i * (card_w + 0.3))
                card = _rect(slide, x, Inches(2.3), Inches(card_w), Inches(3.0), theme["panel"],
                             MSO_SHAPE.ROUNDED_RECTANGLE)
                card.adjustments[0] = 0.06
                _rect(slide, x, Inches(2.3), Inches(card_w), Inches(0.1), theme["accent"])
                _text(slide, x, Inches(2.75), Inches(card_w), Inches(1.4), k.get("value", ""),
                      size=54 if len(k.get("value", "")) < 7 else 36, color=theme["accent"], bold=True,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                _text(slide, x + Inches(0.2), Inches(4.15), Inches(card_w - 0.4), Inches(1.0), k.get("label", ""),
                      size=16, color=theme["text"], align=PP_ALIGN.CENTER)
            if s.get("bullets"):
                _text(slide, Inches(0.6), Inches(5.65), Inches(12), Inches(1.0), " · ".join(s["bullets"]), size=14,
                      color=theme["muted"], align=PP_ALIGN.CENTER)
        elif layout == "quote":
            _text(slide, Inches(1.0), Inches(1.7), Inches(1.5), Inches(1.5), "“", size=120, color=theme["accent"],
                  bold=True)
            _text(slide, Inches(1.8), Inches(2.3), Inches(10), Inches(2.6), s.get("quote", ""), size=30,
                  color=theme["text"], italic=True, anchor=MSO_ANCHOR.MIDDLE)
            if s.get("author"):
                _text(slide, Inches(1.8), Inches(5.0), Inches(10), Inches(0.6), "— " + s["author"], size=18,
                      color=theme["muted"])
        elif layout == "timeline":
            steps = s.get("steps", [])[:6] or ["—"]
            n = len(steps)
            gap = 12.1 / n
            y = Inches(3.3)
            _rect(slide, Inches(0.9), y + Inches(0.27), Inches(11.5), Inches(0.06), theme["muted"])
            for i, step in enumerate(steps):
                cx = Inches(0.6 + gap * i + gap / 2)
                dot = _rect(slide, cx - Inches(0.3), y, Inches(0.6), Inches(0.6), theme["accent"], MSO_SHAPE.OVAL)
                dot.text_frame.text = str(i + 1)
                para = dot.text_frame.paragraphs[0]
                para.alignment = PP_ALIGN.CENTER
                para.runs[0].font.size, para.runs[0].font.bold = Pt(16), True
                para.runs[0].font.color.rgb = _rgb("FFFFFF")
                head, _, rest = step.partition(":")
                label = head if rest else step
                _text(slide, cx - Inches(gap / 2 - 0.1), y + Inches(0.8), Inches(gap - 0.2), Inches(0.6), label.strip(),
                      size=18, color=theme["text"], bold=True, align=PP_ALIGN.CENTER)
                if rest:
                    _text(slide, cx - Inches(gap / 2 - 0.1), y + Inches(1.4), Inches(gap - 0.2), Inches(1.6),
                          rest.strip(), size=15, color=theme["muted"], align=PP_ALIGN.CENTER)
        elif layout == "chart" and s.get("chart", {}).get("categories"):
            _chart(slide, s["chart"], theme, s.get("bullets"))
        elif layout == "table" and s.get("table"):
            _table(slide, s["table"], theme)
        else:
            _bullets(slide, Inches(0.9), top, Inches(11.6), height, s.get("bullets", []), theme)
    _footer(slide, theme, deck_title, number)
    if s.get("notes"):
        slide.notes_slide.notes_text_frame.text = s["notes"]
    return slide


def _chart(slide, chart, theme, bullets=None):
    data = CategoryChartData()
    data.categories = chart["categories"]
    for series in chart["series"]:
        data.add_series(series["name"], [float(v or 0) for v in series["values"]])
    kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}[chart.get("type", "column")]
    width = Inches(8.4) if bullets else Inches(12.1)
    gf = slide.shapes.add_chart(kind, Inches(0.6), Inches(1.6), width, Inches(5.2), data)
    ch = gf.chart
    ch.font.size, ch.font.name = Pt(12), FONT
    ch.font.color.rgb = _rgb(theme["text"])
    multi = len(chart["series"]) > 1 or chart.get("type") == "pie"
    ch.has_legend = multi
    if multi:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    plot = ch.plots[0]
    if chart.get("type") == "pie":
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(theme["series"][i % len(theme["series"])])
    else:
        for i, series in enumerate(plot.series):
            color = _rgb(theme["series"][i % len(theme["series"])])
            if chart.get("type") == "line":
                series.format.line.color.rgb = color
                series.format.line.width = Pt(3)
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color
        if chart.get("type") in ("column", "bar"):
            plot.gap_width = 60
            ch.value_axis.has_major_gridlines = True
            ch.value_axis.major_gridlines.format.line.color.rgb = _rgb(theme["panel"])
    if len(chart["categories"]) <= 12 and chart.get("type") != "line":
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.font.size = Pt(11)
        labels.number_format = "#,##0" if chart.get("type") != "pie" else "0%"
        labels.number_format_is_linked = False
        if chart.get("type") == "pie":
            labels.show_percentage, labels.show_value = True, False
    if bullets:
        _bullets(slide, Inches(9.3), Inches(1.9), Inches(3.5), Inches(4.8), bullets, theme, size=16)


def _table(slide, rows, theme):
    n_rows, n_cols = min(len(rows), 14), max(len(r) for r in rows)
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.6), Inches(1.65), Inches(12.1),
                                   Inches(0.5 * n_rows))
    table = shape.table
    for r in range(n_rows):
        for c in range(n_cols):
            cell = table.cell(r, c)
            value = rows[r][c] if c < len(rows[r]) else ""
            cell.text = str(value)
            para = cell.text_frame.paragraphs[0]
            if para.runs:
                para.runs[0].font.size = Pt(16 if r == 0 else 15)
                para.runs[0].font.bold = r == 0
                para.runs[0].font.name = FONT
                para.runs[0].font.color.rgb = _rgb("FFFFFF" if r == 0 else theme["text"])
            if c > 0 and r > 0:
                para.alignment = PP_ALIGN.RIGHT
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(theme["primary"] if r == 0 else (theme["panel"] if r % 2 else theme["bg"]))


def build_pptx(deck, path, theme_key="corporate", lang="es", author="", agenda=True, closing=True, contact=""):
    theme = THEMES.get(theme_key, THEMES["corporate"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    title = deck.get("title") or "SlidePilot"
    subtitle = deck.get("subtitle", "")
    today = date.today().isoformat()
    _cover(prs, theme, title, subtitle, author or ("" if today in subtitle else today))
    slides = deck.get("slides", [])
    number = 2
    if agenda:
        topics = [s.get("title", "") for s in slides if s.get("layout") == "section"]
        if len(topics) < 3:
            topics = [s.get("title", "") for s in slides][:8]
        if len(topics) >= 3:
            agenda_slide = {"layout": "bullets", "title": AGENDA.get(lang, "Agenda"), "bullets": topics[:8]}
            _render_slide(prs, theme, agenda_slide, number, title)
            number += 1
    for s in slides:
        _render_slide(prs, theme, s, number, title)
        number += 1
    if closing:
        _closing(prs, theme, CLOSING.get(lang, "Thank you!"), contact)
    prs.save(path)
    return path
