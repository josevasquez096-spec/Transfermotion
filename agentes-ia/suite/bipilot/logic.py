"""BIPilot: lógica sin interfaz (modelo, relaciones, DAX, calendario, Power Query y temas de Power BI)."""

import colorsys
import json
import re
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from ..ai import STR, obj
from ..tabular import is_text, read_all_sheets, title_cell, write_df

# ---------------------------------------------------------------- modelo


def load_tables(paths):
    tables = {}
    for path in paths:
        for name, df in read_all_sheets(path).items():
            base, n = name, 2
            while name in tables:
                name = f"{base} ({n})"
                n += 1
            tables[name] = df
    return tables


def _norm(name):
    return re.sub(r"[\s_\-.]", "", str(name)).lower()


def is_key(series):
    s = series.dropna()
    return len(s) > 0 and s.is_unique and len(s) == len(series)


def is_id_name(name):
    n = _norm(name)
    return n == "id" or n.endswith("id") or n.endswith("key") or n.endswith("code") or n.endswith("codigo") \
        or n.startswith("id") or n.endswith("clave") or n.endswith("cod")


def table_role(df):
    numeric = [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c]) and not is_id_name(c)]
    dates = [c for c in df.columns if pd.api.types.is_datetime64_any_dtype(df[c])]
    ids = [c for c in df.columns if is_id_name(c)]
    keys = [c for c in ids if is_key(df[c])]
    if (dates or len(ids) - len(keys) >= 1) and numeric:
        return "fact"
    if keys:
        return "dimension"
    return "other"


def _compatible(a, b):
    num_a, num_b = pd.api.types.is_numeric_dtype(a), pd.api.types.is_numeric_dtype(b)
    return num_a == num_b


def suggest_relationships(tables):
    """Relaciones 1:* entre una columna clave (lado 1) y una columna con los mismos valores (lado *)."""
    rels = []
    names = list(tables)
    for one_name in names:
        one = tables[one_name]
        for one_col in one.columns:
            if not is_key(one[one_col]):
                continue
            keys = set(one[one_col].astype(str))
            for many_name in names:
                if many_name == one_name:
                    continue
                many = tables[many_name]
                for many_col in many.columns:
                    if not _compatible(one[one_col], many[many_col]):
                        continue
                    same_name = _norm(one_col) == _norm(many_col) or (
                        is_id_name(one_col) and _norm(many_col).endswith(_norm(one_col)))
                    values = many[many_col].dropna().astype(str)
                    if values.empty:
                        continue
                    coverage = values.isin(keys).mean()
                    by_values = (not pd.api.types.is_numeric_dtype(many[many_col]) and coverage >= 0.95
                                 and values.nunique() >= 3)
                    if (same_name and coverage >= 0.5) or by_values:
                        rels.append({
                            "one_table": one_name, "one_col": str(one_col),
                            "many_table": many_name, "many_col": str(many_col),
                            "cardinality": "1:1" if many[many_col].is_unique else "1:*",
                            "coverage": round(float(coverage), 4),
                            "orphans": int((~values.isin(keys)).sum()),
                        })
    # Si hay 1:1 en ambos sentidos, se queda uno.
    seen, unique = set(), []
    for r in rels:
        pair = frozenset([(r["one_table"], r["one_col"]), (r["many_table"], r["many_col"])])
        if pair not in seen:
            seen.add(pair)
            unique.append(r)
    return unique


def quality_issues(tables, rels):
    """Lista de (gravedad, clave_de_mensaje, parámetros)."""
    issues = []
    for name, df in tables.items():
        for col in df.columns:
            s = df[col]
            if is_id_name(col) and table_role(df) == "dimension":
                dupes = int(s.dropna().duplicated().sum())
                if dupes:
                    issues.append(("high", "q_dup_keys", {"table": name, "col": col, "n": dupes}))
                nulls = int(s.isna().sum())
                if nulls:
                    issues.append(("high", "q_null_keys", {"table": name, "col": col, "n": nulls}))
            if is_text(s):
                values = s.dropna()
                numeric_like = pd.to_numeric(values, errors="coerce").notna().mean() if len(values) else 0
                if 0.5 <= numeric_like < 1:
                    issues.append(("medium", "q_mixed", {"table": name, "col": col}))
                elif len(values) and values.astype(str).str.match(r"^\d{1,4}[-/]\d{1,2}[-/]\d{1,4}$").mean() > 0.9:
                    issues.append(("medium", "q_text_dates", {"table": name, "col": col}))
        empty = [c for c in df.columns if df[c].isna().all()]
        for col in empty:
            issues.append(("low", "q_empty_col", {"table": name, "col": col}))
    for r in rels:
        if r["orphans"]:
            issues.append(("high", "q_orphans", {"table": r["many_table"], "col": r["many_col"], "n": r["orphans"],
                                                 "dim": r["one_table"]}))
    if not any(table_role(df) == "fact" for df in tables.values()):
        issues.append(("low", "q_no_fact", {}))
    return issues


def date_columns(tables):
    out = []
    for name, df in tables.items():
        for col in df.columns:
            if pd.api.types.is_datetime64_any_dtype(df[col]):
                out.append((name, col))
    return out


def date_range(tables):
    lo, hi = None, None
    for name, col in date_columns(tables):
        s = tables[name][col].dropna()
        if s.empty:
            continue
        lo = s.min() if lo is None else min(lo, s.min())
        hi = s.max() if hi is None else max(hi, s.max())
    if lo is None:
        today = date.today()
        return date(today.year - 1, 1, 1), date(today.year, 12, 31)
    return date(lo.year, 1, 1), date(hi.year, 12, 31)


# ---------------------------------------------------------------- DAX

MEASURE_NAMES = {
    "es": {"total": "Total {c}", "avg": "Promedio {c}", "count": "Nº filas {t}", "dcount": "Nº {c} distintos",
           "ytd": "{c} acumulado año", "py": "{c} año anterior", "yoy": "{c} var. % interanual",
           "pm": "{c} mes anterior", "mom": "{c} var. % mensual", "running": "{c} acumulado",
           "pct": "{c} % del total", "ma3": "{c} media móvil 3M", "rank": "Ranking {r} por {c}"},
    "en": {"total": "Total {c}", "avg": "Average {c}", "count": "Row count {t}", "dcount": "Distinct {c}",
           "ytd": "{c} YTD", "py": "{c} PY", "yoy": "{c} YoY %", "pm": "{c} PM", "mom": "{c} MoM %",
           "running": "{c} running total", "pct": "{c} % of total", "ma3": "{c} 3M moving avg",
           "rank": "{r} rank by {c}"},
    "fr": {"total": "Total {c}", "avg": "Moyenne {c}", "count": "Nb lignes {t}", "dcount": "Nb {c} distincts",
           "ytd": "{c} cumul annuel", "py": "{c} année précédente", "yoy": "{c} var. % annuelle",
           "pm": "{c} mois précédent", "mom": "{c} var. % mensuelle", "running": "{c} cumulé",
           "pct": "{c} % du total", "ma3": "{c} moyenne mobile 3M", "rank": "Rang {r} par {c}"},
}
COMMENTS = {
    "es": {"base": "Medidas base", "time": "Inteligencia de tiempo (requiere una tabla calendario marcada como "
                                           "tabla de fechas y relacionada con {t})", "other": "Análisis"},
    "en": {"base": "Base measures", "time": "Time intelligence (needs a calendar table marked as date table and "
                                           "related to {t})", "other": "Analysis"},
    "fr": {"base": "Mesures de base", "time": "Intelligence temporelle (nécessite une table calendrier marquée comme "
                                             "table de dates et liée à {t})", "other": "Analyse"},
}


def q(name):
    """Nombre de tabla entre comillas simples para DAX."""
    return "'" + str(name).replace("'", "''") + "'"


def col_ref(table, col):
    return f"{q(table)}[{str(col).replace(']', ']]')}]"


def measure_ref(name):
    return f"[{name.replace(']', ']]')}]"


def dax_measures(table, value_col, lang="es", date_ref=None, rank_table=None, rank_col=None, agg="SUM"):
    names = {k: v.format(c=value_col, t=table, r=rank_col or "") for k, v in MEASURE_NAMES[lang].items()}
    total = measure_ref(names["total"])
    lines = [f"-- {COMMENTS[lang]['base']}",
             f"{names['total']} = {agg} ( {col_ref(table, value_col)} )",
             f"{names['avg']} = AVERAGE ( {col_ref(table, value_col)} )",
             f"{names['count']} = COUNTROWS ( {q(table)} )", ""]
    if date_ref:
        lines += [
            f"-- {COMMENTS[lang]['time'].format(t=table)}",
            f"{names['ytd']} = TOTALYTD ( {total}, {date_ref} )",
            f"{names['py']} = CALCULATE ( {total}, SAMEPERIODLASTYEAR ( {date_ref} ) )",
            f"{names['yoy']} = DIVIDE ( {total} - {measure_ref(names['py'])}, {measure_ref(names['py'])} )",
            f"{names['pm']} = CALCULATE ( {total}, DATEADD ( {date_ref}, -1, MONTH ) )",
            f"{names['mom']} = DIVIDE ( {total} - {measure_ref(names['pm'])}, {measure_ref(names['pm'])} )",
            f"{names['running']} =\nCALCULATE (\n    {total},\n    FILTER ( ALL ( {date_ref} ), {date_ref} <= MAX ( "
            f"{date_ref} ) )\n)",
            f"{names['ma3']} =\nDIVIDE (\n    CALCULATE ( {total}, DATESINPERIOD ( {date_ref}, MAX ( {date_ref} ), -3, "
            f"MONTH ) ),\n    3\n)",
            "",
        ]
    lines += [f"-- {COMMENTS[lang]['other']}",
              f"{names['pct']} = DIVIDE ( {total}, CALCULATE ( {total}, ALLSELECTED () ) )"]
    if rank_table and rank_col:
        lines.append(f"{names['rank']} = RANKX ( ALLSELECTED ( {col_ref(rank_table, rank_col)} ), {total}, , DESC, "
                     f"DENSE )")
    return "\n".join(lines) + "\n"


def model_description(tables, rels):
    """Descripción compacta del modelo para la IA (sin filas de datos)."""
    lines = []
    for name, df in tables.items():
        cols = ", ".join(f"{c} ({_kind(df[c])})" for c in df.columns)
        lines.append(f"Table '{name}' [{table_role(df)}, {len(df)} rows]: {cols}")
    for r in rels:
        lines.append(f"Relationship: '{r['one_table']}'[{r['one_col']}] (1) -> '{r['many_table']}'[{r['many_col']}] (*)")
    return "\n".join(lines)


def _kind(series):
    if pd.api.types.is_bool_dtype(series):
        return "boolean"
    if pd.api.types.is_numeric_dtype(series):
        return "number"
    if pd.api.types.is_datetime64_any_dtype(series):
        return "date"
    return "text"


DAX_SCHEMA = obj({"measure_name": STR, "dax": STR, "format_string": STR, "explanation": STR, "notes": STR})
DAX_SYSTEM = """You are BIPilot, a Power BI and DAX expert. Write ONE DAX measure for the request using only the
tables and columns of the model (quote table names with single quotes, columns as 'Table'[Column]).
Prefer DIVIDE over "/", variables (VAR/RETURN) for readability and time-intelligence functions when a calendar
table exists. `dax` must be the full definition "Name = expression". `format_string`: a Power BI format string.
`notes`: prerequisites (e.g. mark the calendar as date table) or caveats; empty if none."""
EXPLAIN_DAX_SYSTEM = """You are BIPilot, a DAX expert. Explain the DAX code to a business user: what it calculates,
how filter context affects it, step by step, possible performance or correctness issues and an improved version if
useful. Plain text with short bullets."""
M_SCHEMA = obj({"m_code": STR, "explanation": STR})
M_SYSTEM = """You are BIPilot, a Power Query (M) expert. Write a complete M query (let ... in ...) that performs the
requested transformation on the described table. Use the source step given by the user when provided; otherwise
start from a step called Source that the user will replace. Use clear English step names without spaces."""

# ---------------------------------------------------------------- calendario

CAL = {
    "es": {"table": "Calendario", "date": "Fecha", "year": "Año", "quarter": "Trimestre", "q": "T",
           "month_n": "Mes Nº", "month": "Mes", "year_month": "Año-Mes", "weekday_n": "Día semana Nº",
           "weekday": "Día semana", "fy": "Año fiscal", "fy_prefix": "AF", "weekend": "Fin de semana",
           "culture": "es-ES",
           "months": ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
                      "octubre", "noviembre", "diciembre"],
           "days": ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]},
    "en": {"table": "Calendar", "date": "Date", "year": "Year", "quarter": "Quarter", "q": "Q",
           "month_n": "Month No", "month": "Month", "year_month": "Year-Month", "weekday_n": "Weekday No",
           "weekday": "Weekday", "fy": "Fiscal year", "fy_prefix": "FY", "weekend": "Weekend", "culture": "en-US",
           "months": ["January", "February", "March", "April", "May", "June", "July", "August", "September",
                      "October", "November", "December"],
           "days": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]},
    "fr": {"table": "Calendrier", "date": "Date", "year": "Année", "quarter": "Trimestre", "q": "T",
           "month_n": "Mois n°", "month": "Mois", "year_month": "Année-Mois", "weekday_n": "Jour semaine n°",
           "weekday": "Jour semaine", "fy": "Exercice", "fy_prefix": "EX", "weekend": "Week-end",
           "culture": "fr-FR",
           "months": ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre",
                      "octobre", "novembre", "décembre"],
           "days": ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]},
}


def _dax_str(text):
    return '"' + text.replace('"', '""') + '"'


def calendar_dax(start, end, lang="es", fiscal_start=1, table=None):
    c = CAL[lang]
    table = table or c["table"]
    months = ", ".join(f"{i}, {_dax_str(m)}" for i, m in enumerate(c["months"], start=1))
    days = ", ".join(f"{i}, {_dax_str(d)}" for i, d in enumerate(c["days"], start=1))
    fiscal = (f'{_dax_str(c["fy_prefix"] + " ")} & ( YEAR ( [Date] ) + IF ( MONTH ( [Date] ) >= {fiscal_start}, 1, 0 ) )'
              if fiscal_start > 1 else f'{_dax_str(c["fy_prefix"] + " ")} & YEAR ( [Date] )')
    return f"""{table} =
VAR _Start = DATE ( {start.year}, {start.month}, {start.day} )
VAR _End = DATE ( {end.year}, {end.month}, {end.day} )
RETURN
ADDCOLUMNS (
    CALENDAR ( _Start, _End ),
    {_dax_str(c["year"])}, YEAR ( [Date] ),
    {_dax_str(c["quarter"])}, {_dax_str(c["q"])} & QUARTER ( [Date] ),
    {_dax_str(c["month_n"])}, MONTH ( [Date] ),
    {_dax_str(c["month"])}, SWITCH ( MONTH ( [Date] ), {months} ),
    {_dax_str(c["year_month"])}, FORMAT ( [Date], "YYYY-MM" ),
    {_dax_str(c["weekday_n"])}, WEEKDAY ( [Date], 2 ),
    {_dax_str(c["weekday"])}, SWITCH ( WEEKDAY ( [Date], 2 ), {days} ),
    {_dax_str(c["fy"])}, {fiscal},
    {_dax_str(c["weekend"])}, WEEKDAY ( [Date], 2 ) > 5
)
"""


def _m_str(text):
    return '"' + text.replace('"', '""') + '"'


def calendar_m(start, end, lang="es", fiscal_start=1):
    c = CAL[lang]
    fy = (f'{_m_str(c["fy_prefix"] + " ")} & Text.From(Date.Year([{c["date"]}]) + '
          f'(if Date.Month([{c["date"]}]) >= {fiscal_start} then 1 else 0))' if fiscal_start > 1
          else f'{_m_str(c["fy_prefix"] + " ")} & Text.From(Date.Year([{c["date"]}]))')
    d = f'[{c["date"]}]'
    return f"""let
    StartDate = #date({start.year}, {start.month}, {start.day}),
    EndDate = #date({end.year}, {end.month}, {end.day}),
    Dates = List.Dates(StartDate, Duration.Days(EndDate - StartDate) + 1, #duration(1, 0, 0, 0)),
    AsTable = Table.FromList(Dates, Splitter.SplitByNothing(), {{{_m_str(c["date"])}}}),
    Typed = Table.TransformColumnTypes(AsTable, {{{{{_m_str(c["date"])}, type date}}}}),
    AddYear = Table.AddColumn(Typed, {_m_str(c["year"])}, each Date.Year({d}), Int64.Type),
    AddQuarter = Table.AddColumn(AddYear, {_m_str(c["quarter"])}, each {_m_str(c["q"])} & Text.From(Date.QuarterOfYear({d})), type text),
    AddMonthNo = Table.AddColumn(AddQuarter, {_m_str(c["month_n"])}, each Date.Month({d}), Int64.Type),
    AddMonth = Table.AddColumn(AddMonthNo, {_m_str(c["month"])}, each Date.MonthName({d}, {_m_str(c["culture"])}), type text),
    AddYearMonth = Table.AddColumn(AddMonth, {_m_str(c["year_month"])}, each Date.ToText({d}, "yyyy-MM"), type text),
    AddWeekdayNo = Table.AddColumn(AddYearMonth, {_m_str(c["weekday_n"])}, each Date.DayOfWeek({d}, Day.Monday) + 1, Int64.Type),
    AddWeekday = Table.AddColumn(AddWeekdayNo, {_m_str(c["weekday"])}, each Date.DayOfWeekName({d}, {_m_str(c["culture"])}), type text),
    AddFiscal = Table.AddColumn(AddWeekday, {_m_str(c["fy"])}, each {fy}, type text),
    AddWeekend = Table.AddColumn(AddFiscal, {_m_str(c["weekend"])}, each Date.DayOfWeek({d}, Day.Monday) >= 5, type logical)
in
    AddWeekend
"""


def calendar_df(start, end, lang="es", fiscal_start=1):
    c = CAL[lang]
    dates = pd.date_range(start, end, freq="D")
    fy = dates.year + ((dates.month >= fiscal_start) & (fiscal_start > 1)).astype(int)
    return pd.DataFrame({
        c["date"]: dates,
        c["year"]: dates.year,
        c["quarter"]: [f"{c['q']}{q}" for q in dates.quarter],
        c["month_n"]: dates.month,
        c["month"]: [c["months"][m - 1] for m in dates.month],
        c["year_month"]: dates.strftime("%Y-%m"),
        c["weekday_n"]: dates.dayofweek + 1,
        c["weekday"]: [c["days"][d] for d in dates.dayofweek],
        c["fy"]: [f"{c['fy_prefix']} {y}" for y in fy],
        c["weekend"]: dates.dayofweek >= 5,
    })


# ---------------------------------------------------------------- Power Query


def m_clean_query(df, table_name, path_hint=None):
    """Consulta M que lee la tabla, pone los tipos detectados, recorta textos y quita duplicados/vacías."""
    path = (path_hint or r"C:\\data\\file.xlsx").replace('"', '""')
    types = []
    trims = []
    for col in df.columns:
        s = df[col]
        name = _m_str(str(col))
        if pd.api.types.is_bool_dtype(s):
            types.append(f"{{{name}, type logical}}")
        elif pd.api.types.is_integer_dtype(s):
            types.append(f"{{{name}, Int64.Type}}")
        elif pd.api.types.is_numeric_dtype(s):
            types.append(f"{{{name}, type number}}")
        elif pd.api.types.is_datetime64_any_dtype(s):
            has_time = (s.dropna().dt.normalize() != s.dropna()).any()
            types.append(f"{{{name}, {'type datetime' if has_time else 'type date'}}}")
        else:
            types.append(f"{{{name}, type text}}")
            trims.append(f"{{{name}, Text.Trim, type text}}")
    if str(path).lower().endswith((".csv", ".txt")):
        source = (f'    Source = Csv.Document(File.Contents("{path}"), [Delimiter=",", Encoding=65001, '
                  f'QuoteStyle=QuoteStyle.Csv]),\n    Promoted = Table.PromoteHeaders(Source, [PromoteAllScalars=true]),')
    else:
        source = (f'    Source = Excel.Workbook(File.Contents("{path}"), null, true),\n'
                  f'    Sheet = Source{{[Item={_m_str(table_name)}, Kind="Sheet"]}}[Data],\n'
                  f'    Promoted = Table.PromoteHeaders(Sheet, [PromoteAllScalars=true]),')
    sep = ",\n        "
    return f"""let
{source}
    Typed = Table.TransformColumnTypes(Promoted, {{
        {sep.join(types)}
    }}),
    Trimmed = Table.TransformColumns(Typed, {{
        {sep.join(trims) if trims else ''}
    }}),
    NoEmptyRows = Table.SelectRows(Trimmed, each not List.IsEmpty(List.RemoveMatchingItems(Record.FieldValues(_), {{"", null}}))),
    NoDuplicates = Table.Distinct(NoEmptyRows)
in
    NoDuplicates
"""


# ---------------------------------------------------------------- temas


def hex_to_rgb(color):
    color = color.lstrip("#")
    return tuple(int(color[i:i + 2], 16) / 255 for i in (0, 2, 4))


def rgb_to_hex(rgb):
    return "#" + "".join(f"{max(0, min(255, round(v * 255))):02X}" for v in rgb)


def palette_from(base, n=8):
    """Paleta armónica: el color base, rotaciones de tono y variaciones de luminosidad."""
    h, l, s = colorsys.rgb_to_hls(*hex_to_rgb(base))
    steps = [0, 0.5, 0.15, 0.65, 0.33, 0.83, 0.08, 0.58, 0.25, 0.75]
    colors = []
    for i in range(n):
        hue = (h + steps[i % len(steps)]) % 1
        light = min(0.62, max(0.42, l + (0.1 if i % 2 else 0)))
        sat = min(0.85, max(0.5, s))
        colors.append(rgb_to_hex(colorsys.hls_to_rgb(hue, light, sat)))
    colors[0] = base.upper() if base.startswith("#") else "#" + base.upper()
    return colors


def theme_json(name, colors, background="#FFFFFF", foreground="#252423", font="Segoe UI"):
    theme = {
        "name": name,
        "dataColors": colors,
        "background": background,
        "foreground": foreground,
        "tableAccent": colors[0],
        "good": "#1AAB40",
        "neutral": "#D9B300",
        "bad": "#D64554",
        "maximum": colors[0],
        "center": "#D9B300",
        "minimum": colors[1] if len(colors) > 1 else colors[0],
        "textClasses": {
            "callout": {"fontSize": 28, "fontFace": font, "color": foreground},
            "title": {"fontSize": 14, "fontFace": font, "color": foreground},
            "header": {"fontSize": 12, "fontFace": font, "color": foreground},
            "label": {"fontSize": 10, "fontFace": font, "color": foreground},
        },
    }
    return json.dumps(theme, indent=2, ensure_ascii=False)


# ---------------------------------------------------------------- exportación


def export_model(path, tables, rels, issues_text, labels, calendar=None, dax=""):
    wb = Workbook()
    doc = wb.active
    doc.title = labels["doc"][:31]
    title_cell(doc, "A1", labels["doc_title"])
    rows = []
    for name, df in tables.items():
        for col in df.columns:
            rows.append({labels["table"]: name, labels["column"]: str(col), labels["type"]: _kind(df[col]),
                         labels["role"]: labels["role_" + table_role(df)]})
    write_df(doc, pd.DataFrame(rows), accent="B8860B", start_row=3, as_table=False)
    r = len(rows) + 6
    doc.cell(row=r, column=1, value=labels["relationships"]).font = Font(bold=True, size=13)
    rel_df = pd.DataFrame([{labels["one"]: f"{x['one_table']}[{x['one_col']}]",
                            labels["many"]: f"{x['many_table']}[{x['many_col']}]",
                            labels["cardinality"]: x["cardinality"], labels["coverage"]: x["coverage"]} for x in rels])
    if not rel_df.empty:
        write_df(doc, rel_df, accent="B8860B", start_row=r + 1, as_table=False, number_format="0.0%")
    r += len(rel_df) + 4
    doc.cell(row=r, column=1, value=labels["issues"]).font = Font(bold=True, size=13)
    for i, line in enumerate(issues_text.splitlines(), start=r + 1):
        doc.cell(row=i, column=1, value=line)
    if dax:
        ws = wb.create_sheet("DAX")
        ws.column_dimensions["A"].width = 120
        for i, line in enumerate(dax.splitlines(), start=1):
            ws.cell(row=i, column=1, value=line).font = Font(name="Consolas")
    for name, df in tables.items():
        ws = wb.create_sheet(_sheet_name(wb, name))
        write_df(ws, df, accent="B8860B")
    if calendar is not None:
        ws = wb.create_sheet(_sheet_name(wb, labels["calendar"]))
        write_df(ws, calendar, accent="B8860B")
    wb.save(path)
    return path


def _sheet_name(wb, name):
    base = re.sub(r"[\[\]:*?/\\]", "_", str(name))[:28]
    candidate, n = base, 2
    while candidate in wb.sheetnames:
        candidate = f"{base}_{n}"
        n += 1
    return candidate
