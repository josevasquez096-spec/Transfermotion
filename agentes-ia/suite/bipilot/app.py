"""BIPilot: agente de IA para Power BI."""

import tkinter as tk
from datetime import date
from pathlib import Path
from tkinter import colorchooser, ttk

import pandas as pd

from .. import samples
from ..tabular import read_all_sheets
from ..ui import DataTable, PilotApp, TextBox
from . import logic
from .strings import STRINGS

MONO = ("Consolas", 10)
FONTS = ("Segoe UI", "Segoe UI Semibold", "Arial", "Calibri", "Verdana", "Tahoma", "Georgia", "DIN")


class BIPilot(PilotApp):
    APP_NAME = "BIPilot"
    ICON_TEXT = "BI"
    ACCENT = "#B8860B"
    STRINGS = STRINGS

    def build(self):
        self.tables = {}
        self.sources = {}
        self.rels = []
        self.issues = []
        self.colors = logic.palette_from("#2B579A")
        self.background = "#FFFFFF"
        self.foreground = "#252423"
        self._build_model()
        self._build_dax()
        self._build_calendar()
        self._build_pq()
        self._build_theme()

    # ------------------------------------------------------------ modelo
    def _build_model(self):
        tab = self.add_tab("tab_model")
        bar = ttk.Frame(tab)
        bar.pack(fill="x")
        self.btn(bar, "add_tables", self.add_tables, accent=True).pack(side="left")
        self.btn(bar, "load_sample", self.load_sample).pack(side="left", padx=6)
        self.btn(bar, "clear_model", self.clear_model).pack(side="left")
        self.btn(bar, "export_model", self.export_model).pack(side="right")
        self.lbl(tab, "model_hint", hint=True).pack(anchor="w", pady=6)

        panes = ttk.PanedWindow(tab, orient="horizontal")
        panes.pack(fill="both", expand=True)
        left = self.lframe(panes, "tables", padding=6)
        self.tables_view = DataTable(left, None, height=10)
        self.tables_view.pack(fill="both", expand=True)
        right = ttk.Frame(panes)
        rel_box = self.lframe(right, "relationships", padding=6)
        rel_box.pack(fill="both", expand=True)
        self.rels_view = DataTable(rel_box, None, height=6)
        self.rels_view.pack(fill="both", expand=True)
        issue_box = self.lframe(right, "issues", padding=6)
        issue_box.pack(fill="both", expand=True, pady=(8, 0))
        self.issues_text = TextBox(issue_box, height=8, readonly=True)
        self.issues_text.pack(fill="both", expand=True)
        panes.add(left, weight=2)
        panes.add(right, weight=3)
        self.tr_call(self._refresh_model_views)

    def add_tables(self):
        paths = self.ask_open(["excel", "csv"], multiple=True)
        if not paths:
            return

        def work():
            loaded = {}
            for path in paths:
                for name, df in read_all_sheets(path).items():
                    loaded[name] = (df, path)
            return loaded

        def done(loaded):
            for name, (df, path) in loaded.items():
                base, n = name, 2
                while name in self.tables:
                    name = f"{base} ({n})"
                    n += 1
                self.tables[name] = df
                self.sources[name] = path
            self._analyze_model()

        self.run_task(work, done)

    def load_sample(self):
        self.tables = samples.bi_model(self.t.lang)
        self.sources = {name: "C:\\data\\sample.xlsx" for name in self.tables}
        self._analyze_model()

    def clear_model(self):
        self.tables, self.sources, self.rels, self.issues = {}, {}, [], []
        self._refresh_model_views()

    def _analyze_model(self):
        self.rels = logic.suggest_relationships(self.tables)
        self.issues = logic.quality_issues(self.tables, self.rels)
        self._refresh_model_views()
        self._refresh_choices()
        start, end = logic.date_range(self.tables)
        self._set_entry(self.cal_start, start.isoformat())
        self._set_entry(self.cal_end, end.isoformat())

    def _issues_text(self):
        if not self.tables:
            return ""
        if not self.issues:
            return self.t("no_issues")
        return "\n".join(f"[{self.t('sev_' + sev)}] {self.t(key, **params)}" for sev, key, params in self.issues)

    def _refresh_model_views(self):
        t = self.t
        rows = [{t("table"): name, t("rows"): len(df), t("columns"): len(df.columns),
                 t("role"): t("role_" + logic.table_role(df))} for name, df in self.tables.items()]
        self.tables_view.show(pd.DataFrame(rows) if rows else None)
        rels = [{t("one"): f"{r['one_table']}[{r['one_col']}]", t("many"): f"{r['many_table']}[{r['many_col']}]",
                 t("cardinality"): r["cardinality"], t("coverage"): f"{r['coverage']:.1%}"} for r in self.rels]
        self.rels_view.show(pd.DataFrame(rels) if rels else None)
        self.issues_text.set(self._issues_text())

    def export_model(self):
        if not self.need(self.tables, "need_model"):
            return
        path = self.ask_save("excel", "PowerBI_model.xlsx")
        if not path:
            return
        keys = ("doc", "doc_title", "table", "column", "type", "role", "role_fact", "role_dimension", "role_other",
                "relationships", "one", "many", "cardinality", "coverage", "issues", "calendar")
        labels = {k: self.t(k) for k in keys}
        start, end = logic.date_range(self.tables)
        calendar = logic.calendar_df(start, end, self.t.lang, self._fiscal())
        dax = self.dax_out.get()
        self.run_task(lambda: logic.export_model(path, self.tables, self.rels, self._issues_text(), labels, calendar,
                                                 dax), self.saved)

    # ------------------------------------------------------------ DAX
    def _build_dax(self):
        tab = self.add_tab("tab_dax")
        left = ttk.Frame(tab)
        left.pack(side="left", fill="y", padx=(0, 10))
        box = self.lframe(left, "templates_box")
        box.pack(fill="x")
        box.columnconfigure(1, weight=1)
        self.fact_combo = self.field(box, "measure_table", ttk.Combobox(box, state="readonly", width=30), 0)
        self.fact_combo.bind("<<ComboboxSelected>>", lambda e: self._refresh_value_cols())
        self.value_combo = self.field(box, "value_col", ttk.Combobox(box, state="readonly", width=30), 1)
        self.agg_combo = self.field(box, "agg_fn", ttk.Combobox(
            box, state="readonly", width=30, values=("SUM", "AVERAGE", "MIN", "MAX", "DISTINCTCOUNT")), 2)
        self.agg_combo.current(0)
        self.date_entry = self.field(box, "date_col", ttk.Entry(box, width=32), 3)
        self.lbl(box, "date_hint", hint=True).grid(row=4, column=0, columnspan=2, sticky="w")
        self.rank_combo = self.field(box, "rank_by", ttk.Combobox(box, state="readonly", width=30), 5)
        self.btn(box, "gen_measures", self.gen_measures, accent=True).grid(row=6, column=0, columnspan=2, sticky="ew",
                                                                           pady=(8, 0))
        ai = self.lframe(left, "ai_dax_box")
        ai.pack(fill="x", pady=(12, 0))
        self.dax_request = TextBox(ai, height=4, width=46)
        self.dax_request.pack(fill="x")
        self.lbl(ai, "ai_dax_hint", hint=True, wraplength=380, justify="left").pack(anchor="w", pady=4)
        self.btn(ai, "gen_dax_ai", self.gen_dax_ai, accent=True).pack(fill="x")

        right = self.lframe(tab, "dax_out", padding=6)
        right.pack(side="left", fill="both", expand=True)
        self.dax_out = TextBox(right, height=20, font=MONO, wrap="none")
        self.dax_out.pack(fill="both", expand=True)
        bar = ttk.Frame(right)
        bar.pack(fill="x", pady=(6, 0))
        self.btn(bar, "copy_all", lambda: self.copy(self.dax_out.get())).pack(side="left")
        self.btn(bar, "explain_dax", self.explain_dax).pack(side="left", padx=6)
        self.tr_call(self._default_date_ref)

    def _default_date_ref(self):
        c = logic.CAL[self.t.lang]
        current = self.date_entry.get()
        if not current or any(current == f"'{x['table']}'[{x['date']}]" for x in logic.CAL.values()):
            self._set_entry(self.date_entry, f"'{c['table']}'[{c['date']}]")

    @staticmethod
    def _set_entry(entry, value):
        entry.delete(0, "end")
        entry.insert(0, value)

    def _refresh_choices(self):
        names = list(self.tables)
        facts = [n for n in names if logic.table_role(self.tables[n]) == "fact"] or names
        self.fact_combo["values"] = names
        self.pq_combo["values"] = names
        if names:
            self.fact_combo.set(facts[0])
            self.pq_combo.set(names[0])
        ranks = [""]
        for name, df in self.tables.items():
            if logic.table_role(df) == "dimension":
                ranks += [f"{name}[{c}]" for c in df.columns if not logic.is_id_name(c) and
                          not pd.api.types.is_numeric_dtype(df[c])]
        self.rank_combo["values"] = ranks
        self.rank_combo.set(ranks[1] if len(ranks) > 1 else "")
        self._refresh_value_cols()

    def _refresh_value_cols(self):
        df = self.tables.get(self.fact_combo.get())
        if df is None:
            return
        numeric = [str(c) for c in df.columns if pd.api.types.is_numeric_dtype(df[c]) and not logic.is_id_name(c)]
        others = [str(c) for c in df.columns if str(c) not in numeric]
        self.value_combo["values"] = numeric + others
        self.value_combo.set(numeric[-1] if numeric else (others[0] if others else ""))

    def gen_measures(self):
        table, value = self.fact_combo.get(), self.value_combo.get()
        if not self.need(table and value, "need_model"):
            return
        rank = self.rank_combo.get()
        rank_table = rank_col = None
        if rank and "[" in rank:
            rank_table, rank_col = rank[:-1].split("[", 1)
        self.dax_out.set(logic.dax_measures(table, value, self.t.lang, self.date_entry.get().strip() or None,
                                            rank_table, rank_col, self.agg_combo.get()))

    def gen_dax_ai(self):
        request = self.dax_request.get()
        if not self.need(request, "need_text"):
            return
        model = logic.model_description(self.tables, self.rels) if self.tables else "(no model loaded)"
        date_ref = self.date_entry.get().strip()
        prompt = f"Model:\n{model}\nCalendar date column: {date_ref or 'none'}\n\nRequest: {request}"

        def done(data):
            text = f"-- {data.get('explanation', '')}\n{data.get('dax', '')}\n"
            if data.get("format_string"):
                text += f"-- Format: {data['format_string']}\n"
            if data.get("notes"):
                text += f"-- {data['notes']}\n"
            self.dax_out.set(text)

        self.run_ai(lambda ai: ai.ask_json(logic.DAX_SYSTEM, prompt, logic.DAX_SCHEMA), done)

    def explain_dax(self):
        code = self.dax_out.get()
        if not self.need(code, "need_text"):
            return

        def done(text):
            self.dax_out.set(code + "\n\n/*\n" + text + "\n*/\n")

        self.run_ai(lambda ai: ai.ask(logic.EXPLAIN_DAX_SYSTEM, code, max_tokens=6000), done)

    # ------------------------------------------------------------ calendario
    def _build_calendar(self):
        tab = self.add_tab("tab_calendar")
        box = self.lframe(tab, "cal_box")
        box.pack(fill="x")
        self.lbl(box, "cal_hint", hint=True, wraplength=1050, justify="left").grid(row=0, column=0, columnspan=6,
                                                                                    sticky="w", pady=(0, 6))
        self.lbl(box, "start").grid(row=1, column=0, sticky="w")
        self.cal_start = ttk.Entry(box, width=14)
        self.cal_start.grid(row=1, column=1, padx=(6, 18))
        self.lbl(box, "end").grid(row=1, column=2, sticky="w")
        self.cal_end = ttk.Entry(box, width=14)
        self.cal_end.grid(row=1, column=3, padx=(6, 18))
        self.lbl(box, "fiscal").grid(row=1, column=4, sticky="w")
        self.fiscal = ttk.Spinbox(box, from_=1, to=12, width=5)
        self.fiscal.set(1)
        self.fiscal.grid(row=1, column=5, padx=6)
        start, end = logic.date_range({})
        self._set_entry(self.cal_start, start.isoformat())
        self._set_entry(self.cal_end, end.isoformat())
        bar = ttk.Frame(box)
        bar.grid(row=2, column=0, columnspan=6, sticky="w", pady=(10, 0))
        self.btn(bar, "cal_dax", lambda: self.gen_calendar("dax"), accent=True).pack(side="left")
        self.btn(bar, "cal_m", lambda: self.gen_calendar("m")).pack(side="left", padx=6)
        self.btn(bar, "cal_excel", self.export_calendar).pack(side="left")
        self.btn(bar, "copy_all", lambda: self.copy(self.cal_out.get())).pack(side="left", padx=6)
        self.cal_out = TextBox(tab, height=20, font=MONO, wrap="none")
        self.cal_out.pack(fill="both", expand=True, pady=(10, 0))

    def _fiscal(self):
        try:
            return max(1, min(12, int(self.fiscal.get())))
        except ValueError:
            return 1

    def _cal_dates(self):
        try:
            start = date.fromisoformat(self.cal_start.get().strip())
            end = date.fromisoformat(self.cal_end.get().strip())
        except ValueError:
            self.need(False, "bad_date")
            return None
        return (start, end) if start <= end else (end, start)

    def gen_calendar(self, kind):
        dates = self._cal_dates()
        if dates:
            fn = logic.calendar_dax if kind == "dax" else logic.calendar_m
            self.cal_out.set(fn(dates[0], dates[1], self.t.lang, self._fiscal()))

    def export_calendar(self):
        dates = self._cal_dates()
        if not dates:
            return
        path = self.ask_save("excel", f"{logic.CAL[self.t.lang]['table']}.xlsx")
        if path:
            from openpyxl import Workbook

            from ..tabular import write_df

            wb = Workbook()
            wb.active.title = logic.CAL[self.t.lang]["table"]
            write_df(wb.active, logic.calendar_df(dates[0], dates[1], self.t.lang, self._fiscal()), accent="B8860B")
            wb.save(path)
            self.saved(path)

    # ------------------------------------------------------------ Power Query
    def _build_pq(self):
        tab = self.add_tab("tab_pq")
        left = ttk.Frame(tab)
        left.pack(side="left", fill="y", padx=(0, 10))
        box = self.lframe(left, "pq_box")
        box.pack(fill="x")
        self.lbl(box, "pq_hint", hint=True, wraplength=380, justify="left").grid(row=0, column=0, columnspan=2,
                                                                                  sticky="w", pady=(0, 6))
        self.pq_combo = self.field(box, "table", ttk.Combobox(box, state="readonly", width=28), 1)
        self.pq_combo.bind("<<ComboboxSelected>>", lambda e: self._pq_path())
        self.pq_path = self.field(box, "file_path", ttk.Entry(box, width=30), 2)
        self.btn(box, "gen_m", self.gen_m, accent=True).grid(row=3, column=0, columnspan=2, sticky="ew", pady=(8, 0))
        ai = self.lframe(left, "ai_m_box")
        ai.pack(fill="x", pady=(12, 0))
        self.m_request = TextBox(ai, height=4, width=46)
        self.m_request.pack(fill="x")
        self.lbl(ai, "ai_m_hint", hint=True, wraplength=380, justify="left").pack(anchor="w", pady=4)
        self.btn(ai, "gen_m_ai", self.gen_m_ai, accent=True).pack(fill="x")
        right = self.lframe(tab, "m_out", padding=6)
        right.pack(side="left", fill="both", expand=True)
        self.m_out = TextBox(right, height=20, font=MONO, wrap="none")
        self.m_out.pack(fill="both", expand=True)
        self.btn(right, "copy_all", lambda: self.copy(self.m_out.get())).pack(anchor="w", pady=(6, 0))

    def _pq_path(self):
        self._set_entry(self.pq_path, str(self.sources.get(self.pq_combo.get(), "")))

    def gen_m(self):
        name = self.pq_combo.get()
        if not self.need(name in self.tables, "need_model"):
            return
        if not self.pq_path.get():
            self._pq_path()
        self.m_out.set(logic.m_clean_query(self.tables[name], name, self.pq_path.get().strip() or None))

    def gen_m_ai(self):
        request = self.m_request.get()
        if not self.need(request, "need_text"):
            return
        name = self.pq_combo.get()
        context = ""
        if name in self.tables:
            df = self.tables[name]
            context = f"Table '{name}' columns: " + ", ".join(f"{c} ({logic._kind(df[c])})" for c in df.columns)
            if self.m_out.get().startswith("let"):
                context += f"\nCurrent query (reuse its Source steps):\n{self.m_out.get()}"
        prompt = f"{context}\n\nRequest: {request}"

        def done(data):
            self.m_out.set(f"// {data.get('explanation', '')}\n{data.get('m_code', '')}")

        self.run_ai(lambda ai: ai.ask_json(logic.M_SYSTEM, prompt, logic.M_SCHEMA), done)

    # ------------------------------------------------------------ tema
    def _build_theme(self):
        tab = self.add_tab("tab_theme")
        box = self.lframe(tab, "theme_box")
        box.pack(fill="x")
        self.lbl(box, "theme_hint", hint=True, wraplength=1050, justify="left").pack(anchor="w")
        row = ttk.Frame(box)
        row.pack(fill="x", pady=8)
        self.lbl(row, "theme_name").pack(side="left")
        self.theme_name = ttk.Entry(row, width=22)
        self.theme_name.insert(0, "Corporate")
        self.theme_name.pack(side="left", padx=6)
        self.btn(row, "brand_color", self.pick_brand).pack(side="left", padx=6)
        self.btn(row, "background", lambda: self.pick_extra("background")).pack(side="left")
        self.btn(row, "foreground", lambda: self.pick_extra("foreground")).pack(side="left", padx=6)
        self.lbl(row, "font").pack(side="left", padx=(12, 4))
        self.font_combo = ttk.Combobox(row, values=FONTS, width=16)
        self.font_combo.set(FONTS[0])
        self.font_combo.bind("<<ComboboxSelected>>", lambda e: self._draw_preview())
        self.font_combo.pack(side="left")
        self.btn(row, "save_theme", self.save_theme, accent=True).pack(side="right")

        self.swatches = ttk.Frame(box)
        self.swatches.pack(fill="x")
        self.lbl(box, "click_to_edit", hint=True).pack(anchor="w", pady=(4, 0))
        prev = self.lframe(tab, "preview_chart", padding=6)
        prev.pack(fill="both", expand=True, pady=(10, 0))
        self.canvas = tk.Canvas(prev, highlightthickness=0, height=330)
        self.canvas.pack(fill="both", expand=True)
        self.canvas.bind("<Configure>", lambda e: self._draw_preview())
        self._draw_swatches()

    def pick_brand(self):
        color = colorchooser.askcolor(self.colors[0], parent=self)[1]
        if color:
            self.colors = logic.palette_from(color)
            self._draw_swatches()

    def pick_extra(self, attr):
        color = colorchooser.askcolor(getattr(self, attr), parent=self)[1]
        if color:
            setattr(self, attr, color.upper())
            self._draw_preview()

    def _edit_color(self, index):
        color = colorchooser.askcolor(self.colors[index], parent=self)[1]
        if color:
            self.colors[index] = color.upper()
            self._draw_swatches()

    def _draw_swatches(self):
        for child in self.swatches.winfo_children():
            child.destroy()
        for i, color in enumerate(self.colors):
            sw = tk.Label(self.swatches, text=color, bg=color, fg="white", width=11, height=2, cursor="hand2",
                          font=("Segoe UI", 9, "bold"))
            sw.pack(side="left", padx=3, pady=4)
            sw.bind("<Button-1>", lambda e, i=i: self._edit_color(i))
        self._draw_preview()

    def _draw_preview(self):
        c = self.canvas
        c.delete("all")
        w, h = max(c.winfo_width(), 600), max(c.winfo_height(), 300)
        c.configure(background=self.background)
        font = self.font_combo.get() or "Segoe UI"
        c.create_text(20, 18, text=self.theme_name.get() or "Theme", anchor="w", fill=self.foreground,
                      font=(font, 14, "bold"))
        values = [62, 48, 75, 33, 54, 41, 68, 29]
        n = len(self.colors)
        bar_w = (w - 80) / (n * 1.6)
        base = h - 40
        for i, (color, v) in enumerate(zip(self.colors, values)):
            x = 40 + i * bar_w * 1.6
            top = base - (base - 60) * v / 80
            c.create_rectangle(x, top, x + bar_w, base, fill=color, outline="")
            c.create_text(x + bar_w / 2, top - 10, text=str(v), fill=self.foreground, font=(font, 9))
        c.create_line(30, base, w - 30, base, fill=self.foreground)

    def save_theme(self):
        name = self.theme_name.get().strip() or "Theme"
        path = self.ask_save("json", f"{name}.json")
        if path:
            Path(path).write_text(logic.theme_json(name, self.colors, self.background, self.foreground,
                                                   self.font_combo.get() or "Segoe UI"), encoding="utf-8")
            self.saved(path)


def main():
    BIPilot().mainloop()


if __name__ == "__main__":
    main()
