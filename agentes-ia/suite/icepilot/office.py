"""Exportación de los juegos a PowerPoint (presentación con respuestas) y Word (guía, cartones e impresos)."""

from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PPt

from .games import GAMES, prepare, theme_colors
from .web_text import RULES

TEXT = {
    "es": {"how": "Cómo se juega", "answer": "Respuesta", "explain": "Explicación", "question": "Pregunta",
           "round": "Ronda", "thanks": "¡Gracias por jugar!", "guide": "Guía del moderador", "key": "Clave de respuestas",
           "sheet": "Hoja de respuestas", "name": "Nombre / equipo", "tips": "Consejos para Teams", "card": "Cartón",
           "free": "LIBRE", "words": "Palabras", "find": "Encuentra estas palabras", "lie": "Mentira",
           "board": "Tablero", "category": "Categoría", "points": "Puntos", "terms": "Términos para cantar",
           "would": "¿Qué prefieres?", "spin": "Pregunta rompehielos", "statements": "Afirmaciones",
           "content": "Contenido del juego", "true": "Verdadero", "false": "Falso", "prompt": "Pista visual",
           "date": "Fecha", "options": "Opciones",
           "tips_list": ["Comparte solo la ventana del navegador o de PowerPoint y usa el botón de pantalla completa.",
                         "Al compartir, activa «Incluir sonido» si quieres que se oigan los efectos.",
                         "Pide las respuestas por el chat para que participe todo el mundo.",
                         "Bingo: envía el archivo HTML por el chat; cada persona lo abre y elige su número de cartón.",
                         "Los juegos funcionan sin internet: no se conectan a ningún servidor ni envían datos."]},
    "en": {"how": "How to play", "answer": "Answer", "explain": "Explanation", "question": "Question", "round": "Round",
           "thanks": "Thanks for playing!", "guide": "Host guide", "key": "Answer key", "sheet": "Answer sheet",
           "name": "Name / team", "tips": "Tips for Teams", "card": "Card", "free": "FREE", "words": "Words",
           "find": "Find these words", "lie": "Lie", "board": "Board", "category": "Category", "points": "Points",
           "terms": "Terms to call", "would": "Would you rather…?", "spin": "Icebreaker question",
           "statements": "Statements", "content": "Game content", "true": "True", "false": "False",
           "prompt": "Visual clue", "date": "Date", "options": "Options",
           "tips_list": ["Share only the browser or PowerPoint window and use the full-screen button.",
                         "When sharing, turn on “Include sound” if you want the effects to be heard.",
                         "Ask for answers in the chat so everyone takes part.",
                         "Bingo: send the HTML file in the chat; each person opens it and picks their card number.",
                         "The games work offline: they don't connect to any server or send any data."]},
    "fr": {"how": "Comment jouer", "answer": "Réponse", "explain": "Explication", "question": "Question",
           "round": "Manche", "thanks": "Merci d'avoir joué !", "guide": "Guide de l'animateur",
           "key": "Corrigé", "sheet": "Feuille de réponses", "name": "Nom / équipe", "tips": "Conseils pour Teams",
           "card": "Carton", "free": "LIBRE", "words": "Mots", "find": "Trouvez ces mots", "lie": "Mensonge",
           "board": "Tableau", "category": "Catégorie", "points": "Points", "terms": "Termes à tirer",
           "would": "Tu préfères… ?", "spin": "Question brise-glace", "statements": "Affirmations",
           "content": "Contenu du jeu", "true": "Vrai", "false": "Faux", "prompt": "Indice visuel", "date": "Date",
           "options": "Options",
           "tips_list": ["Partagez uniquement la fenêtre du navigateur ou de PowerPoint et utilisez le bouton plein écran.",
                         "Au partage, activez « Inclure le son » si vous voulez entendre les effets.",
                         "Demandez les réponses dans le chat pour que tout le monde participe.",
                         "Bingo : envoyez le fichier HTML dans le chat ; chacun l'ouvre et choisit son numéro de carton.",
                         "Les jeux fonctionnent sans internet : ils ne se connectent à aucun serveur."]},
}
OPTION_COLORS = ["E21B3C", "1368CE", "D89E00", "26890C"]
SHAPES = ["▲", "◆", "●", "■"]


def _hex(color):
    return color.lstrip("#").upper()


# ---------------------------------------------------------------- PowerPoint

class Deck:
    W, H = Inches(13.333), Inches(7.5)

    def __init__(self, cfg):
        self.cfg = cfg
        self.theme = theme_colors(cfg)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = self.W, self.H
        self.font = "Segoe UI"

    def slide(self, dark=False, notes=""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = PRGB.from_string(_hex(self.theme["primary"] if dark else
                                                                 ("0F172A" if self.theme["dark"] else "FFFFFF")))
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def rect(self, s, x, y, w, h, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        r = s.shapes.add_shape(shape, x, y, w, h)
        r.fill.solid()
        r.fill.fore_color.rgb = PRGB.from_string(_hex(color))
        r.line.fill.background()
        r.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            r.adjustments[0] = 0.12
        return r

    def text(self, s, x, y, w, h, text, size=24, color=None, bold=False, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, font=None):
        box = s.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        lines = str(text).split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            run = p.add_run()
            run.text = line
            run.font.size, run.font.bold, run.font.name = PPt(size), bold, font or self.font
            run.font.color.rgb = PRGB.from_string(_hex(color or ("E2E8F0" if self.theme["dark"] else "0F172A")))
        return box

    def header(self, s, label):
        self.rect(s, 0, 0, self.W, Inches(0.9), self.theme["primary"], MSO_SHAPE.RECTANGLE)
        self.text(s, Inches(0.5), Inches(0.1), Inches(9), Inches(0.7), f"{self.cfg.get('emoji', '')}  {self.title}",
                  18, "FFFFFF", True, PP_ALIGN.LEFT)
        self.text(s, Inches(9), Inches(0.1), Inches(3.8), Inches(0.7), label, 16, "FFFFFF", False, PP_ALIGN.RIGHT)

    @property
    def title(self):
        return self.cfg.get("title") or "IcePilot"

    def cover(self, rules):
        s = self.slide(dark=True)
        self.text(s, 0, Inches(1.2), self.W, Inches(1.6), self.cfg.get("emoji", ""), 80, "FFFFFF", font="Segoe UI Emoji")
        self.text(s, Inches(0.8), Inches(2.9), Inches(11.7), Inches(1.4), self.title, 48, "FFFFFF", True)
        if self.cfg.get("subtitle"):
            self.text(s, Inches(0.8), Inches(4.3), Inches(11.7), Inches(0.8), self.cfg["subtitle"], 22, "DCE6F5")
        self.rect(s, Inches(5.9), Inches(5.4), Inches(1.5), Inches(0.1), self.theme["accent"], MSO_SHAPE.RECTANGLE)
        s = self.slide()
        self.header(s, "")
        t = TEXT[self.cfg.get("lang", "es")]
        self.text(s, Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.9), "📖 " + t["how"], 34, None, True, PP_ALIGN.LEFT)
        self.text(s, Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.5), "\n".join(f"•  {r}" for r in rules), 24, None,
                  False, PP_ALIGN.LEFT, MSO_ANCHOR.TOP)

    def closing(self):
        s = self.slide(dark=True)
        self.text(s, 0, Inches(2.2), self.W, Inches(1.5), "🏆", 80, "FFFFFF", font="Segoe UI Emoji")
        self.text(s, 0, Inches(3.8), self.W, Inches(1.2), TEXT[self.cfg.get("lang", "es")]["thanks"], 48, "FFFFFF", True)

    def options(self, s, options, correct=None, top=Inches(3.0)):
        n = len(options)
        cols = 2
        w, h, gap = Inches(5.9), Inches(1.6) if n > 2 else Inches(2.2), Inches(0.3)
        for k, opt in enumerate(options):
            x = Inches(0.6) + (k % cols) * (w + gap)
            y = top + (k // cols) * (h + gap)
            color = OPTION_COLORS[k % 4]
            if correct is not None and k != correct:
                color = "CBD5E1"
            box = self.rect(s, x, y, w, h, color)
            if correct is not None and k == correct:
                box.line.color.rgb = PRGB.from_string("16A34A")
                box.line.width = PPt(6)
            self.text(s, x + Inches(0.2), y, Inches(0.8), h, SHAPES[k % 4], 28, "FFFFFF", True)
            label = opt + ("   ✔" if correct is not None and k == correct else "")
            self.text(s, x + Inches(0.9), y, w - Inches(1.1), h, label, 24 if len(opt) < 60 else 18, "FFFFFF", True,
                      PP_ALIGN.LEFT)


def build_pptx(cfg, path):
    data, warnings = prepare(cfg)
    game, lang = cfg["type"], cfg.get("lang", "es")
    t = TEXT[lang]
    deck = Deck(cfg)
    deck.cover(RULES[game][lang])
    if game in ("quiz", "truefalse"):
        for i, q in enumerate(data["questions"], start=1):
            label = f"{t['question']} {i}/{len(data['questions'])}"
            for reveal in (False, True):
                s = deck.slide(notes=f"{t['answer']}: {q['answer']}" + (f"\n{q['explain']}" if q["explain"] else ""))
                deck.header(s, label)
                deck.text(s, Inches(0.6), Inches(1.1), Inches(12.1), Inches(1.7), q["q"], 32 if len(q["q"]) < 90 else 26,
                          None, True)
                if q["options"]:
                    deck.options(s, q["options"], q["correct"] if reveal else None)
                elif reveal:
                    deck.rect(s, Inches(2.5), Inches(3.4), Inches(8.3), Inches(1.4), "16A34A")
                    deck.text(s, Inches(2.5), Inches(3.4), Inches(8.3), Inches(1.4), "✔ " + q["answer"], 32, "FFFFFF", True)
                if reveal and q["explain"]:
                    deck.text(s, Inches(0.6), Inches(6.6), Inches(12.1), Inches(0.8), "💡 " + q["explain"], 18, None,
                              False, PP_ALIGN.CENTER)
    elif game == "trivia":
        board = data["board"]
        s = deck.slide()
        deck.header(s, t["board"])
        cols = max(1, len(board))
        w = (12.1 - 0.2 * (cols - 1)) / cols
        for ci, cat in enumerate(board):
            x = Inches(0.6 + ci * (w + 0.2))
            deck.rect(s, x, Inches(1.2), Inches(w), Inches(0.9), deck.theme["primary"])
            deck.text(s, x, Inches(1.2), Inches(w), Inches(0.9), cat["category"].upper(), 18, "FFFFFF", True)
            for r, q in enumerate(cat["questions"]):
                y = Inches(2.3 + r * 1.0)
                deck.rect(s, x, y, Inches(w), Inches(0.85), "E2E8F0")
                deck.text(s, x, y, Inches(w), Inches(0.85), str(q["points"]), 28, deck.theme["primary"], True)
        for cat in board:
            for q in cat["questions"]:
                for reveal in (False, True):
                    s = deck.slide(notes=f"{t['answer']}: {q['a']}")
                    deck.header(s, f"{cat['category']} · {q['points']}")
                    deck.text(s, Inches(0.8), Inches(1.5), Inches(11.7), Inches(2.6), q["q"], 36, None, True)
                    if reveal:
                        deck.rect(s, Inches(2.5), Inches(4.4), Inches(8.3), Inches(1.4), "16A34A")
                        deck.text(s, Inches(2.5), Inches(4.4), Inches(8.3), Inches(1.4), "✔ " + q["a"], 30, "FFFFFF", True)
    elif game in ("emoji", "scramble"):
        for i, r in enumerate(data["rounds"], start=1):
            for reveal in (False, True):
                s = deck.slide(notes=f"{t['answer']}: {r['answer']}" + (f"\n💡 {r['hint']}" if r["hint"] else ""))
                deck.header(s, f"{t['round']} {i}/{len(data['rounds'])}")
                font = "Segoe UI Emoji" if game == "emoji" else "Consolas"
                deck.text(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(2.8), r["prompt"], 96 if game == "emoji" else 60,
                          deck.theme["primary"], True, font=font)
                if reveal:
                    deck.rect(s, Inches(2.5), Inches(4.7), Inches(8.3), Inches(1.4), "16A34A")
                    deck.text(s, Inches(2.5), Inches(4.7), Inches(8.3), Inches(1.4), "✔ " + r["answer"], 32, "FFFFFF", True)
    elif game == "wouldrather":
        for i, p in enumerate(data["pairs"], start=1):
            s = deck.slide()
            deck.header(s, f"{t['would']} {i}/{len(data['pairs'])}")
            for k, (text, color) in enumerate(((p["a"], deck.theme["primary"]), (p["b"], deck.theme["accent"]))):
                x = Inches(0.6 + k * 6.3)
                deck.rect(s, x, Inches(1.4), Inches(5.8), Inches(5.4), color)
                deck.text(s, x, Inches(1.6), Inches(5.8), Inches(0.9), "AB"[k], 40, "FFFFFF", True)
                deck.text(s, x + Inches(0.3), Inches(2.5), Inches(5.2), Inches(3.8), text, 32, "FFFFFF", True)
            deck.rect(s, Inches(6.07), Inches(3.5), Inches(1.2), Inches(1.2), "FFFFFF", MSO_SHAPE.OVAL)
            deck.text(s, Inches(6.07), Inches(3.5), Inches(1.2), Inches(1.2), "VS", 26, "0F172A", True)
    elif game == "truths":
        for p in data["people"]:
            for reveal in (False, True):
                s = deck.slide(notes=f"{t['lie']}: {p['lie'] + 1}. {p['statements'][p['lie']]}")
                deck.header(s, t["statements"])
                deck.text(s, Inches(0.6), Inches(1.1), Inches(12.1), Inches(0.9), "🙋 " + p["person"], 34, None, True)
                for k, st in enumerate(p["statements"]):
                    x = Inches(0.6 + k * 4.1)
                    color = ("DC2626" if k == p["lie"] else "16A34A") if reveal else deck.theme["accent"]
                    deck.rect(s, x, Inches(2.3), Inches(3.8), Inches(4.4), color)
                    deck.text(s, x, Inches(2.4), Inches(3.8), Inches(0.9), str(k + 1), 40, "FFFFFF", True)
                    deck.text(s, x + Inches(0.2), Inches(3.3), Inches(3.4), Inches(3.2), st, 24, "FFFFFF", True)
    elif game == "wheel":
        for i, entry in enumerate(data["entries"], start=1):
            s = deck.slide()
            deck.header(s, f"{t['spin']} {i}/{len(data['entries'])}")
            deck.text(s, 0, Inches(1.3), deck.W, Inches(1.6), "🎯", 72, None, font="Segoe UI Emoji")
            deck.text(s, Inches(0.8), Inches(3.0), Inches(11.7), Inches(3.0), entry, 40, deck.theme["primary"], True)
    else:
        raise ValueError(game)
    deck.closing()
    deck.prs.save(path)
    return path, warnings


# ---------------------------------------------------------------- Word

def _shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex(color))
    tc_pr.append(shd)


def _cell_text(cell, text, size=11, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.font.size, run.bold = Pt(size), bold
    if color:
        run.font.color.rgb = RGBColor.from_string(_hex(color))


def _new_doc(cfg):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.6)
        sec.left_margin = sec.right_margin = Cm(1.8)
    primary = RGBColor.from_string(_hex(theme_colors(cfg)["primary"]))
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = primary
        doc.styles[name].font.name = "Calibri"
    return doc


def _title(doc, cfg, label):
    doc.add_paragraph(f"{cfg.get('emoji', '')} {cfg.get('title') or 'IcePilot'}", style="Title")
    p = doc.add_paragraph()
    run = p.add_run(" · ".join(x for x in (label, cfg.get("subtitle", ""), date.today().isoformat()) if x))
    run.italic = True
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _content_table(doc, header, rows):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, header):
        _cell_text(cell, text, 10, True)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            _cell_text(cell, str(text), 10, align=WD_ALIGN_PARAGRAPH.LEFT)
    return table


def _guide(doc, cfg, data, t):
    game, lang = cfg["type"], cfg.get("lang", "es")
    doc.add_heading(t["how"], level=1)
    for rule in RULES[game][lang]:
        doc.add_paragraph(rule, style="List Bullet")
    doc.add_heading(t["tips"], level=1)
    for tip in t["tips_list"]:
        doc.add_paragraph(tip, style="List Bullet")
    doc.add_heading(t["content"], level=1)
    if game in ("quiz", "truefalse", "snakes"):
        rows = [(i, q["q"], " | ".join(q["options"]), q["answer"], q.get("explain", ""))
                for i, q in enumerate(data["questions"], start=1)]
        _content_table(doc, ["#", t["question"], t["options"], t["answer"], t["explain"]], rows)
    elif game == "trivia":
        rows = [(c["category"], q["points"], q["q"], q["a"]) for c in data["board"] for q in c["questions"]]
        _content_table(doc, [t["category"], t["points"], t["question"], t["answer"]], rows)
    elif game == "hangman":
        _content_table(doc, ["#", t["answer"], "💡"], [(i, w["answer"], w["hint"]) for i, w in enumerate(data["words"], 1)])
    elif game in ("emoji", "scramble"):
        _content_table(doc, ["#", t["prompt"], t["answer"], "💡"],
                       [(i, r["prompt"], r["answer"], r["hint"]) for i, r in enumerate(data["rounds"], 1)])
    elif game == "wheel":
        _content_table(doc, ["#", t["spin"]], list(enumerate(data["entries"], 1)))
    elif game == "bingo":
        _content_table(doc, ["#", t["terms"]], list(enumerate(data["terms"], 1)))
    elif game == "wordsearch":
        _content_table(doc, [t["words"], "💡"], [(w["label"], w["clue"]) for w in data["words"]])
    elif game == "truths":
        _content_table(doc, ["🙋", "1", "2", "3", t["lie"]],
                       [(p["person"], *p["statements"], p["lie"] + 1) for p in data["people"]])
    elif game == "wouldrather":
        _content_table(doc, ["#", "A", "B"], [(i, p["a"], p["b"]) for i, p in enumerate(data["pairs"], 1)])


def _bingo_cards(doc, cfg, data, t):
    size, terms = data["size"], data["terms"]
    for n, card in enumerate(data["cards"], start=1):
        _page_break(doc)
        p = doc.add_paragraph(style="Title")
        p.add_run(f"{cfg.get('emoji', '🎱')} {cfg.get('title') or 'Bingo'}")
        doc.add_paragraph(f"{t['card']} #{n}").runs[0].bold = True
        if size == 5:
            head = doc.add_table(rows=1, cols=5)
            head.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, ch in zip(head.rows[0].cells, "BINGO"):
                _cell_text(cell, ch, 22, True, theme_colors(cfg)["primary"])
        table = doc.add_table(rows=size, cols=size)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, idx in enumerate(card):
            cell = table.cell(k // size, k % size)
            cell.width = Cm(17 / size)
            if idx < 0:
                _cell_text(cell, "⭐ " + t["free"], 12, True, "FFFFFF")
                _shade(cell, theme_colors(cfg)["primary"])
            else:
                _cell_text(cell, terms[idx], 12 if size == 5 else 14, True)
        for row in table.rows:
            row.height = Cm(17 / size * 0.75)


def _wordsearch(doc, cfg, data, t):
    grid, placements = data["grid"], data["placements"]
    on_path = {}
    for k, p in enumerate(placements):
        for i in range(len(p["word"])):
            on_path[(p["r"] + p["dr"] * i, p["c"] + p["dc"] * i)] = k
    for solution in (False, True):
        _page_break(doc)
        doc.add_heading(t["key"] if solution else t["find"], level=1)
        n = len(grid)
        table = doc.add_table(rows=n, cols=n)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        size = 16 if n <= 12 else 12
        for r in range(n):
            for c in range(n):
                cell = table.cell(r, c)
                cell.width = Cm(16 / n)
                hit = solution and (r, c) in on_path
                _cell_text(cell, grid[r][c], size, hit, "FFFFFF" if hit else None)
                if hit:
                    _shade(cell, theme_colors(cfg)["accent"])
        if not solution:
            doc.add_paragraph()
            words = [w["clue"] if cfg.get("options", {}).get("use_clues") and w["clue"] else w["label"]
                     for w in data["words"]]
            doc.add_paragraph("   ☐ " + "      ☐ ".join(words))


def _worksheet(doc, cfg, data, t):
    game = cfg["type"]
    _page_break(doc)
    doc.add_heading(t["sheet"], level=1)
    doc.add_paragraph(f"{t['name']}: ______________________________     {t['date']}: ____________")
    if game in ("quiz", "truefalse"):
        for i, q in enumerate(data["questions"], start=1):
            doc.add_paragraph(f"{i}. {q['q']}").runs[0].bold = True
            if q["options"]:
                doc.add_paragraph("     " + "     ".join(f"☐ {'ABCD'[k]}) {o}" for k, o in enumerate(q["options"])))
            else:
                doc.add_paragraph("     ____________________________________________")
    else:
        for i, r in enumerate(data["rounds"], start=1):
            doc.add_paragraph(f"{i}.  {r['prompt']}     →  ______________________________")
    _page_break(doc)
    doc.add_heading(t["key"], level=1)
    items = data.get("questions") or data.get("rounds") or []
    for i, q in enumerate(items, start=1):
        doc.add_paragraph(f"{i}. {q['answer']}")


def build_docx(cfg, path):
    data, warnings = prepare(cfg)
    lang = cfg.get("lang", "es")
    t = TEXT[lang]
    doc = _new_doc(cfg)
    _title(doc, cfg, t["guide"])
    _guide(doc, cfg, data, t)
    game = cfg["type"]
    if game == "bingo" and data.get("cards"):
        _bingo_cards(doc, cfg, data, t)
    elif game == "wordsearch":
        _wordsearch(doc, cfg, data, t)
    elif game in ("quiz", "truefalse", "emoji", "scramble"):
        _worksheet(doc, cfg, data, t)
    doc.save(path)
    return path, warnings


def supports(game, kind):
    return kind == "docx" or (kind == "pptx" and GAMES[game]["pptx"]) or kind == "html"

