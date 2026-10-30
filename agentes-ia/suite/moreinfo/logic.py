"""MoreInfo Pilot: builds the "MoreInfo" evidence file of a fraud package from screenshots.

Two outputs:
- Word/PDF: every screenshot goes into a Word document (case header, section titles, captions) that is saved as
  "CASE,MOREINFO.pdf". The PDF is made with Microsoft Word when it is installed (best quality) or with the
  built-in renderer otherwise, so it also works on PCs without Word and without internet.
- Excel: the analyst's downloaded transaction file (.xlsx/.xlsm/.xls) gets a new "Additional info" tab with the
  screenshots, saved as "CASE,MOREINFO.xlsx".
Sections are free names chosen by the analyst; the screenshots are grouped by section in the order given.

Files are numbered when the case already has one in any destination folder: CASE,MOREINFO.pdf, CASE,MOREINFO1.pdf,
CASE,MOREINFO2.pdf… Extra documents of the case get the next numbers too.
"""

import re
import shutil
import sys
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

EXCEL_SHEET = "Additional info"  # pestaña que se añade al Excel de transacciones
IMAGE_TYPES = (".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff")
EXCEL_TYPES = (".xlsx", ".xlsm", ".xls")
DEFAULT_BASE = "MOREINFO"


@dataclass
class Shot:
    section: str
    path: str
    caption: str = ""


def clean_case(case):
    """Case number safe for a file name (no \\ / : * ? " < > |)."""
    return re.sub(r'[\\/:*?"<>|\s]+', "", (case or "").strip())


def taken_numbers(case, base, folders):
    """Numbers already used by this case in any folder: "" counts as 0, "1" as 1… (any extension)."""
    pattern = re.compile(rf"^{re.escape(case)},{re.escape(base)}(\d*)\.[^.]+$", re.I)
    used = set()
    for folder in folders:
        try:
            names = [p.name for p in Path(folder).iterdir()]
        except OSError:
            continue
        for name in names:
            m = pattern.match(name)
            if m:
                used.add(int(m.group(1) or 0))
    return used


def next_names(case, base, folders, extensions):
    """Free names for several files, in order: CASE,BASE.ext, then CASE,BASE1.ext, CASE,BASE2.ext…"""
    used = taken_numbers(case, base, folders)
    names, n = [], 0
    for ext in extensions:
        while n in used:
            n += 1
        names.append(f"{case},{base}{n or ''}{ext.lower()}")
        used.add(n)
    return names


# ---------------------------------------------------------------- Word

def build_docx(path, case, analyst, shots, title="MoreInfo"):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    width = sec.page_width - sec.left_margin - sec.right_margin
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10)
    head = doc.add_paragraph()
    run = head.add_run(f"{title} – Case {case}")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(16), RGBColor(0x0E, 0x74, 0x90)
    meta = doc.add_paragraph(f"Prepared by: {analyst or '-'}    Date: {datetime.now():%Y-%m-%d %H:%M}")
    meta.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    current = None
    for shot in shots:
        if shot.section != current:
            current = shot.section
            doc.add_heading(current, level=2)
        with Image.open(shot.path) as img:
            w = img.width
        pic_width = min(width, Cm(w / 96 * 2.54)) if w else width
        doc.add_picture(shot.path, width=pic_width)
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if shot.caption:
            cap = doc.add_paragraph(shot.caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
    doc.save(path)
    return path


def word_to_pdf(docx_path, pdf_path):
    """Converts with Microsoft Word (Windows). Raises an exception if Word is not available."""
    if sys.platform != "win32":
        raise RuntimeError("Microsoft Word is only available on Windows")
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(Path(docx_path).resolve()), ReadOnly=True)
        try:
            doc.ExportAsFixedFormat(str(Path(pdf_path).resolve()), 17)  # wdExportFormatPDF
        finally:
            doc.Close(False)
    finally:
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()
    return pdf_path


# ---------------------------------------------------------------- built-in PDF (no Word needed)

PAGE = (1240, 1754)  # A4 at 150 dpi
MARGIN = 70


def _font(size, bold=False):
    names = ("arialbd.ttf", "DejaVuSans-Bold.ttf") if bold else ("arial.ttf", "DejaVuSans.ttf")
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def render_pdf(path, case, analyst, shots, title="MoreInfo"):
    """Draws the same layout as the Word file straight into a PDF (one image per page area)."""
    big, mid, small = _font(34, True), _font(24, True), _font(18)
    pages, page, y = [], None, 0

    def new_page():
        nonlocal page, y
        page = Image.new("RGB", PAGE, "white")
        pages.append(page)
        y = MARGIN
        draw = ImageDraw.Draw(page)
        draw.text((MARGIN, PAGE[1] - MARGIN + 20), f"Case {case} – page {len(pages)}", fill="#888888", font=small)
        return draw

    draw = new_page()
    draw.text((MARGIN, y), f"{title} – Case {case}", fill="#0E7490", font=big)
    y += 52
    draw.text((MARGIN, y), f"Prepared by: {analyst or '-'}    Date: {datetime.now():%Y-%m-%d %H:%M}",
              fill="#555555", font=small)
    y += 50
    bottom = PAGE[1] - MARGIN
    current = None
    for shot in shots:
        with Image.open(shot.path) as src:
            img = src.convert("RGB")
        max_w = PAGE[0] - 2 * MARGIN
        scale = min(1.0, max_w / img.width, (bottom - MARGIN - 120) / img.height)
        img = img.resize((max(1, int(img.width * scale)), max(1, int(img.height * scale))), Image.LANCZOS)
        need = img.height + (50 if shot.section != current else 0) + (40 if shot.caption else 0) + 20
        if y + need > bottom:
            draw = new_page()
        if shot.section != current:
            current = shot.section
            draw.text((MARGIN, y), current, fill="#1F2937", font=mid)
            y += 44
        page.paste(img, ((PAGE[0] - img.width) // 2, y))
        draw.rectangle([(PAGE[0] - img.width) // 2 - 1, y - 1, (PAGE[0] + img.width) // 2, y + img.height],
                       outline="#BBBBBB")
        y += img.height + 8
        if shot.caption:
            draw.text((MARGIN, y), shot.caption, fill="#333333", font=small)
            y += 34
        y += 16
    pages[0].save(path, "PDF", resolution=150, save_all=True, append_images=pages[1:])
    return path


def image_to_pdf(src, dst):
    with Image.open(src) as img:
        img.convert("RGB").save(dst, "PDF", resolution=150)
    return dst


# ---------------------------------------------------------------- Excel

def build_excel(src, dst, case, analyst, shots, sheet=EXCEL_SHEET):
    """Copies the transaction file and adds a sheet with the screenshots. Returns the saved path."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font

    src = Path(src)
    if src.suffix.lower() == ".xls":
        wb = _xls_workbook(src)
    else:
        wb = load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
    if wb is None:
        wb = Workbook()
    name, n = sheet[:31] or EXCEL_SHEET, 2
    while name in wb.sheetnames:
        name, n = f"{sheet[:28]} {n}", n + 1
    ws = wb.create_sheet(name)
    ws["A1"] = f"{name} – Case {case}"
    ws["A1"].font = Font(bold=True, size=14, color="0E7490")
    ws["A2"] = f"Prepared by: {analyst or '-'}    Date: {datetime.now():%Y-%m-%d %H:%M}"
    row, current = 4, None
    for shot in shots:
        if shot.section != current:
            current = shot.section
            ws.cell(row=row, column=1, value=current).font = Font(bold=True, size=12)
            row += 1
        if shot.caption:
            ws.cell(row=row, column=1, value=shot.caption).font = Font(italic=True)
            row += 1
        image = XLImage(shot.path)
        scale = min(1.0, 900 / max(1, image.width))
        image.width, image.height = int(image.width * scale), int(image.height * scale)
        ws.add_image(image, f"A{row}")
        row += int(image.height / 20) + 2  # filas de ~20 px
    wb.save(dst)
    return dst


def _xls_workbook(path):
    """Old .xls files: read the cells (with xlrd) into a new workbook."""
    from openpyxl import Workbook

    import xlrd

    book = xlrd.open_workbook(str(path))
    wb = Workbook()
    wb.remove(wb.active)
    for sh in book.sheets():
        ws = wb.create_sheet(sh.name[:31])
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    continue
                ws.cell(row=r + 1, column=c + 1, value=value)
    return wb


# ---------------------------------------------------------------- case documents (No Pay Decision Letter…)

WORD_TYPES = (".docx", ".doc", ".rtf", ".odt")
KEEP_TYPES = (".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".msg", ".eml")  # se suben tal cual


def document_target(path):
    """Extension a document is saved with: Word and images become PDF; Excel, text and emails keep theirs."""
    ext = Path(path).suffix.lower()
    if ext in WORD_TYPES or ext in IMAGE_TYPES or ext == ".pdf":
        return ".pdf"
    return ext


def render_docx_pdf(src, dst, title=""):
    """Text-only PDF of a .docx (paragraphs and tables) for PCs without Microsoft Word."""
    from docx import Document

    doc = Document(str(src))
    lines = []
    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append("".join(t.text or "" for t in block.iter() if t.tag.endswith("}t")))
        elif tag == "tbl":
            for row in block.iter():
                if row.tag.endswith("}tr"):
                    cells = ["".join(t.text or "" for t in cell.iter() if t.tag.endswith("}t"))
                             for cell in row if cell.tag.endswith("}tc")]
                    lines.append(" | ".join(cells))
            lines.append("")
    font, bold = _font(20), _font(26, True)
    pages, page, draw, y = [], None, None, 0
    max_w = PAGE[0] - 2 * MARGIN

    def new_page():
        nonlocal page, draw, y
        page = Image.new("RGB", PAGE, "white")
        pages.append(page)
        draw = ImageDraw.Draw(page)
        y = MARGIN

    new_page()
    if title:
        draw.text((MARGIN, y), title, fill="#1F2937", font=bold)
        y += 50
    for line in lines:
        words, current = line.split(" "), ""
        wrapped = []
        for word in words:
            candidate = f"{current} {word}".strip()
            if draw.textlength(candidate, font=font) <= max_w or not current:
                current = candidate
            else:
                wrapped.append(current)
                current = word
        wrapped.append(current)
        for text in wrapped:
            if y + 30 > PAGE[1] - MARGIN:
                new_page()
            draw.text((MARGIN, y), text, fill="black", font=font)
            y += 30
        y += 8
    pages[0].save(dst, "PDF", resolution=150, save_all=True, append_images=pages[1:])
    return dst


def prepare_document(src, dest_pdf, use_word=True):
    """Saves one case document in the work folder. Returns (saved path, note key or None)."""
    src, ext = Path(src), Path(src).suffix.lower()
    if ext == ".pdf":
        shutil.copy2(src, dest_pdf)
        return dest_pdf, None
    if ext in IMAGE_TYPES:
        return image_to_pdf(src, dest_pdf), None
    if ext in WORD_TYPES:
        if use_word:
            try:
                return word_to_pdf(src, dest_pdf), None
            except Exception:  # Word no disponible
                pass
        if ext == ".docx":
            return render_docx_pdf(src, dest_pdf), "doc_text_only"
        kept = dest_pdf.with_suffix(ext)  # .doc/.rtf sin Word: se sube en su formato
        shutil.copy2(src, kept)
        return kept, "doc_kept"
    dest = dest_pdf.with_suffix(ext)
    shutil.copy2(src, dest)
    return dest, None


# ---------------------------------------------------------------- delivery

def deliver(plan):
    """Copies files to their folders: plan = [(file, [folders])]. Returns (copied paths, [(folder, error)])."""
    copied, errors, failed = [], [], set()
    for f, folders in plan:
        for folder in folders:
            if folder in failed:
                continue
            target = Path(folder)
            try:
                target.mkdir(parents=True, exist_ok=True)
                dest = target / Path(f).name
                if dest.exists():
                    raise FileExistsError(f"{dest.name} already exists")
                shutil.copy2(f, dest)
                copied.append(dest)
            except OSError as exc:
                failed.add(folder)
                errors.append((str(folder), str(exc)))
    return copied, errors


def create_package(case, analyst, shots, documents, folder, letter_folders=(), mode="pdf", excel_source=None,
                   base=DEFAULT_BASE, use_word=True, keep_docx=False, sheet=EXCEL_SHEET, include_moreinfo=True):
    """Builds CASE,MOREINFO (+ the case documents, numbered after it) and saves them.

    The MoreInfo and the documents go to `folder`. Documents marked as the No Pay Decision Letter
    (documents = [(path, is_letter)]) go to every folder in `letter_folders` instead (or to `folder` when that
    list is empty). Numbers never repeat in any of those folders.
    With include_moreinfo=False only the documents are uploaded.
    The files are made in a temporary folder that is deleted afterwards (no local copy is kept). Only if nothing
    could be written is that folder kept, so nothing is lost: its path is returned in "kept"."""
    case = clean_case(case)
    documents = [(d, False) if isinstance(d, (str, Path)) else (d[0], bool(d[1])) for d in documents]
    letter_folders = [f for f in letter_folders if f] or ([folder] if folder else [])
    if not case:
        raise ValueError("Enter the case number.")
    if not folder and (include_moreinfo or any(not is_letter for _, is_letter in documents)):
        raise ValueError("Choose the destination folder.")
    if include_moreinfo and mode == "excel" and not excel_source:
        raise ValueError("Choose the transaction Excel file.")
    if include_moreinfo and mode == "pdf" and not shots:
        raise ValueError("Add at least one screenshot.")
    if not include_moreinfo and not documents:
        raise ValueError("Add at least one document.")
    work = Path(tempfile.mkdtemp(prefix=f"moreinfo_{case}_"))
    exts = []
    if include_moreinfo:
        main_ext = ".xlsx" if mode == "excel" else ".pdf"
        if mode == "excel" and Path(excel_source).suffix.lower() == ".xlsm":
            main_ext = ".xlsm"
        exts.append(main_ext)
    exts += [document_target(d) for d, _ in documents]
    every = list(dict.fromkeys(str(f) for f in [folder, *letter_folders] if f))
    names = next_names(case, base, [Path(f) for f in every], exts)
    plan, engine, notes = [], None, []
    if include_moreinfo:
        main = work / names[0]
        if mode == "excel":
            build_excel(excel_source, main, case, analyst, shots, sheet)
            engine = "excel"
        else:
            docx = work / (Path(names[0]).stem + ".docx")
            build_docx(docx, case, analyst, shots)
            engine = "builtin"
            if use_word:
                try:
                    word_to_pdf(docx, main)
                    engine = "word"
                except Exception:  # Word no disponible: se usa el generador propio
                    engine = "builtin"
            if engine == "builtin":
                render_pdf(main, case, analyst, shots)
            if keep_docx:
                plan.append((docx, [folder]))
        plan.insert(0, (main, [folder]))
        names = names[1:]
    for (doc, is_letter), name in zip(documents, names, strict=True):
        saved, note = prepare_document(doc, work / Path(name).with_suffix(".pdf").name, use_word)
        if note:
            notes.append((note, Path(doc).name))
        plan.append((saved, letter_folders if is_letter else [folder]))
    copied, errors = deliver(plan)
    kept = None
    if copied:
        shutil.rmtree(work, ignore_errors=True)
    else:
        kept = work  # ninguna carpeta disponible: se conserva para no perder el trabajo
    return {"names": [f.name for f, _ in plan], "copied": copied, "errors": errors, "engine": engine,
            "case": case, "kept": kept, "notes": notes}
