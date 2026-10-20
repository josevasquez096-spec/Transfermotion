"""DocPilot: lógica sin interfaz (redacción, marcado → Word, legibilidad y combinación de correspondencia)."""

import copy
import math
import re
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..ai import STR, STR_LIST, obj
from ..i18n import LANGS

# ---------------------------------------------------------------- tipos de documento y plantillas

DOC_TYPES = ("proposal", "report", "letter", "minutes", "procedure", "email")
TONES = ("formal", "friendly", "persuasive", "technical")
LENGTHS = ("short", "medium", "long")

SKELETONS = {
    "proposal": {
        "es": ["Resumen ejecutivo", "Necesidad del cliente", "Solución propuesta", "Plan de trabajo y plazos",
               "Inversión", "Por qué nosotros", "Próximos pasos"],
        "en": ["Executive summary", "Client need", "Proposed solution", "Work plan and timeline", "Investment",
               "Why us", "Next steps"],
        "fr": ["Synthèse", "Besoin du client", "Solution proposée", "Plan de travail et calendrier", "Investissement",
               "Pourquoi nous", "Prochaines étapes"],
    },
    "report": {
        "es": ["Resumen", "Contexto y objetivos", "Metodología", "Resultados", "Conclusiones", "Recomendaciones"],
        "en": ["Summary", "Context and objectives", "Methodology", "Results", "Conclusions", "Recommendations"],
        "fr": ["Résumé", "Contexte et objectifs", "Méthodologie", "Résultats", "Conclusions", "Recommandations"],
    },
    "letter": {
        "es": ["Saludo", "Motivo de la carta", "Detalles", "Despedida"],
        "en": ["Greeting", "Purpose of the letter", "Details", "Closing"],
        "fr": ["Formule d'appel", "Objet de la lettre", "Détails", "Formule de politesse"],
    },
    "minutes": {
        "es": ["Datos de la reunión", "Asistentes", "Orden del día", "Temas tratados", "Acuerdos",
               "Tareas (responsable y fecha)", "Próxima reunión"],
        "en": ["Meeting details", "Attendees", "Agenda", "Discussion", "Decisions", "Action items (owner and date)",
               "Next meeting"],
        "fr": ["Informations de la réunion", "Participants", "Ordre du jour", "Points abordés", "Décisions",
               "Actions (responsable et date)", "Prochaine réunion"],
    },
    "procedure": {
        "es": ["Objetivo", "Alcance", "Responsables", "Definiciones", "Procedimiento paso a paso", "Controles",
               "Registros y anexos"],
        "en": ["Purpose", "Scope", "Responsibilities", "Definitions", "Step-by-step procedure", "Controls",
               "Records and appendices"],
        "fr": ["Objectif", "Périmètre", "Responsabilités", "Définitions", "Procédure pas à pas", "Contrôles",
               "Enregistrements et annexes"],
    },
    "email": {
        "es": ["Asunto", "Saludo", "Mensaje", "Llamada a la acción", "Firma"],
        "en": ["Subject", "Greeting", "Message", "Call to action", "Signature"],
        "fr": ["Objet", "Salutation", "Message", "Appel à l'action", "Signature"],
    },
}
PLACEHOLDER = {"es": "[Escribe aquí: {h}]", "en": "[Write here: {h}]", "fr": "[Écrire ici : {h}]"}


def skeleton(doc_type, lang, title, company="", recipient=""):
    lines = [f"# {title}"]
    meta = " · ".join(x for x in (company, recipient, format_date(date.today(), lang)) if x)
    lines.append(f"> {meta}")
    for heading in SKELETONS[doc_type][lang]:
        lines += ["", f"## {heading}", PLACEHOLDER[lang].format(h=heading.lower())]
    return "\n".join(lines)


DRAFT_SCHEMA = obj({
    "title": STR,
    "subtitle": STR,
    "sections": {"type": "array", "items": obj({"heading": STR, "paragraphs": STR_LIST, "bullets": STR_LIST})},
})
DRAFT_SYSTEM = """You are DocPilot, a professional business writer. Draft the requested document with a clear
structure, concrete content based ONLY on the facts provided, and the requested tone and length
(short ≈ 300 words, medium ≈ 700, long ≈ 1500). Where a fact is missing, write a clearly marked placeholder in
square brackets instead of inventing it (e.g. [amount], [date]). Paragraphs are plain text; you may use **bold**
for key terms. Use bullets only when listing items. `subtitle`: one line with context (company, recipient, date)."""


def draft_to_markup(data):
    lines = [f"# {data.get('title', '').strip()}"]
    if data.get("subtitle"):
        lines.append(f"> {data['subtitle'].strip()}")
    for section in data.get("sections", []):
        lines += ["", f"## {section.get('heading', '').strip()}"]
        lines += [p.strip() for p in section.get("paragraphs", []) if p.strip()]
        lines += [f"- {b.strip()}" for b in section.get("bullets", []) if b.strip()]
    return "\n".join(lines)


IMPROVE_ACTIONS = ("summarize", "proofread", "tone", "translate", "actions", "simplify")
MARKUP_RULES = ("Format the output as plain text using only this markup: '# ' title, '## ' headings, '- ' bullets, "
                "'1. ' numbered items and **bold**. No tables, no other markdown, no preamble.")
IMPROVE_SYSTEMS = {
    "summarize": "You are DocPilot. Write an executive summary of the document: a one-line purpose, the key points "
                 "as bullets, decisions or figures that matter, and open questions. " + MARKUP_RULES,
    "proofread": "You are DocPilot, a meticulous editor. Return the full document corrected for spelling, grammar, "
                 "punctuation and style, keeping its meaning, language and structure. After it, add a section "
                 "'## Changes' listing the main corrections. " + MARKUP_RULES,
    "tone": "You are DocPilot. Rewrite the full document in the requested tone, keeping all facts, figures, language "
            "and structure. " + MARKUP_RULES,
    "translate": "You are DocPilot, a professional business translator. Translate the full document into the requested "
                 "language with natural business style, keeping structure, names, figures and formatting. "
                 + MARKUP_RULES,
    "actions": "You are DocPilot. Extract every action item, decision and deadline from the document. Group them under "
               "'## Decisions', '## Action items' (with owner and due date when known, otherwise [to define]) and "
               "'## Open questions'. " + MARKUP_RULES,
    "simplify": "You are DocPilot. Rewrite the document in plain language: short sentences, common words, active voice, "
                "no jargon, same facts and language. " + MARKUP_RULES,
}

# ---------------------------------------------------------------- marcado → Word

_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _add_text(paragraph, text, color=None):
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        run = paragraph.add_run(m.group(1))
        run.bold = True
        if color is not None:
            run.font.color.rgb = color
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _add_field(paragraph, instr):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


PAGE_WORDS = {"es": ("Página", "de"), "en": ("Page", "of"), "fr": ("Page", "sur")}


def _hex(color):
    color = color.lstrip("#")
    return RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))


def markup_to_docx(text, path, lang="es", company="", accent="#2B579A", font="Calibri", disclaimer=""):
    """Convierte el marcado sencillo (# título, ## sección, - viñeta, 1. numerada, **negrita**) en un .docx."""
    doc = Document()
    accent_rgb = _hex(accent)
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Title", 26)):
        style = doc.styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.color.rgb = accent_rgb
        style.font.bold = True
    section = doc.sections[0]
    if company:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run(company)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page, of = PAGE_WORDS.get(lang, PAGE_WORDS["en"])
    footer.add_run(f"{page} ")
    _add_field(footer, "PAGE")
    footer.add_run(f" {of} ")
    _add_field(footer, "NUMPAGES")

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:].strip(), style="Title")
        elif line.startswith("> "):
            p = doc.add_paragraph()
            run = p.add_run(line[2:].strip())
            run.italic = True
            run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif re.match(r"^\s*[-•*] ", line):
            _add_text(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-•*] ", "", line))
        elif re.match(r"^\s*\d+[.)] ", line):
            _add_text(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)] ", "", line))
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            _add_text(p, line.strip())
    if disclaimer:
        p = doc.add_paragraph()
        run = p.add_run(disclaimer)
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
    doc.save(path)
    return path


def read_docx_text(path):
    """Texto de un .docx con marcado sencillo (títulos, viñetas y tablas como líneas)."""
    doc = Document(path)
    lines = []
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "").lower()
            if style == "title":
                lines.append(f"# {text}")
            elif style.startswith("heading"):
                lines.append(f"{'##' if style.endswith('1') else '###'} {text}")
            elif "list" in style:
                lines.append(f"- {text}")
            else:
                lines.append(text)
        elif child.tag == qn("w:tbl"):
            from docx.table import Table

            for row in Table(child, doc).rows:
                cells = []
                for cell in row.cells:
                    value = cell.text.strip()
                    if not cells or cells[-1] != value:  # celdas combinadas se repiten
                        cells.append(value)
                lines.append(" | ".join(cells))
    return "\n".join(lines)


# ---------------------------------------------------------------- legibilidad

VOWELS = {"es": "aeiouáéíóúü", "en": "aeiouy", "fr": "aeiouyàâéèêëîïôûùü"}
STOPWORDS = {
    "es": set("el la los las un una unos unas de del al a y o que en por para con sin se su sus es son lo como más "
              "pero este esta estos estas ese esa le les ya no si sí muy también entre sobre hasta desde cuando donde "
              "fue ser ha han hay nos me te mi tu".split()),
    "en": set("the a an and or of to in for on with by at from is are was were be been it its this that these those "
              "as but not no if so than then there their they we you he she i our your his her will would can could "
              "should has have had do does did".split()),
    "fr": set("le la les un une des de du d l et ou que qui en pour par avec sans se sa son ses est sont ce cette ces "
              "au aux ne pas plus mais comme sur dans il elle ils elles nous vous je tu on y a été être ont leur leurs "
              "très aussi entre".split()),
}


def count_syllables(word, lang):
    word = word.lower()
    groups = re.findall(f"[{VOWELS[lang]}]+", word)
    n = len(groups)
    if lang == "en" and n > 1 and (
            (word.endswith("e") and not word.endswith(("le", "ee")))
            or (word.endswith(("es", "ed")) and not word.endswith(("ted", "ded", "ses", "ces", "zes", "ges")))):
        n -= 1
    if lang == "fr" and word.endswith(("e", "es", "ent")) and n > 1:
        n -= 1
    return max(1, n)


def text_stats(text, lang):
    clean = re.sub(r"^[#>\-•*\d.)\s]+", "", text, flags=re.MULTILINE)
    sentences = [s for s in re.split(r"(?<=[.!?¿¡:;])\s+|\n+", clean) if re.search(r"\w", s)]
    words = re.findall(r"[^\W\d_]+(?:['’-][^\W\d_]+)*", clean)
    n_words, n_sent = len(words), max(1, len(sentences))
    syllables = sum(count_syllables(w, lang) for w in words) if words else 0
    wps = n_words / n_sent
    spw = syllables / max(1, n_words)
    if lang == "es":  # Szigriszt-Pazos (escala INFLESZ)
        score = 206.835 - 62.3 * spw - wps
    elif lang == "fr":  # Kandel y Moles
        score = 207 - 1.015 * wps - 73.6 * spw
    else:  # Flesch
        score = 206.835 - 1.015 * wps - 84.6 * spw
    freq = {}
    for w in words:
        lw = w.lower()
        if len(lw) > 3 and lw not in STOPWORDS[lang]:
            freq[lw] = freq.get(lw, 0) + 1
    repeated = sorted(((c, w) for w, c in freq.items() if c >= 3), reverse=True)[:10]
    long_sentences = [s.strip() for s in sentences if len(s.split()) > 30]
    return {
        "words": n_words, "sentences": len(sentences), "paragraphs": len([p for p in text.split("\n") if p.strip()]),
        "minutes": max(1, math.ceil(n_words / 200)) if n_words else 0, "avg_sentence": round(wps, 1),
        "score": round(max(0, min(100, score)), 1), "long": long_sentences, "repeated": repeated,
    }


def readability_level(score):
    if score >= 80:
        return "very_easy"
    if score >= 60:
        return "easy"
    if score >= 40:
        return "medium"
    return "hard"


# ---------------------------------------------------------------- combinación de correspondencia

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
MONTHS = {
    "es": ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
           "noviembre", "diciembre"],
    "en": ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
           "November", "December"],
    "fr": ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre",
           "novembre", "décembre"],
}


def format_date(value, lang):
    month = MONTHS.get(lang, MONTHS["en"])[value.month - 1]
    if lang == "es":
        return f"{value.day} de {month} de {value.year}"
    if lang == "fr":
        return f"{value.day} {month} {value.year}"
    return f"{month} {value.day}, {value.year}"


def format_number(value, lang, decimals=2):
    text = f"{value:,.{decimals}f}"
    if lang == "es":
        return text.replace(",", "X").replace(".", ",").replace("X", ".")
    if lang == "fr":
        return text.replace(",", " ").replace(".", ",")
    return text


def format_value(value, lang):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return format_date(value, lang)
    if isinstance(value, (bool, np.bool_)):
        return str(value)
    if isinstance(value, (float, np.floating)):  # importes: siempre con 2 decimales
        return format_number(float(value), lang, 2)
    if isinstance(value, np.integer):
        return str(int(value))
    return str(value)


def _iter_paragraphs(container):
    for p in container.paragraphs:
        yield p
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _all_paragraphs(doc):
    yield from _iter_paragraphs(doc)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            if part is not None and not part.is_linked_to_previous:
                yield from _iter_paragraphs(part)


def find_placeholders(path):
    doc = Document(path)
    found = []
    for p in _all_paragraphs(doc):
        for name in PLACEHOLDER_RE.findall(p.text):
            if name not in found:
                found.append(name)
    return found


def _replace_in_paragraph(paragraph, values):
    if "{{" not in paragraph.text:
        return

    def repl(m):
        key = m.group(1).strip()
        return values[key] if key in values else m.group(0)

    for run in paragraph.runs:  # primero dentro de cada fragmento (conserva el formato)
        if "{{" in run.text:
            run.text = PLACEHOLDER_RE.sub(repl, run.text)
    full = paragraph.text
    if "{{" in full and PLACEHOLDER_RE.sub(repl, full) != full:
        # El marcador quedó partido entre fragmentos: se une todo en el primero.
        new = PLACEHOLDER_RE.sub(repl, full)
        runs = paragraph.runs
        runs[0].text = new
        for run in runs[1:]:
            run.text = ""


def fill_document(template, row_values):
    doc = Document(template)
    for p in _all_paragraphs(doc):
        _replace_in_paragraph(p, row_values)
    return doc


def _safe_name(text):
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(text)).strip(" .")
    return text[:80] or "document"


def mail_merge(template, df, out_dir, lang="es", name_col=None, single_file=False, prefix=""):
    """Genera un .docx por fila (o uno solo con saltos de página). Devuelve la lista de archivos."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = [{str(k).strip(): format_value(v, lang) for k, v in row.items()} for row in df.to_dict("records")]
    if single_file:
        master = None
        for row in rows:
            doc = fill_document(template, row)
            if master is None:
                master = doc
                continue
            body = master.element.body
            sect = body.find(qn("w:sectPr"))
            br = master.add_paragraph().add_run()
            br.add_break(WD_BREAK.PAGE)
            for el in list(doc.element.body.iterchildren()):
                if el.tag == qn("w:sectPr"):
                    continue
                if sect is not None:
                    sect.addprevious(copy.deepcopy(el))
                else:
                    body.append(copy.deepcopy(el))
        path = out_dir / f"{prefix or Path(template).stem}_merged.docx"
        if master is not None:
            master.save(path)
        return [path]
    paths, used = [], set()
    for i, row in enumerate(rows, start=1):
        base = _safe_name(row.get(name_col) or f"{prefix or Path(template).stem}_{i:03d}") if name_col else \
            f"{prefix or Path(template).stem}_{i:03d}"
        name, n = base, 2
        while name.lower() in used:
            name = f"{base}_{n}"
            n += 1
        used.add(name.lower())
        path = out_dir / f"{name}.docx"
        fill_document(template, row).save(path)
        paths.append(path)
    return paths


SAMPLE_LETTER = {
    "es": ("Recordatorio de pago", "Estimado/a {{Nombre}}:",
           "Le escribimos en nombre del equipo de facturación para recordarle que la factura de **{{Producto}}** "
           "emitida a {{Empresa}} por un importe de **{{Importe}} €** vence el **{{FechaLimite}}**.",
           "Si ya realizó el pago, por favor ignore este mensaje. Para cualquier consulta, responda a este correo "
           "({{Email}}).", "Atentamente,", "Departamento de Cobros"),
    "en": ("Payment reminder", "Dear {{Name}},",
           "We are writing on behalf of the billing team to remind you that the invoice for **{{Product}}** issued to "
           "{{Company}} for **${{Amount}}** is due on **{{DueDate}}**.",
           "If you have already paid, please disregard this message. For any question, reply to this email "
           "({{Email}}).", "Kind regards,", "Accounts Receivable"),
    "fr": ("Rappel de paiement", "Madame, Monsieur {{Nom}},",
           "Nous vous écrivons au nom du service facturation pour vous rappeler que la facture de **{{Produit}}** "
           "émise à {{Entreprise}} pour un montant de **{{Montant}} €** arrive à échéance le **{{DateLimite}}**.",
           "Si vous avez déjà réglé, merci de ne pas tenir compte de ce message. Pour toute question, répondez à cet "
           "e-mail ({{Email}}).", "Cordialement,", "Service recouvrement"),
}


def sample_template(path, lang="es", company="Pilot Suite"):
    title, greeting, body, note, closing, signature = SAMPLE_LETTER[lang if lang in LANGS else "es"]
    text = "\n".join([f"# {title}", f"> {company}", greeting, body, note, closing, f"**{signature}**"])
    return markup_to_docx(text, path, lang, company)
